"""员工管理学习版后端：使用 MySQL 持久化员工、入职待办和入职档案。"""

import json
import re
import hmac
import logging
import secrets
import threading
import time
from io import BytesIO
from datetime import date, timedelta
from functools import wraps
from pathlib import Path
from urllib.parse import quote

import bcrypt
import pymysql
from flask import Blueprint, Flask, Response, after_this_request, g, jsonify, request, send_from_directory, session
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash, generate_password_hash
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from config import (
    ACCOUNT_LOCK_MINUTES, APP_DEBUG, AUTO_INITIALIZE_DATABASE, CREATE_DEMO_DATA, DATABASE_CONFIG, IS_PRODUCTION,
    MAX_CONTENT_LENGTH, MAX_IMPORT_FILE_BYTES, MAX_LOGIN_FAILURES, ROLE_NAMES,
    SECRET_KEY, SESSION_COOKIE_NAME, SESSION_COOKIE_SECURE, SESSION_LIFETIME, TRUST_PROXY_HEADERS, AGENT_API_KEY,
    APP_PORT,
)
from services.employee_service import OnboardingPolicy, find_employee, normalize_list, public_employee, serialize_record
from services.ai_assistant import (
    ask_assistant, delete_conversation, get_conversation_messages, list_conversations,
    save_conversation_exchange, stream_assistant, update_conversation_summary, update_conversation_title,
)
from services.database_query import DatabaseQueryError, run_database_intent, run_database_query
from services.notification_service import create_notifications, serialize_notification
from services.schedule_rules import find_minimum_rest_violation
from database.bootstrap import DEFAULT_DEPARTMENTS, seed_department_demo_accounts
from utils.schedule_import_reader import ScheduleImportReader
from utils.mysql_storage import connect_database, execute, query_all, query_one


def create_app():
    """创建 Flask 应用；配置集中定义在 config.py。"""
    application = Flask(__name__)
    application.config.update(
        SECRET_KEY=SECRET_KEY,
        SESSION_COOKIE_NAME=SESSION_COOKIE_NAME,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=SESSION_COOKIE_SECURE,
        PERMANENT_SESSION_LIFETIME=SESSION_LIFETIME,
        MAX_CONTENT_LENGTH=MAX_CONTENT_LENGTH,
    )
    if TRUST_PROXY_HEADERS:
        application.wsgi_app = ProxyFix(application.wsgi_app, x_for=1, x_proto=1, x_host=1)
    return application


app = create_app()
app.logger.setLevel(logging.INFO)
api_bp = Blueprint("api", __name__)
PHONE_PATTERN = re.compile(r"^1\d{10}$")
PASSWORD_PATTERN = re.compile(r"^(?=.*[a-z])(?=.*[A-Z])(?=.*\d)(?=.*[^\w\s]).{8,}$")
REQUEST_ID_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{7,63}$")
schedule_import_reader = ScheduleImportReader(MAX_IMPORT_FILE_BYTES)
OFFICE_TOKEN_MAX_AGE_SECONDS = 300

# 注册员工的待确认数据只保存在服务端内存中，避免把明文密码写入 Flask 客户端会话。
_assistant_registration_lock = threading.Lock()
_assistant_pending_registrations = {}
_ASSISTANT_REGISTRATION_TTL_SECONDS = 300

# 便于本地直接访问 Flask 的 5000 端口；生产环境仍由 Nginx 提供前端静态文件。
FRONTEND_DIR = Path(__file__).resolve().parent.parent / "frontend"


@app.route("/")
def frontend_index():
    return send_from_directory(FRONTEND_DIR, "login.html")


@app.route("/<path:asset_path>")
def frontend_asset(asset_path):
    if asset_path.startswith(("api/",)):
        return jsonify({"message": "接口不存在"}), 404
    return send_from_directory(FRONTEND_DIR, asset_path)

def initialize_database():
    """创建业务表并执行兼容迁移；仅在数据库尚不存在时尝试建库。"""
    try:
        with connect_database():
            pass
    except pymysql.err.OperationalError as error:
        if error.args[0] != 1049:
            raise
        with connect_database(False) as connection:
            with connection.cursor() as cursor:
                cursor.execute(f"CREATE DATABASE `{DATABASE_CONFIG['database']}` CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci")
    statements = [
        """CREATE TABLE IF NOT EXISTS employees (
            id INT PRIMARY KEY AUTO_INCREMENT, name VARCHAR(50) NOT NULL,
            employee_no VARCHAR(50) NOT NULL UNIQUE, password VARCHAR(255) NOT NULL,
            phone VARCHAR(30) NOT NULL, email VARCHAR(100) NOT NULL DEFAULT '',
            department VARCHAR(100) NOT NULL DEFAULT '', position VARCHAR(100) NOT NULL DEFAULT '',
            role VARCHAR(30) NOT NULL DEFAULT 'employee',
            login_failed_attempts INT NOT NULL DEFAULT 0, locked_until DATETIME NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS onboarding_records (
            record_id INT PRIMARY KEY AUTO_INCREMENT, employee_id INT NOT NULL,
            employee_name VARCHAR(50) NOT NULL, department VARCHAR(100) NOT NULL,
            position VARCHAR(100) NOT NULL, onboarding_type VARCHAR(100) NOT NULL,
            priority VARCHAR(20) NOT NULL, queue_number INT NOT NULL UNIQUE,
            status VARCHAR(20) NOT NULL DEFAULT '待办理', created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            CONSTRAINT fk_onboarding_employee FOREIGN KEY (employee_id) REFERENCES employees(id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS employee_records (
            record_id INT PRIMARY KEY AUTO_INCREMENT, employee_id INT NOT NULL UNIQUE,
            employee_name VARCHAR(50) NOT NULL, department VARCHAR(100) NOT NULL DEFAULT '',
            position VARCHAR(100) NOT NULL DEFAULT '', role VARCHAR(30) NOT NULL,
            training_items JSON NOT NULL, assets JSON NOT NULL, remark TEXT NOT NULL,
            status VARCHAR(20) NOT NULL DEFAULT '已归档', created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            CONSTRAINT fk_employee_record_employee FOREIGN KEY (employee_id) REFERENCES employees(id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS employee_tasks (
            task_id INT PRIMARY KEY AUTO_INCREMENT, employee_id INT NOT NULL,
            task_type VARCHAR(30) NOT NULL, title VARCHAR(100) NOT NULL, content TEXT NOT NULL,
            status VARCHAR(20) NOT NULL DEFAULT '待处理', priority VARCHAR(20) NOT NULL DEFAULT '普通',
            due_date DATE NULL, publisher_id INT NULL, client_request_id VARCHAR(64) NULL, published_at TIMESTAMP NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            submitted_at TIMESTAMP NULL, completed_at TIMESTAMP NULL,
            FOREIGN KEY (employee_id) REFERENCES employees(id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS admin_todos (
            todo_id INT PRIMARY KEY AUTO_INCREMENT, task_id INT NOT NULL UNIQUE, employee_id INT NOT NULL,
            employee_name VARCHAR(50) NOT NULL, title VARCHAR(100) NOT NULL, content TEXT NOT NULL,
            status VARCHAR(20) NOT NULL DEFAULT '待确认', created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            reviewed_at TIMESTAMP NULL, FOREIGN KEY (task_id) REFERENCES employee_tasks(task_id),
            FOREIGN KEY (employee_id) REFERENCES employees(id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS departments (
            department_id INT PRIMARY KEY AUTO_INCREMENT, name VARCHAR(100) NOT NULL UNIQUE,
            parent_id INT NULL, description VARCHAR(255) NOT NULL DEFAULT '',
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (parent_id) REFERENCES departments(department_id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS shift_types (
            shift_id INT PRIMARY KEY AUTO_INCREMENT, name VARCHAR(100) NOT NULL UNIQUE,
            start_time TIME NOT NULL, end_time TIME NOT NULL, duration_hours DECIMAL(4,2) NOT NULL,
            color VARCHAR(20) NOT NULL DEFAULT '#2563eb', is_active TINYINT(1) NOT NULL DEFAULT 1,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS work_calendar (
            work_date DATE PRIMARY KEY, is_workday TINYINT(1) NOT NULL DEFAULT 1,
            label VARCHAR(100) NOT NULL DEFAULT '', updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS scheduling_rules (
            rule_id INT PRIMARY KEY AUTO_INCREMENT, rule_name VARCHAR(100) NOT NULL UNIQUE,
            max_consecutive_days INT NOT NULL DEFAULT 6, max_daily_overtime_hours DECIMAL(4,2) NOT NULL DEFAULT 3,
            max_monthly_overtime_hours DECIMAL(5,2) NOT NULL DEFAULT 36, min_staff_per_position INT NOT NULL DEFAULT 1,
            min_rest_hours DECIMAL(4,2) NOT NULL DEFAULT 10,
            updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS schedule_plans (
            plan_id INT PRIMARY KEY AUTO_INCREMENT, name VARCHAR(100) NOT NULL,
            department VARCHAR(100) NOT NULL DEFAULT '', start_date DATE NOT NULL, end_date DATE NOT NULL,
            cycle_id INT NULL,
            plan_key VARCHAR(255) NULL UNIQUE, version INT NOT NULL DEFAULT 1,
            status VARCHAR(20) NOT NULL DEFAULT '草稿', created_by INT NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
            published_at TIMESTAMP NULL,
            FOREIGN KEY (created_by) REFERENCES employees(id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS schedule_assignments (
            assignment_id INT PRIMARY KEY AUTO_INCREMENT, plan_id INT NOT NULL, employee_id INT NOT NULL,
            work_date DATE NOT NULL, shift_id INT NOT NULL, created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE KEY uq_plan_employee_date (plan_id, employee_id, work_date),
            FOREIGN KEY (plan_id) REFERENCES schedule_plans(plan_id),
            FOREIGN KEY (employee_id) REFERENCES employees(id),
            FOREIGN KEY (shift_id) REFERENCES shift_types(shift_id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS schedule_templates (
            template_id BIGINT PRIMARY KEY AUTO_INCREMENT, department VARCHAR(100) NOT NULL,
            name VARCHAR(100) NOT NULL, template_json JSON NOT NULL, created_by INT NOT NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
            INDEX idx_schedule_templates_owner (created_by, department, updated_at),
            FOREIGN KEY (created_by) REFERENCES employees(id) ON DELETE CASCADE
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS schedule_plan_versions (
            snapshot_id BIGINT PRIMARY KEY AUTO_INCREMENT, plan_id INT NOT NULL,
            plan_version INT NOT NULL, action VARCHAR(30) NOT NULL, status VARCHAR(20) NOT NULL,
            actor_id INT NULL, note VARCHAR(500) NOT NULL DEFAULT '', snapshot_json JSON NOT NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            INDEX idx_plan_versions (plan_id, plan_version, created_at),
            FOREIGN KEY (plan_id) REFERENCES schedule_plans(plan_id) ON DELETE CASCADE,
            FOREIGN KEY (actor_id) REFERENCES employees(id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS scheduling_cycles (
            cycle_id INT PRIMARY KEY AUTO_INCREMENT, start_date DATE NOT NULL UNIQUE, end_date DATE NOT NULL,
            start_day TINYINT NOT NULL, cycle_kind VARCHAR(20) NOT NULL DEFAULT 'standard',
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE KEY uq_cycle_range (start_date, end_date)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS scheduling_cycle_change_requests (
            request_id INT PRIMARY KEY AUTO_INCREMENT, proposed_start_day TINYINT NOT NULL,
            status VARCHAR(20) NOT NULL DEFAULT '待老板审批', requested_by INT NOT NULL,
            requested_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            reviewed_by INT NULL, reviewed_at TIMESTAMP NULL, review_note VARCHAR(500) NOT NULL DEFAULT '',
            transition_start DATE NULL, effective_start DATE NULL,
            FOREIGN KEY (requested_by) REFERENCES employees(id),
            FOREIGN KEY (reviewed_by) REFERENCES employees(id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS schedule_requests (
            request_id INT PRIMARY KEY AUTO_INCREMENT, employee_id INT NOT NULL,
            department VARCHAR(100) NOT NULL, request_type VARCHAR(20) NOT NULL,
            work_date DATE NOT NULL, desired_shift_id INT NULL, reason VARCHAR(500) NOT NULL,
            assignment_id INT NULL, original_shift_id INT NULL, assignment_version INT NOT NULL DEFAULT 0,
            status VARCHAR(30) NOT NULL DEFAULT '待部门审批',
            manager_id INT NULL, manager_note VARCHAR(500) NOT NULL DEFAULT '', manager_reviewed_at TIMESTAMP NULL,
            admin_id INT NULL, admin_note VARCHAR(500) NOT NULL DEFAULT '', admin_reviewed_at TIMESTAMP NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (employee_id) REFERENCES employees(id),
            FOREIGN KEY (desired_shift_id) REFERENCES shift_types(shift_id),
            FOREIGN KEY (original_shift_id) REFERENCES shift_types(shift_id),
            FOREIGN KEY (manager_id) REFERENCES employees(id),
            FOREIGN KEY (admin_id) REFERENCES employees(id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS notifications (
            notification_id BIGINT PRIMARY KEY AUTO_INCREMENT,
            recipient_id INT NOT NULL, notification_type VARCHAR(40) NOT NULL DEFAULT 'system',
            title VARCHAR(120) NOT NULL, content VARCHAR(500) NOT NULL,
            target_url VARCHAR(255) NOT NULL DEFAULT '/pages/profile.html',
            related_type VARCHAR(40) NOT NULL DEFAULT '', related_id BIGINT NULL,
            dedupe_key VARCHAR(190) NULL, is_read TINYINT(1) NOT NULL DEFAULT 0,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP, read_at TIMESTAMP NULL,
            INDEX idx_notifications_recipient (recipient_id,is_read,created_at),
            UNIQUE KEY uq_notifications_dedupe (recipient_id,dedupe_key),
            FOREIGN KEY (recipient_id) REFERENCES employees(id) ON DELETE CASCADE
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS audit_logs (
            audit_id BIGINT PRIMARY KEY AUTO_INCREMENT,
            request_id VARCHAR(64) NOT NULL DEFAULT '',
            actor_id INT NULL, actor_employee_no VARCHAR(50) NOT NULL DEFAULT 'anonymous',
            action VARCHAR(120) NOT NULL, http_method VARCHAR(10) NOT NULL,
            request_path VARCHAR(255) NOT NULL, status_code SMALLINT NOT NULL,
            client_ip VARCHAR(45) NOT NULL DEFAULT '', user_agent VARCHAR(255) NOT NULL DEFAULT '',
            details TEXT NOT NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            INDEX idx_audit_created_at (created_at), INDEX idx_audit_request (request_id), INDEX idx_audit_actor (actor_id, created_at),
            INDEX idx_audit_action (action, created_at)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS ai_usage_logs (
            usage_id BIGINT PRIMARY KEY AUTO_INCREMENT, employee_id INT NOT NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            INDEX idx_ai_usage_employee_time (employee_id, created_at),
            FOREIGN KEY (employee_id) REFERENCES employees(id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS ai_conversations (
            conversation_id CHAR(36) PRIMARY KEY, employee_id INT NOT NULL,
            title VARCHAR(100) NOT NULL DEFAULT '新对话', prompt_version VARCHAR(40) NOT NULL,
            memory_type VARCHAR(20) NOT NULL DEFAULT 'buffer',
            memory_summary TEXT NOT NULL,
            memory_facts TEXT NOT NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
            INDEX idx_ai_conversation_employee_time (employee_id, updated_at),
            FOREIGN KEY (employee_id) REFERENCES employees(id) ON DELETE CASCADE
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
        """CREATE TABLE IF NOT EXISTS ai_conversation_messages (
            message_id BIGINT PRIMARY KEY AUTO_INCREMENT, conversation_id CHAR(36) NOT NULL,
            role VARCHAR(12) NOT NULL, content TEXT NOT NULL,
            output_format VARCHAR(10) NOT NULL DEFAULT '', prompt_version VARCHAR(40) NOT NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            INDEX idx_ai_message_conversation (conversation_id, message_id),
            FOREIGN KEY (conversation_id) REFERENCES ai_conversations(conversation_id) ON DELETE CASCADE
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""",
    ]
    with connect_database() as connection:
        with connection.cursor() as cursor:
            for statement in statements:
                cursor.execute(statement)
            cursor.execute("SHOW COLUMNS FROM employees LIKE 'employment_status'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE employees ADD COLUMN employment_status VARCHAR(20) NOT NULL DEFAULT '待分配' AFTER role")
                cursor.execute("UPDATE employees SET employment_status = '在职' WHERE department <> '' OR position <> ''")
            cursor.execute("SHOW COLUMNS FROM employees LIKE 'login_failed_attempts'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE employees ADD COLUMN login_failed_attempts INT NOT NULL DEFAULT 0 AFTER role")
            cursor.execute("SHOW COLUMNS FROM employees LIKE 'locked_until'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE employees ADD COLUMN locked_until DATETIME NULL AFTER login_failed_attempts")
            cursor.execute("SHOW COLUMNS FROM schedule_plans LIKE 'submitted_at'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE schedule_plans ADD COLUMN submitted_at TIMESTAMP NULL AFTER published_at")
                cursor.execute("ALTER TABLE schedule_plans ADD COLUMN hr_id INT NULL AFTER submitted_at")
                cursor.execute("ALTER TABLE schedule_plans ADD COLUMN hr_note VARCHAR(500) NOT NULL DEFAULT '' AFTER hr_id")
            cursor.execute("SHOW COLUMNS FROM schedule_plans LIKE 'cycle_id'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE schedule_plans ADD COLUMN cycle_id INT NULL AFTER end_date")
            cursor.execute("SHOW COLUMNS FROM schedule_plans LIKE 'updated_at'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE schedule_plans ADD COLUMN updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP AFTER created_at")
            cursor.execute("SHOW COLUMNS FROM schedule_plans LIKE 'plan_key'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE schedule_plans ADD COLUMN plan_key VARCHAR(255) NULL AFTER cycle_id")
                cursor.execute("SELECT plan_id,department,start_date,end_date FROM schedule_plans ORDER BY plan_id DESC")
                active_keys = set()
                for plan in cursor.fetchall():
                    plan_key = schedule_plan_key(plan["department"], plan["start_date"], plan["end_date"])
                    if plan_key in active_keys:
                        continue
                    cursor.execute("UPDATE schedule_plans SET plan_key=%s WHERE plan_id=%s", (plan_key, plan["plan_id"]))
                    active_keys.add(plan_key)
                cursor.execute("CREATE UNIQUE INDEX uq_schedule_plans_plan_key ON schedule_plans(plan_key)")
            cursor.execute("SHOW COLUMNS FROM schedule_plans LIKE 'version'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE schedule_plans ADD COLUMN version INT NOT NULL DEFAULT 1 AFTER plan_key")
            cursor.execute("SHOW COLUMNS FROM schedule_plans LIKE 'force_submit_reason'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE schedule_plans ADD COLUMN force_submit_reason VARCHAR(500) NOT NULL DEFAULT '' AFTER hr_note")
                cursor.execute("ALTER TABLE schedule_plans ADD COLUMN force_submitted_by INT NULL AFTER force_submit_reason")
                cursor.execute("ALTER TABLE schedule_plans ADD COLUMN force_submitted_at TIMESTAMP NULL AFTER force_submitted_by")
            cursor.execute("SHOW COLUMNS FROM scheduling_rules LIKE 'min_rest_hours'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE scheduling_rules ADD COLUMN min_rest_hours DECIMAL(4,2) NOT NULL DEFAULT 10 AFTER min_staff_per_position")
            cursor.execute("SHOW COLUMNS FROM schedule_requests LIKE 'assignment_id'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE schedule_requests ADD COLUMN assignment_id INT NULL AFTER reason")
                cursor.execute("ALTER TABLE schedule_requests ADD COLUMN original_shift_id INT NULL AFTER assignment_id")
                cursor.execute("ALTER TABLE schedule_requests ADD COLUMN assignment_version INT NOT NULL DEFAULT 0 AFTER original_shift_id")
                cursor.execute("CREATE INDEX idx_schedule_requests_assignment ON schedule_requests(assignment_id)")
            cursor.execute(
                """SELECT CONSTRAINT_NAME FROM information_schema.KEY_COLUMN_USAGE
                   WHERE TABLE_SCHEMA=%s AND TABLE_NAME='schedule_requests'
                     AND COLUMN_NAME='assignment_id' AND REFERENCED_TABLE_NAME='schedule_assignments'""",
                (DATABASE_CONFIG["database"],),
            )
            assignment_foreign_key = cursor.fetchone()
            if assignment_foreign_key:
                cursor.execute(f"ALTER TABLE schedule_requests DROP FOREIGN KEY `{assignment_foreign_key['CONSTRAINT_NAME']}`")
            cursor.execute("SHOW COLUMNS FROM audit_logs LIKE 'request_id'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE audit_logs ADD COLUMN request_id VARCHAR(64) NOT NULL DEFAULT '' AFTER audit_id")
                cursor.execute("CREATE INDEX idx_audit_request ON audit_logs(request_id)")
            cursor.execute("SHOW COLUMNS FROM audit_logs LIKE 'details'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE audit_logs ADD COLUMN details TEXT NOT NULL AFTER user_agent")
            cursor.execute("SHOW COLUMNS FROM ai_conversations LIKE 'memory_type'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE ai_conversations ADD COLUMN memory_type VARCHAR(20) NOT NULL DEFAULT 'buffer' AFTER prompt_version")
            cursor.execute("SHOW COLUMNS FROM ai_conversations LIKE 'memory_summary'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE ai_conversations ADD COLUMN memory_summary TEXT NULL AFTER memory_type")
                cursor.execute("UPDATE ai_conversations SET memory_summary = '' WHERE memory_summary IS NULL")
                cursor.execute("ALTER TABLE ai_conversations MODIFY COLUMN memory_summary TEXT NOT NULL")
            cursor.execute("SHOW COLUMNS FROM ai_conversations LIKE 'memory_facts'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE ai_conversations ADD COLUMN memory_facts TEXT NULL AFTER memory_summary")
                cursor.execute("UPDATE ai_conversations SET memory_facts = '' WHERE memory_facts IS NULL")
                cursor.execute("ALTER TABLE ai_conversations MODIFY COLUMN memory_facts TEXT NOT NULL")
            for column, definition in (
                ("priority", "VARCHAR(20) NOT NULL DEFAULT '普通' AFTER status"),
                ("due_date", "DATE NULL AFTER priority"),
                ("publisher_id", "INT NULL AFTER due_date"),
                ("client_request_id", "VARCHAR(64) NULL AFTER publisher_id"),
                ("published_at", "TIMESTAMP NULL AFTER publisher_id"),
            ):
                cursor.execute(f"SHOW COLUMNS FROM employee_tasks LIKE '{column}'")
                if not cursor.fetchone():
                    cursor.execute(f"ALTER TABLE employee_tasks ADD COLUMN {column} {definition}")
            cursor.execute(
                """SELECT COLUMN_NAME FROM information_schema.statistics
                   WHERE table_schema=%s AND table_name='employee_tasks' AND index_name='uq_employee_tasks_request'
                   ORDER BY SEQ_IN_INDEX""",
                (DATABASE_CONFIG["database"],),
            )
            request_index_columns = [row["COLUMN_NAME"] for row in cursor.fetchall()]
            expected_request_index_columns = ["publisher_id", "client_request_id", "employee_id"]
            if request_index_columns and request_index_columns != expected_request_index_columns:
                cursor.execute("DROP INDEX uq_employee_tasks_request ON employee_tasks")
            if request_index_columns != expected_request_index_columns:
                cursor.execute(
                    "CREATE UNIQUE INDEX uq_employee_tasks_request "
                    "ON employee_tasks(publisher_id,client_request_id,employee_id)"
                )
            cursor.execute("""INSERT IGNORE INTO scheduling_rules
                (rule_name, max_consecutive_days, max_daily_overtime_hours, max_monthly_overtime_hours, min_staff_per_position)
                VALUES ('默认排班规则', 6, 3, 36, 1)""")
            for department_name in DEFAULT_DEPARTMENTS:
                cursor.execute("INSERT IGNORE INTO departments (name) VALUES (%s)", (department_name,))
            for shift in (("早班", "08:00", "17:00", 8, "#2563eb"), ("中班", "10:00", "19:00", 8, "#0f9f6e"), ("晚班", "14:00", "23:00", 8, "#8b5cf6")):
                cursor.execute("""INSERT IGNORE INTO shift_types (name,start_time,end_time,duration_hours,color)
                               VALUES (%s,%s,%s,%s,%s)""", shift)
            # 演示账号仅允许在显式开发模式创建，生产环境绝不写入固定凭据。
            if CREATE_DEMO_DATA and not IS_PRODUCTION:
                seed_department_demo_accounts(cursor)
    ensure_schedule_cycles(date.today() + timedelta(days=1095))


def add_months(value, months):
    """Return the first day of a month offset without relying on locale-specific parsing."""
    absolute_month = value.year * 12 + value.month - 1 + months
    return date(absolute_month // 12, absolute_month % 12 + 1, 1)


def cycle_end_for_start(start, start_day):
    return add_months(start.replace(day=1), 1).replace(day=start_day) - timedelta(days=1)


def ensure_schedule_cycles(until_date):
    """Keep a rolling calendar. The default policy is the 21st through the following 20th."""
    last = query_one("SELECT * FROM scheduling_cycles ORDER BY end_date DESC LIMIT 1")
    if last:
        start = last["end_date"] + timedelta(days=1)
        start_day = int(last["start_day"])
    else:
        start, start_day = date(2000, 1, 21), 21
    while start <= until_date:
        end = cycle_end_for_start(start, start_day)
        execute("""INSERT IGNORE INTO scheduling_cycles (start_date,end_date,start_day,cycle_kind)
                   VALUES (%s,%s,%s,'standard')""", (start, end, start_day))
        start = end + timedelta(days=1)


def get_schedule_cycle(anchor_date):
    ensure_schedule_cycles(anchor_date + timedelta(days=1095))
    return query_one("""SELECT * FROM scheduling_cycles
                        WHERE start_date<=%s AND end_date>=%s ORDER BY start_date DESC LIMIT 1""",
                     (anchor_date, anchor_date))


def validate_cycle_range(start_date, end_date):
    cycle = get_schedule_cycle(start_date)
    if not cycle or cycle["start_date"] != start_date or cycle["end_date"] != end_date:
        return None, "班表必须完整覆盖一个已生效的排班周期，请通过“周期定位日”加载系统周期。"
    return cycle, None


def next_day_of_month_on_or_after(value, day):
    candidate = value.replace(day=day) if value.day <= day else add_months(value.replace(day=1), 1).replace(day=day)
    return candidate


def cycle_change_preview(start_day):
    current = get_schedule_cycle(date.today())
    transition_start = current["end_date"] + timedelta(days=1)
    effective_start = next_day_of_month_on_or_after(transition_start, start_day)
    return current, transition_start, effective_start


def serialize_cycle(cycle):
    if not cycle:
        return None
    result = cycle.copy()
    for field in ("start_date", "end_date", "created_at", "updated_at"):
        if hasattr(result.get(field), "isoformat"):
            result[field] = result[field].isoformat()
    return result


def serialize_cycle_request(record):
    if not record:
        return None
    result = record.copy()
    for field in ("requested_at", "reviewed_at", "transition_start", "effective_start"):
        if hasattr(result.get(field), "isoformat"):
            result[field] = result[field].isoformat()
    return result


def validate_password(password):
    if not isinstance(password, str) or not PASSWORD_PATTERN.fullmatch(password):
        return "密码至少 8 位，且必须包含大写字母、小写字母、数字和特殊字符"
    return None


def validate_registration_data(data):
    """统一清洗和校验员工开户数据，避免路由中散落重复校验。"""
    required_fields = ("name", "employee_no", "password", "confirm_password", "phone")
    if any(not str(data.get(field, "")).strip() for field in required_fields):
        return None, "姓名、工号、密码、确认密码和手机号不能为空"
    employee = {
        "name": str(data["name"]).strip(),
        "employee_no": str(data["employee_no"]).strip(),
        "password": data["password"],
        "phone": str(data["phone"]).strip(),
        "email": str(data.get("email", "")).strip(),
    }
    if not re.fullmatch(r"EMP\d+", employee["employee_no"]):
        return None, "工号格式必须为 EMP 加数字，例如 EMP006"
    if employee["password"] != data["confirm_password"]:
        return None, "两次输入的密码不一致"
    password_error = validate_password(employee["password"])
    if password_error:
        return None, password_error
    if not PHONE_PATTERN.fullmatch(employee["phone"]):
        return None, "手机号格式无效，请输入 11 位手机号"
    if employee["email"] and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", employee["email"]):
        return None, "邮箱格式无效"
    return employee, None


def account_is_locked(employee):
    return bool(employee["locked_until"] and query_one("SELECT NOW() < %s AS is_locked", (employee["locked_until"],))["is_locked"])


def record_login_failure(employee_id):
    """累计失败次数；第 5 次失败起锁定账户 30 分钟。"""
    execute(
        """UPDATE employees
           SET locked_until = CASE WHEN login_failed_attempts + 1 >= %s
                                   THEN DATE_ADD(NOW(), INTERVAL %s MINUTE)
                                   ELSE locked_until END,
               login_failed_attempts = login_failed_attempts + 1
           WHERE id = %s""",
        (MAX_LOGIN_FAILURES, ACCOUNT_LOCK_MINUTES, employee_id),
    )


def reset_login_failures(employee_id):
    execute("UPDATE employees SET login_failed_attempts = 0, locked_until = NULL WHERE id = %s", (employee_id,))


def parse_iso_date(value, field_name="日期"):
    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError):
        return None


def validate_shift_data(data):
    name = str(data.get("name", "")).strip()
    start_time = str(data.get("start_time", "")).strip()
    end_time = str(data.get("end_time", "")).strip()
    color = str(data.get("color", "#2563eb")).strip()
    if not name or not start_time or not end_time:
        return None, "班次名称、开始时间和结束时间不能为空"
    if not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", start_time) or not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", end_time):
        return None, "班次时间格式必须为 HH:MM"
    if not re.fullmatch(r"#[0-9A-Fa-f]{6}", color):
        return None, "班次颜色必须为 #RRGGBB 格式"
    try:
        duration_hours = float(data.get("duration_hours"))
    except (TypeError, ValueError):
        return None, "班次时长必须是数字"
    if not 0 < duration_hours <= 24:
        return None, "班次时长必须在 0 到 24 小时之间"
    start_minutes = int(start_time[:2]) * 60 + int(start_time[3:])
    end_minutes = int(end_time[:2]) * 60 + int(end_time[3:])
    elapsed_hours = ((end_minutes - start_minutes) % (24 * 60)) / 60
    if elapsed_hours == 0:
        elapsed_hours = 24
    if duration_hours > elapsed_hours:
        return None, f"计薪时长不能超过班次跨度 {elapsed_hours:g} 小时；跨夜班次请保持结束时间早于开始时间"
    return {"name": name, "start_time": start_time, "end_time": end_time, "duration_hours": duration_hours, "color": color}, None


def serialize_shift(shift):
    if shift is None:
        return None
    result = shift.copy()
    for field in ("start_time", "end_time"):
        value = result[field]
        if hasattr(value, "total_seconds"):
            total_seconds = int(value.total_seconds())
            result[field] = f"{total_seconds // 3600:02d}:{total_seconds % 3600 // 60:02d}"
    result["duration_hours"] = float(result["duration_hours"])
    return result


def get_plan_assignments(plan_id, serialize_dates=False):
    assignments = query_all(
        """SELECT a.assignment_id, a.plan_id, a.employee_id, e.name AS employee_name,
                  e.department, e.position, a.work_date, s.shift_id, s.name AS shift_name,
                  s.start_time, s.end_time, s.duration_hours, s.color,
                  EXISTS(SELECT 1 FROM schedule_requests r
                         WHERE r.employee_id=a.employee_id AND r.work_date=a.work_date
                           AND r.request_type='leave' AND r.status='已批准') AS is_on_leave
           FROM schedule_assignments a
           JOIN employees e ON e.id=a.employee_id JOIN shift_types s ON s.shift_id=a.shift_id
           WHERE a.plan_id=%s ORDER BY a.work_date, a.employee_id""", (plan_id,)
    )
    for assignment in assignments:
        if serialize_dates and hasattr(assignment.get("work_date"), "isoformat"):
            assignment["work_date"] = assignment["work_date"].isoformat()
        for field in ("start_time", "end_time"):
            value = assignment[field]
            if hasattr(value, "total_seconds"):
                total_seconds = int(value.total_seconds())
                assignment[field] = f"{total_seconds // 3600:02d}:{total_seconds % 3600 // 60:02d}"
        assignment["duration_hours"] = float(assignment["duration_hours"])
        assignment["is_on_leave"] = bool(assignment.get("is_on_leave"))
    return assignments


def validate_plan_assignments(plan, assignments, allow_existing=False, allow_consecutive_warning=False):
    """校验同日多班、最短班间休息、跨周期连续工作天数与所属部门。"""
    if not isinstance(assignments, list) or not assignments:
        return None, "assignments 必须是非空数组"
    plan_start = date.fromisoformat(plan["start_date"].isoformat() if isinstance(plan["start_date"], date) else str(plan["start_date"]))
    plan_end = date.fromisoformat(plan["end_date"].isoformat() if isinstance(plan["end_date"], date) else str(plan["end_date"]))
    _, cycle_error = validate_cycle_range(plan_start, plan_end)
    if cycle_error:
        return None, cycle_error
    rule = query_one("SELECT * FROM scheduling_rules WHERE rule_name='默认排班规则'")
    cleaned, employee_dates, shift_cache, employee_names = [], {}, {}, {}
    existing_keys = {(row["employee_id"], row["work_date"]) for row in get_plan_assignments(plan["plan_id"])}
    for item in assignments:
        if not isinstance(item, dict): return None, "每条排班记录必须是对象"
        try:
            employee_id, shift_id = int(item.get("employee_id")), int(item.get("shift_id"))
        except (TypeError, ValueError): return None, "employee_id 和 shift_id 必须是有效编号"
        work_date = parse_iso_date(item.get("work_date"))
        if not work_date or not plan_start <= work_date <= plan_end:
            return None, "排班日期必须在计划周期内"
        employee = find_employee(employee_id)
        if employee is not None and employee.get("employment_status") != "在职":
            return None, f"员工 {employee['name']} 当前不是在职状态，不能排班"
        shift = query_one("SELECT * FROM shift_types WHERE shift_id=%s AND is_active=1", (shift_id,))
        if employee is None: return None, f"员工 {employee_id} 不存在"
        if shift is None: return None, f"班次 {shift_id} 不存在或已停用"
        shift_cache[shift_id] = shift
        employee_names[employee_id] = employee["name"]
        if plan["department"] and employee["department"] != plan["department"]:
            return None, f"员工 {employee['name']} 不属于计划部门"
        key = (employee_id, work_date)
        if key in employee_dates: return None, f"员工 {employee['name']} 在 {work_date} 已安排多个班次"
        if not allow_existing and key in existing_keys: return None, f"员工 {employee['name']} 在 {work_date} 已有排班"
        employee_dates[key] = True
        cleaned.append({
            "employee_id": employee_id, "shift_id": shift_id, "work_date": work_date,
            "is_on_leave": bool(item.get("is_on_leave")),
        })

    for item in cleaned:
        collision = query_one("""SELECT p.plan_id FROM schedule_assignments a JOIN schedule_plans p ON p.plan_id=a.plan_id
                                 WHERE a.employee_id=%s AND a.work_date=%s AND p.plan_id<>%s AND p.status='已发布' LIMIT 1""",
                              (item["employee_id"], item["work_date"], plan["plan_id"]))
        if collision:
            return None, "员工在该日期已有已发布班次，不能产生跨部门或跨版本的重复排班"

    existing = get_plan_assignments(plan["plan_id"])
    all_dates = {}
    for row in existing:
        if not row.get("is_on_leave"):
            all_dates.setdefault(row["employee_id"], set()).add(row["work_date"])
    for item in cleaned:
        if not item.get("is_on_leave"):
            all_dates.setdefault(item["employee_id"], set()).add(item["work_date"])
    employee_ids = sorted(all_dates)
    if employee_ids:
        context_days = int(rule["max_consecutive_days"]) + 1
        placeholders = ",".join(["%s"] * len(employee_ids))
        adjacent = query_all(
            f"""SELECT a.employee_id,a.work_date FROM schedule_assignments a
                 JOIN schedule_plans p ON p.plan_id=a.plan_id
                 LEFT JOIN (SELECT employee_id,work_date FROM schedule_requests
                            WHERE request_type='leave' AND status='已批准' GROUP BY employee_id,work_date) lr
                   ON lr.employee_id=a.employee_id AND lr.work_date=a.work_date
                 WHERE a.employee_id IN ({placeholders}) AND p.plan_id<>%s AND p.status='已发布'
                   AND lr.employee_id IS NULL AND a.work_date BETWEEN %s AND %s""",
            [*employee_ids, plan["plan_id"], plan_start - timedelta(days=context_days), plan_end + timedelta(days=context_days)],
        )
        for row in adjacent:
            all_dates.setdefault(row["employee_id"], set()).add(row["work_date"])
    for employee_id, dates in all_dates.items():
        consecutive = 0
        for work_date in sorted(dates):
            consecutive = consecutive + 1 if work_date - timedelta(days=1) in dates else 1
            if consecutive > rule["max_consecutive_days"] and not allow_consecutive_warning:
                employee = find_employee(employee_id)
                return None, f"员工 {employee['name']} 连续工作超过 {rule['max_consecutive_days']} 天"

    minimum_rest_hours = float(rule.get("min_rest_hours") or 0)
    if minimum_rest_hours > 0 and cleaned:
        rest_assignments = []
        if not allow_existing:
            rest_assignments.extend({**item, "_candidate": False} for item in existing if not item.get("is_on_leave"))
        for item in cleaned:
            if item.get("is_on_leave"):
                continue
            shift = shift_cache[item["shift_id"]]
            rest_assignments.append({
                **item,
                "employee_name": employee_names[item["employee_id"]],
                "shift_name": shift["name"],
                "start_time": shift["start_time"],
                "end_time": shift["end_time"],
                "_candidate": True,
            })
        employee_ids = sorted({item["employee_id"] for item in cleaned})
        placeholders = ",".join(["%s"] * len(employee_ids))
        published_assignments = query_all(
            f"""SELECT a.employee_id,e.name AS employee_name,a.work_date,
                       s.name AS shift_name,s.start_time,s.end_time
                FROM schedule_assignments a
                JOIN schedule_plans p ON p.plan_id=a.plan_id
                JOIN employees e ON e.id=a.employee_id
                JOIN shift_types s ON s.shift_id=a.shift_id
                LEFT JOIN (SELECT employee_id,work_date FROM schedule_requests
                           WHERE request_type='leave' AND status='已批准' GROUP BY employee_id,work_date) lr
                  ON lr.employee_id=a.employee_id AND lr.work_date=a.work_date
                WHERE a.employee_id IN ({placeholders}) AND p.plan_id<>%s
                  AND p.status='已发布' AND lr.employee_id IS NULL AND a.work_date BETWEEN %s AND %s""",
            [*employee_ids, plan["plan_id"], plan_start - timedelta(days=1), plan_end + timedelta(days=1)],
        )
        rest_assignments.extend({**item, "_candidate": False} for item in published_assignments)
        violation = find_minimum_rest_violation(
            rest_assignments, minimum_rest_hours, require_candidate=True
        )
        if violation:
            previous, current = violation["previous"], violation["current"]
            employee_name = current.get("employee_name") or previous.get("employee_name") or "员工"
            if violation["rest_minutes"] < 0:
                detail = "两个班次发生时间重叠"
            else:
                detail = f"实际仅休息 {violation['rest_minutes'] / 60:g} 小时"
            return None, (
                f"员工 {employee_name} 在 {previous['work_date']} {previous.get('shift_name', '班次')} "
                f"与 {current['work_date']} {current.get('shift_name', '班次')} 之间{detail}，"
                f"少于规则要求的 {minimum_rest_hours:g} 小时"
            )
    return cleaned, None


def collect_schedule_policy_issues(plan, assignments):
    """返回可由人事复核的软规则问题；请假后的班次不计入有效覆盖。"""
    if not assignments:
        return ["班表没有任何班次"]
    rule = query_one("SELECT * FROM scheduling_rules WHERE rule_name='默认排班规则'")
    max_days = int(rule["max_consecutive_days"])
    daily_overtime_limit = float(rule["max_daily_overtime_hours"])
    monthly_overtime_limit = float(rule["max_monthly_overtime_hours"])
    minimum_staff = int(rule["min_staff_per_position"])
    plan_start = parse_iso_date(plan["start_date"])
    plan_end = parse_iso_date(plan["end_date"])
    leave_rows = query_all(
        """SELECT employee_id,work_date FROM schedule_requests
           WHERE request_type='leave' AND status='已批准' AND work_date BETWEEN %s AND %s""",
        (plan_start, plan_end),
    )
    leave_keys = {(int(item["employee_id"]), parse_iso_date(item["work_date"])) for item in leave_rows}
    effective_assignments = [
        item for item in assignments
        if not item.get("is_on_leave")
        and (int(item["employee_id"]), parse_iso_date(item["work_date"])) not in leave_keys
    ]
    if not effective_assignments:
        return ["班表没有任何实际在岗班次"]
    employee_ids = sorted({int(item["employee_id"]) for item in effective_assignments})
    shift_ids = sorted({int(item["shift_id"]) for item in effective_assignments})
    employees = query_all(
        f"SELECT id,name,position FROM employees WHERE id IN ({','.join(['%s'] * len(employee_ids))})",
        employee_ids,
    )
    shifts = query_all(
        f"SELECT shift_id,name,duration_hours FROM shift_types WHERE shift_id IN ({','.join(['%s'] * len(shift_ids))})",
        shift_ids,
    )
    employee_map = {int(item["id"]): item for item in employees}
    shift_map = {int(item["shift_id"]): item for item in shifts}
    employee_dates, monthly_overtime, staffing = {}, {}, {}
    issues = []
    calendar_rows = query_all(
        "SELECT work_date,is_workday,label FROM work_calendar WHERE work_date BETWEEN %s AND %s",
        (plan_start, plan_end),
    )
    calendar_map = {parse_iso_date(item["work_date"]): item for item in calendar_rows}
    for item in effective_assignments:
        employee_id, shift_id = int(item["employee_id"]), int(item["shift_id"])
        work_date = parse_iso_date(item["work_date"])
        employee, shift = employee_map.get(employee_id, {}), shift_map.get(shift_id, {})
        duration = float(shift.get("duration_hours") or 0)
        day_meta = calendar_map.get(work_date)
        if day_meta and not day_meta["is_workday"]:
            day_label = f"（{day_meta.get('label')}）" if day_meta.get("label") else ""
            issues.append(f"员工 {employee.get('name', employee_id)} 在休息日 {work_date} 排班{day_label}，需说明值班、调休或补偿安排")
        overtime = max(duration - 8, 0)
        if overtime > daily_overtime_limit:
            issues.append(f"员工 {employee.get('name', employee_id)} 在 {work_date} 的加班 {overtime:g} 小时，超过每日上限 {daily_overtime_limit:g} 小时")
        month_key = (employee_id, work_date.strftime("%Y-%m"))
        monthly_overtime[month_key] = monthly_overtime.get(month_key, 0) + overtime
        employee_dates.setdefault(employee_id, set()).add(work_date)
        position = str(employee.get("position") or "未设置岗位")
        staffing[(work_date, position)] = staffing.get((work_date, position), 0) + 1
    month_start = plan_start.replace(day=1)
    month_end = add_months(plan_end.replace(day=1), 1) - timedelta(days=1)
    placeholders = ",".join(["%s"] * len(employee_ids))
    other_overtime = query_all(
        f"""SELECT a.employee_id,a.work_date,s.duration_hours FROM schedule_assignments a
             JOIN schedule_plans p ON p.plan_id=a.plan_id JOIN shift_types s ON s.shift_id=a.shift_id
             LEFT JOIN (SELECT employee_id,work_date FROM schedule_requests
                        WHERE request_type='leave' AND status='已批准' GROUP BY employee_id,work_date) lr
               ON lr.employee_id=a.employee_id AND lr.work_date=a.work_date
             WHERE a.employee_id IN ({placeholders}) AND p.plan_id<>%s AND p.status='已发布'
               AND a.work_date BETWEEN %s AND %s AND lr.employee_id IS NULL""",
        [*employee_ids, plan["plan_id"], month_start, month_end],
    )
    for item in other_overtime:
        work_date = parse_iso_date(item["work_date"])
        month_key = (int(item["employee_id"]), work_date.strftime("%Y-%m"))
        monthly_overtime[month_key] = monthly_overtime.get(month_key, 0) + max(float(item["duration_hours"] or 0) - 8, 0)
    for employee_id, work_dates in employee_dates.items():
        continuous = 0
        for work_date in sorted(work_dates):
            continuous = continuous + 1 if work_date - timedelta(days=1) in work_dates else 1
            if continuous == max_days + 1:
                issues.append(f"员工 {employee_map[employee_id]['name']} 连续工作超过 {max_days} 天")
    for (employee_id, month), overtime in monthly_overtime.items():
        if overtime > monthly_overtime_limit:
            issues.append(f"员工 {employee_map[employee_id]['name']} {month} 月累计加班 {overtime:g} 小时，超过上限 {monthly_overtime_limit:g} 小时")
    if minimum_staff > 1:
        for (work_date, position), total in sorted(staffing.items()):
            if total < minimum_staff:
                issues.append(f"{work_date} 岗位“{position}”仅安排 {total} 人，低于最低 {minimum_staff} 人")
    return issues


def record_schedule_snapshot(plan_id, action, note=""):
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if not plan:
        return
    assignments = get_plan_assignments(plan_id, serialize_dates=True)
    payload = {
        "department": plan["department"], "name": plan["name"],
        "start_date": plan["start_date"].isoformat() if hasattr(plan["start_date"], "isoformat") else str(plan["start_date"]),
        "end_date": plan["end_date"].isoformat() if hasattr(plan["end_date"], "isoformat") else str(plan["end_date"]),
        "assignments": [{
            "employee_id": item["employee_id"], "employee_name": item["employee_name"],
            "position": item["position"], "work_date": item["work_date"],
            "shift_id": item["shift_id"], "shift_name": item["shift_name"],
            "duration_hours": item["duration_hours"], "is_on_leave": item.get("is_on_leave", False),
        } for item in assignments],
    }
    execute(
        """INSERT INTO schedule_plan_versions
           (plan_id,plan_version,action,status,actor_id,note,snapshot_json)
           VALUES (%s,%s,%s,%s,%s,%s,%s)""",
        (plan_id, plan["version"], action, plan["status"], current_user()["id"] if current_user() else None,
         str(note or "")[:500], json.dumps(payload, ensure_ascii=False, default=str)),
    )


def build_batch_assignments(plan, employee_ids, shift_id, start_date, end_date, skip_rest_days=True, skip_weekdays=None):
    if not isinstance(employee_ids, list) or not employee_ids:
        return None, "请至少选择一名员工"
    try:
        employee_ids = list({int(employee_id) for employee_id in employee_ids})
        shift_id = int(shift_id)
    except (TypeError, ValueError):
        return None, "员工和班次编号无效"
    start_date, end_date = parse_iso_date(start_date), parse_iso_date(end_date)
    if not start_date or not end_date or start_date > end_date:
        return None, "批量排班日期范围无效"
    if skip_weekdays is None: skip_weekdays = []
    if not isinstance(skip_weekdays, list): return None, "休息星期必须是数组"
    try:
        skip_weekdays = {int(day) for day in skip_weekdays}
    except (TypeError, ValueError): return None, "休息星期格式无效"
    if not skip_weekdays.issubset(set(range(7))): return None, "休息星期仅支持 0（周日）至 6（周六）"
    assignments, current_date = [], start_date
    while current_date <= end_date:
        calendar = query_one("SELECT is_workday FROM work_calendar WHERE work_date=%s", (current_date,))
        if current_date.weekday() not in {(day - 1) % 7 for day in skip_weekdays} and not (skip_rest_days and calendar and not calendar["is_workday"]):
            for employee_id in employee_ids:
                assignments.append({"employee_id": employee_id, "shift_id": shift_id, "work_date": current_date.isoformat()})
        current_date += timedelta(days=1)
    if not assignments: return None, "当前休息规则排除了全部日期，请调整设置"
    return validate_plan_assignments(plan, assignments)


def build_auto_schedule_preview(plan, shift_ids, skip_rest_days=True):
    """按员工轮换班次生成自动排班候选，只用于预览，不写入数据库。"""
    if not isinstance(shift_ids, list) or not shift_ids:
        return None, "请至少选择一个可用班次"
    try:
        shift_ids = list(dict.fromkeys(int(shift_id) for shift_id in shift_ids))
    except (TypeError, ValueError):
        return None, "班次编号无效"
    active_shifts = query_all(
        f"SELECT shift_id FROM shift_types WHERE is_active=1 AND shift_id IN ({','.join(['%s'] * len(shift_ids))})",
        shift_ids,
    )
    if {int(item["shift_id"]) for item in active_shifts} != set(shift_ids):
        return None, "所选班次中包含不存在或已停用的班次"
    employees = query_all("""SELECT id FROM employees WHERE department=%s AND employment_status='在职'
                            AND role<>'department_admin' ORDER BY id""", (plan["department"],))
    if not employees:
        return None, "当前部门没有可排班员工"
    start_date, end_date = parse_iso_date(plan["start_date"]), parse_iso_date(plan["end_date"])
    leave_rows = query_all("""SELECT employee_id,work_date FROM schedule_requests
                            WHERE request_type='leave' AND status='已批准' AND work_date BETWEEN %s AND %s""", (start_date, end_date))
    leave_keys = {(int(item["employee_id"]), parse_iso_date(item["work_date"])) for item in leave_rows}
    assignments, current_date, shift_index = [], start_date, 0
    while current_date <= end_date:
        calendar = query_one("SELECT is_workday FROM work_calendar WHERE work_date=%s", (current_date,))
        if not (skip_rest_days and calendar and not calendar["is_workday"]):
            for employee in employees:
                if (int(employee["id"]), current_date) not in leave_keys:
                    assignments.append({"employee_id": employee["id"], "shift_id": shift_ids[shift_index % len(shift_ids)], "work_date": current_date})
                    shift_index += 1
        current_date += timedelta(days=1)
    return validate_plan_assignments(plan, assignments, allow_existing=True, allow_consecutive_warning=True)


def parse_import_rows(uploaded_file):
    """读取并校验 CSV/XLSX 文件的大小、编码、表头和数据行数。"""
    return schedule_import_reader.read_rows(uploaded_file)


def import_row_value(row, chinese_name, english_name):
    return row.get(chinese_name, row.get(english_name, ""))


# ==================== 员工、入职与档案业务函数 ====================

def authenticate_employee(data):
    """校验登录信息并维护失败次数；返回 (员工, 错误信息, HTTP 状态码)。"""
    if not isinstance(data, dict):
        return None, "请使用 JSON 格式提交登录信息", 400
    employee_no, password = str(data.get("employee_no", "")).strip(), str(data.get("password", ""))
    if not employee_no or not password:
        return None, "工号和密码不能为空", 400
    employee = query_one("SELECT * FROM employees WHERE employee_no = %s", (employee_no,))
    if employee is None:
        return None, "工号或密码错误", 401
    if account_is_locked(employee):
        return None, f"账户已锁定，请于 {employee['locked_until']} 后重试", 423
    if not verify_password(password, employee["password"]):
        record_login_failure(employee["id"])
        return None, "工号或密码错误；连续失败 5 次将锁定账户 30 分钟", 401
    if not employee["password"].startswith(("$2a$", "$2b$", "$2y$", "pbkdf2:", "scrypt:")):
        execute("UPDATE employees SET password=%s WHERE id=%s", (generate_password_hash(password), employee["id"]))
    reset_login_failures(employee["id"])
    return employee, None, None


def create_employee(data):
    """创建员工账号，统一处理注册校验、唯一性校验和密码哈希。"""
    if not isinstance(data, dict):
        return None, "请使用 JSON 格式提交注册信息", 400
    employee_data, error_message = validate_registration_data(data)
    if error_message:
        return None, error_message, 400
    if query_one("SELECT id FROM employees WHERE employee_no = %s", (employee_data["employee_no"],)):
        return None, "该工号已注册，请更换工号", 409
    employee_id, _ = execute(
        "INSERT INTO employees (name, employee_no, password, phone, email) VALUES (%s, %s, %s, %s, %s)",
        (employee_data["name"], employee_data["employee_no"], generate_password_hash(employee_data["password"]), employee_data["phone"], employee_data["email"]),
    )
    return public_employee(find_employee(employee_id)), None, 201


def next_employee_no():
    """按现有 EMP 数字序列生成下一个工号；注册落库时仍会再次校验唯一性。"""
    row = query_one(
        "SELECT employee_no FROM employees WHERE employee_no REGEXP '^EMP[0-9]+$' "
        "ORDER BY CAST(SUBSTRING(employee_no, 4) AS UNSIGNED) DESC LIMIT 1"
    )
    current = 0
    if row:
        match = re.fullmatch(r"EMP(\d+)", str(row.get("employee_no", "")), re.IGNORECASE)
        if match:
            current = int(match.group(1))
    candidate = current + 1
    while query_one("SELECT id FROM employees WHERE employee_no=%s", (f"EMP{candidate:03d}",)):
        candidate += 1
    return f"EMP{candidate:03d}"


def _assistant_registration_intent(question):
    """识别员工开户请求；只在明确提到注册/新增员工或员工账号时命中。"""
    return bool(re.search(
        r"(?:注册|新增|创建|添加|开通).{0,12}(?:员工|账号|用户)|(?:员工|账号|用户).{0,8}(?:注册|新增|创建|添加|开通)",
        question,
        re.IGNORECASE,
    ))


def _assistant_registration_data(question):
    """从自然语言中提取开户字段，最终仍交给 validate_registration_data 做严格校验。"""
    def field(pattern):
        match = re.search(pattern, question, re.IGNORECASE)
        return match.group(1).strip() if match else ""

    name = field(r"(?:员工姓名|姓名|员工名称)\s*(?:是|为|叫|[:：=])?\s*([\u4e00-\u9fffA-Za-z·]{2,20})")
    if not name:
        name = field(r"(?:注册|新增|创建|添加|开通)\s*(?:一个|一名)?\s*员工\s*[,，、:\：]?\s*([\u4e00-\u9fffA-Za-z·]{2,20})")
    employee_no = field(r"\b(EMP\d+)\b").upper()
    password = field(r"(?:初始密码|登录密码|密码)\s*(?:是|为|[:：=])?\s*([^\s，,。；;]+)")
    confirm_password = field(r"(?:确认密码|再次输入密码)\s*(?:是|为|[:：=])?\s*([^\s，,。；;]+)") or password
    phone = field(r"(?:手机号|手机|电话|联系电话)\s*(?:是|为|[:：=])?\s*(1\d{10})")
    if not phone:
        phone = field(r"(?<!\d)(1\d{10})(?!\d)")
    email = field(r"(?:邮箱|电子邮箱|email)\s*(?:是|为|[:：=])?\s*([^\s，,。；;]+)")
    return {
        "name": name,
        "employee_no": employee_no,
        "password": password,
        "confirm_password": confirm_password,
        "phone": phone,
        "email": email,
    }


def _assistant_operation_history_question(question):
    """保存业务操作会话时脱敏密码和手机号，避免敏感信息进入历史记录。"""
    safe = re.sub(
        r"((?:初始密码|登录密码|密码)\s*(?:是|为|[:：=])?\s*)[^\s，,。；;]+",
        r"\1***",
        question,
        flags=re.IGNORECASE,
    )
    safe = re.sub(r"(?<!\d)1\d{2}\d{4}(\d{4})(?!\d)", r"1***\1", safe)
    return safe


def _persist_assistant_operation_response(user, question, conversation_id, response):
    """把已处理的助手操作写入普通会话历史，并把会话 ID返回给前端。"""
    if response.status_code < 200 or response.status_code >= 300 or not response.is_json:
        return response
    payload = response.get_json(silent=True)
    if not isinstance(payload, dict) or not payload.get("handled"):
        return response
    result = {"answer": str(payload.get("answer") or "操作完成。"), "output_format": "normal"}
    try:
        saved_id = save_conversation_exchange(
            user["id"], conversation_id, _assistant_operation_history_question(question), result,
            "buffer", "", None,
        )
    except Exception:
        app.logger.exception("AI 助手业务操作会话保存失败 request_id=%s", getattr(g, "request_id", ""))
        return response
    payload["conversation_id"] = saved_id
    response.set_data(json.dumps(payload, ensure_ascii=False))
    response.content_type = "application/json"
    return response


def _assistant_store_registration(user_id, data):
    """保存短时待确认注册数据；明文密码不会进入客户端 session。"""
    nonce = secrets.token_urlsafe(24)
    now = time.monotonic()
    with _assistant_registration_lock:
        expired = [key for key, item in _assistant_pending_registrations.items() if item["expires_at"] <= now]
        for key in expired:
            _assistant_pending_registrations.pop(key, None)
        _assistant_pending_registrations[nonce] = {
            "user_id": int(user_id), "data": dict(data), "expires_at": now + _ASSISTANT_REGISTRATION_TTL_SECONDS,
        }
    return nonce


def _assistant_take_registration(user_id, nonce):
    if not nonce:
        return None
    with _assistant_registration_lock:
        item = _assistant_pending_registrations.pop(str(nonce), None)
    if not item or item["user_id"] != int(user_id) or item["expires_at"] <= time.monotonic():
        return None
    return item["data"]


def _assistant_profile_update_intent(question):
    return bool(re.search(
        r"(?:修改|更新|更换|变更|设置).{0,12}(?:手机号|手机|电话|邮箱|email)|(?:手机号|手机|电话|邮箱).{0,8}(?:修改|更新|更换|变更|改)",
        question,
        re.IGNORECASE,
    ))


def _assistant_profile_update_data(question):
    phone_match = re.search(r"(?<!\d)(1\d{10})(?!\d)", question)
    email_match = re.search(r"(?:邮箱|电子邮箱|email)\s*(?:改为|更新为|更换为|是|为|[:：=])?\s*([^\s，,。；;]+)", question, re.IGNORECASE)
    data = {}
    if phone_match:
        data["phone"] = phone_match.group(1)
    if email_match:
        data["email"] = email_match.group(1)
    return data


def _assistant_employee_target(question):
    """提取 HR 操作的目标员工；未明确目标时返回 None，避免误操作当前用户。"""
    employee_no = re.search(r"\b(EMP\d+)\b", question, re.IGNORECASE)
    if employee_no:
        return {"employee_no": employee_no.group(1).upper()}
    name_match = re.search(
        r"(?:员工姓名|姓名)\s*(?:是|为|叫|[:：=])?\s*([\u4e00-\u9fffA-Za-z·]{2,12}?)(?=的?(?:手机号|手机|电话|邮箱|部门|岗位|职位|姓名)|\s|[,，。；;]|$)",
        question,
    )
    if not name_match:
        name_match = re.search(
            r"(?:员工|给|为)\s*(?!的?(?:手机号|手机|电话|邮箱|部门|岗位|职位|姓名))([\u4e00-\u9fffA-Za-z·]{2,12}?)(?=的?(?:手机号|手机|电话|邮箱|部门|岗位|职位|姓名)|\s|[,，。；;]|$)",
            question,
        )
    if not name_match:
        name_match = re.search(
            r"(?:修改|更新|更换|变更|设置)\s*(?:把|对)?\s*(?!我的|自己|本人)([\u4e00-\u9fffA-Za-z·]{2,12}?)(?=的?(?:手机号|手机|电话|邮箱|部门|岗位|职位|姓名))",
            question,
        )
    if name_match and name_match.group(1) not in ("自己", "本人", "我的"):
        return {"name": name_match.group(1)}
    return None


def _assistant_find_employee_target(target):
    if target.get("employee_no"):
        rows = query_all("SELECT * FROM employees WHERE employee_no=%s LIMIT 2", (target["employee_no"],))
    else:
        rows = query_all("SELECT * FROM employees WHERE name=%s ORDER BY id LIMIT 2", (target.get("name", ""),))
    if len(rows) != 1:
        return rows
    return rows


def _assistant_hr_profile_changes(question):
    changes = _assistant_profile_update_data(question)
    name_match = re.search(r"(?:姓名|名字)\s*(?:改为|更新为|更换为|是|为|[:：=])\s*([\u4e00-\u9fffA-Za-z·]{2,20})", question)
    department_match = re.search(r"(?:部门)\s*(?:改为|更新为|更换为|调整为|是|为|[:：=])\s*([\u4e00-\u9fffA-Za-z·]{2,30})", question)
    position_match = re.search(r"(?:岗位|职位)\s*(?:改为|更新为|更换为|调整为|是|为|[:：=])\s*([\u4e00-\u9fffA-Za-z·]{2,30})", question)
    role_match = re.search(r"(?:角色)\s*(?:改为|更新为|更换为|调整为|是|为|[:：=])\s*([a-z_]+|普通员工|部门经理|人事|超级管理员)", question, re.IGNORECASE)
    if name_match:
        changes["name"] = name_match.group(1)
    if department_match:
        changes["department"] = department_match.group(1)
    if position_match:
        changes["position"] = position_match.group(1)
    if role_match:
        role_alias = {"普通员工": "employee", "部门经理": "department_admin", "人事": "hr_admin", "超级管理员": "super_admin"}
        changes["role"] = role_alias.get(role_match.group(1), role_match.group(1))
    return changes


def _assistant_update_own_profile(user, changes):
    if not changes:
        return None, "请提供新的手机号或邮箱"
    if "phone" in changes:
        if not PHONE_PATTERN.fullmatch(changes["phone"]):
            return None, "手机号格式无效，请输入 11 位手机号"
    if "email" in changes and changes["email"] and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", changes["email"]):
        return None, "邮箱格式无效"
    assignments = []
    parameters = []
    for field in ("phone", "email"):
        if field in changes:
            assignments.append(f"{field}=%s")
            parameters.append(changes[field])
    parameters.append(user["id"])
    before = {field: user.get(field) for field in changes}
    _, changed = execute(f"UPDATE employees SET {', '.join(assignments)} WHERE id=%s", parameters)
    after = find_employee(user["id"])
    g.audit_details = {"tool": "employee_update", "target_employee_id": user["id"], "before": before, "after": {field: after.get(field) for field in changes}}
    return changed, None


def _assistant_apply_hr_profile_update(user, target_id, changes):
    if user["role"] not in ("hr_admin", "super_admin"):
        return None, "当前登录账号无权修改其他员工档案", 403
    target = find_employee(target_id)
    if target is None:
        return None, "目标员工不存在", 404
    if changes.get("role") in ROLE_NAMES and changes["role"] != target["role"] and changes["role"] != "employee" and user["role"] != "super_admin":
        return None, "只有超级管理员可以授予或调整管理角色", 403
    if changes.get("role") and changes["role"] not in ROLE_NAMES:
        return None, "角色类型无效", 400
    before = {field: target.get(field) for field in changes}
    assignments, parameters = [], []
    for field in ("name", "phone", "email", "department", "position", "role"):
        if field in changes:
            value = str(changes[field]).strip()
            if field == "phone" and not PHONE_PATTERN.fullmatch(value):
                return None, "手机号格式无效，请输入 11 位手机号", 400
            if field == "email" and value and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", value):
                return None, "邮箱格式无效", 400
            if not value:
                return None, f"{field} 不能为空", 400
            assignments.append(f"{field}=%s")
            parameters.append(value)
    if not assignments:
        return None, "请提供要修改的档案字段", 400
    parameters.append(target_id)
    _, changed = execute(f"UPDATE employees SET {', '.join(assignments)} WHERE id=%s", parameters)
    updated = find_employee(target_id)
    g.audit_details = {"tool": "employee_update", "target_employee_id": target_id, "before": before, "after": {field: updated.get(field) for field in changes}}
    return updated, None, 200 if changed else 200


def _assistant_hr_start_offboarding(user, target_id, reason):
    if user["role"] not in ("hr_admin", "super_admin"):
        return None, "当前登录账号无权发起其他员工离职", 403
    target = find_employee(target_id)
    if target is None:
        return None, "目标员工不存在", 404
    if target["role"] == "super_admin":
        return None, "不能通过人事离职流程处理超级管理员账号", 403
    if target["employment_status"] != "在职":
        return None, "目标员工当前不处于在职状态", 409
    before = {"employment_status": target.get("employment_status")}
    updated, error_message, status_code = update_employee_data(
        target_id, {"action": "offboarding", "task_content": reason[:500]},
    )
    if not error_message:
        g.audit_details = {"tool": "employee_offboarding", "target_employee_id": target_id, "before": before, "after": {"employment_status": updated.get("employment_status")}, "reason": reason[:500]}
    return updated, error_message, status_code


def _assistant_create_offboarding_request(user, reason):
    if user["role"] != "employee":
        return None, "当前角色不能发起个人离职申请", 403
    if user["employment_status"] != "在职":
        return None, "当前账户不处于可申请离职状态", 409
    if query_one("SELECT task_id FROM employee_tasks WHERE employee_id=%s AND task_type='offboarding_request' AND status IN ('待部门审批','待老板批准')", (user["id"],)):
        return None, "已有进行中的离职申请", 409
    before = {"employment_status": user.get("employment_status")}
    task_id = create_employee_task(user, "offboarding_request", "离职申请", f"离职原因：{reason}")
    execute("UPDATE employee_tasks SET status='待部门审批' WHERE task_id=%s", (task_id,))
    execute("UPDATE employees SET employment_status='离职申请中' WHERE id=%s", (user["id"],))
    updated = find_employee(user["id"])
    g.audit_details = {"tool": "employee_offboarding", "target_employee_id": user["id"], "before": before, "after": {"employment_status": updated.get("employment_status")}, "reason": reason[:500]}
    return task_id, None, 201


def _assistant_handover_task(user):
    return query_one(
        "SELECT * FROM employee_tasks WHERE employee_id=%s AND task_type='offboarding_handover' AND status='待处理' ORDER BY task_id DESC LIMIT 1",
        (user["id"],),
    )


def _assistant_submit_handover(user, task_id, content):
    before = query_one("SELECT status FROM employee_tasks WHERE task_id=%s AND employee_id=%s", (task_id, user["id"])) or {}
    _, changed = execute(
        "UPDATE employee_tasks SET status='待人事审核', content=CONCAT(content, '\\n员工交接说明：', %s), submitted_at=NOW() "
        "WHERE task_id=%s AND employee_id=%s AND task_type='offboarding_handover' AND status='待处理'",
        (content, task_id, user["id"]),
    )
    after = query_one("SELECT status FROM employee_tasks WHERE task_id=%s AND employee_id=%s", (task_id, user["id"])) or {}
    g.audit_details = {"tool": "employee_handover", "target_employee_id": user["id"], "task_id": task_id, "before": before, "after": after}
    return changed


def _assistant_create_handover_request(user, content):
    if user["role"] != "employee":
        return None, "当前角色不能发起个人交接", 403
    if user["employment_status"] == "已离职":
        return None, "已离职员工不能发起交接", 409
    if query_one(
        "SELECT task_id FROM employee_tasks WHERE employee_id=%s AND task_type IN ('handover_request','offboarding_handover') AND status IN ('待处理','待人事审核')",
        (user["id"],),
    ):
        return None, "已有进行中的交接任务，请先处理现有任务", 409
    task_id = create_employee_task(user, "handover_request", "工作交接申请", content[:2000])
    execute("UPDATE employee_tasks SET status='待人事审核', submitted_at=NOW() WHERE task_id=%s", (task_id,))
    notify_users_safely(
        notification_recipient_ids("hr_admin", "super_admin", exclude_id=user["id"]),
        "handover_request", "新的工作交接申请", f"{user['name']}提交了工作交接申请，请及时审核。",
        "/pages/queue.html", related_type="employee_task", related_id=task_id,
        dedupe_key=f"assistant-handover-request:{task_id}",
    )
    g.audit_details = {"tool": "employee_handover", "target_employee_id": user["id"], "task_id": task_id, "before": {"task_status": None}, "after": {"task_status": "待人事审核"}}
    return task_id, None, 201


def update_employee_data(employee_id, data):
    """更新现有员工资料，并按既有规则创建员工待办。"""
    employee = find_employee(employee_id)
    if employee is None:
        return None, "员工不存在", 404
    if not isinstance(data, dict):
        return None, "请使用 JSON 格式提交员工信息", 400
    department = str(data.get("department", employee["department"])).strip()
    position = str(data.get("position", employee["position"])).strip()
    role = str(data.get("role", employee["role"])).strip()
    if role not in ROLE_NAMES:
        return None, "角色类型无效", 400
    if role == "super_admin":
        return None, "超级管理员账户不由人事页面创建或调整", 403
    execute("UPDATE employees SET department=%s, position=%s, role=%s WHERE id=%s", (department, position, role, employee_id))
    action = data.get("action", "profile_update")
    if action in ("profile_update", "handover", "offboarding"):
        titles = {"profile_update": "员工档案信息已更新，请确认", "handover": "请完成工作交接", "offboarding": "请完成离职交接"}
        content = str(data.get("task_content", "请查看最新档案信息并提交处理结果。"))
        task_id = create_employee_task(find_employee(employee_id), action, titles[action], content)
        if action in ("handover", "offboarding"):
            execute("UPDATE employees SET employment_status=%s WHERE id=%s", ("待超管批准" if action == "offboarding" else "交接中", employee_id))
        if action == "offboarding":
            execute("UPDATE employee_tasks SET status='已提交',submitted_at=NOW() WHERE task_id=%s", (task_id,))
            execute("INSERT INTO admin_todos (task_id,employee_id,employee_name,title,content) VALUES (%s,%s,%s,%s,%s)", (task_id, employee_id, employee["name"], "离职申请审批", content))
    return public_employee(find_employee(employee_id)), None, 200


def create_onboarding(data):
    """为员工创建入职待办并同步部门、岗位、角色状态。"""
    if not isinstance(data, dict):
        return None, "请使用 JSON 格式提交入职信息", 400
    try:
        employee_id = int(data.get("employee_id"))
    except (TypeError, ValueError):
        return None, "employee_id 必须是有效编号", 400
    employee = find_employee(employee_id)
    if employee is None:
        return None, "员工不存在", 404
    department, position, role = str(data.get("department", "")).strip(), str(data.get("position", "")).strip(), str(data.get("role", "employee")).strip()
    if not department or not position:
        return None, "部门和岗位不能为空", 400
    if role not in ROLE_NAMES:
        return None, "角色类型无效", 400
    onboarding_type = OnboardingPolicy.determine_onboarding_type(department, position)
    if query_one("SELECT record_id FROM onboarding_records WHERE employee_id = %s AND status IN ('待超管批准', '待办理')", (employee_id,)):
        return None, "该员工已有进行中的入职申请", 400
    queue_number = query_one("SELECT COALESCE(MAX(queue_number), 0) + 1 AS next_number FROM onboarding_records")["next_number"]
    record_id, _ = execute("""INSERT INTO onboarding_records (employee_id, employee_name, department, position, onboarding_type, priority, queue_number, status)
        VALUES (%s, %s, %s, %s, %s, %s, %s, '待超管批准')""", (employee_id, employee["name"], department, position, onboarding_type, OnboardingPolicy.determine_priority(position), queue_number))
    execute("UPDATE employees SET department=%s, position=%s, role=%s, employment_status='待超管批准' WHERE id=%s", (department, position, role, employee_id))
    task_id = create_employee_task(find_employee(employee_id), "onboarding", "入职申请审批", f"人事已提交 {department} {position} 的入职申请。")
    execute("UPDATE employee_tasks SET status='已提交',submitted_at=NOW() WHERE task_id=%s", (task_id,))
    execute("INSERT INTO admin_todos (task_id,employee_id,employee_name,title,content) VALUES (%s,%s,%s,%s,%s)", (task_id, employee_id, employee["name"], "入职申请审批", f"人事申请为其办理 {department} {position} 入职。"))
    return query_one("SELECT * FROM onboarding_records WHERE record_id=%s", (record_id,)), None, 201


def create_employee_archive(data):
    """创建入职档案，保留原有的入职办结和唯一档案限制。"""
    if not isinstance(data, dict):
        return None, "请使用 JSON 格式提交档案信息", 400
    try:
        employee_id = int(data.get("employee_id"))
    except (TypeError, ValueError):
        return None, "employee_id 必须是有效编号", 400
    employee = find_employee(employee_id)
    if employee is None:
        return None, "员工不存在", 404
    if not query_one("SELECT record_id FROM onboarding_records WHERE employee_id=%s AND status='已办理'", (employee_id,)):
        return None, "请先完成该员工的入职待办，再创建档案", 400
    if query_one("SELECT record_id FROM employee_records WHERE employee_id=%s", (employee_id,)):
        return None, "该员工已存在入职档案", 409
    training_items, error_message = normalize_list(data.get("training_items", []), "培训项目")
    if error_message:
        return None, error_message, 400
    assets, error_message = normalize_list(data.get("assets", []), "领用资产")
    if error_message:
        return None, error_message, 400
    record_id, _ = execute("""INSERT INTO employee_records (employee_id, employee_name, department, position, role, training_items, assets, remark)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)""", (employee_id, employee["name"], employee["department"], employee["position"], employee["role"], json.dumps(training_items, ensure_ascii=False), json.dumps(assets, ensure_ascii=False), str(data.get("remark", "")).strip()))
    return serialize_record(query_one("SELECT * FROM employee_records WHERE record_id=%s", (record_id,))), None, 201


def update_employee_archive(record_id, data):
    """更新档案可编辑字段，保留原有部分更新语义。"""
    if not query_one("SELECT record_id FROM employee_records WHERE record_id=%s", (record_id,)):
        return None, "档案不存在", 404
    if not isinstance(data, dict):
        return None, "请使用 JSON 格式提交更新信息", 400
    updates, parameters = [], []
    for key, field_name in (("training_items", "培训项目"), ("assets", "领用资产")):
        if key in data:
            items, error_message = normalize_list(data[key], field_name)
            if error_message:
                return None, error_message, 400
            updates.append(f"{key}=%s")
            parameters.append(json.dumps(items, ensure_ascii=False))
    if "remark" in data:
        updates.append("remark=%s")
        parameters.append(str(data["remark"]).strip())
    if updates:
        parameters.append(record_id)
        execute(f"UPDATE employee_records SET {', '.join(updates)} WHERE record_id=%s", parameters)
    return serialize_record(query_one("SELECT * FROM employee_records WHERE record_id=%s", (record_id,))), None, 200


def prepare_schedule_import(plan_id, uploaded_file):
    """只解析并校验导入文件，不写数据库；供预览和正式导入共同使用。"""
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None:
        return None, "排班计划不存在", 404
    if plan["status"] not in ("草稿", "已打回"):
        return None, "仅草稿或已打回计划可以导入排班", 409
    if uploaded_file is None or not uploaded_file.filename:
        return None, "请选择要导入的 CSV 或 Excel 文件", 400
    try:
        rows, error_message = parse_import_rows(uploaded_file)
    except Exception:
        return None, "无法读取文件，请确认 CSV 使用 UTF-8 编码或 Excel 文件未损坏", 400
    if error_message:
        return None, error_message, 400
    assignments = []
    for index, row in enumerate(rows, start=2):
        employee_no = str(import_row_value(row, "员工工号", "employee_no") or "").strip()
        work_date_value = import_row_value(row, "排班日期", "work_date")
        work_date = work_date_value.isoformat() if hasattr(work_date_value, "isoformat") else str(work_date_value or "").strip()
        shift_name = str(import_row_value(row, "班次名称", "shift_name") or "").strip()
        if not employee_no or not work_date or not shift_name:
            return None, f"第 {index} 行必须包含员工工号、排班日期、班次名称", 400
        employee = query_one("SELECT id, name, employee_no FROM employees WHERE employee_no=%s", (employee_no,))
        shift = query_one("SELECT shift_id, name FROM shift_types WHERE name=%s AND is_active=1", (shift_name,))
        if employee is None:
            return None, f"第 {index} 行员工工号不存在：{employee_no}", 400
        if shift is None:
            return None, f"第 {index} 行班次不存在或已停用：{shift_name}", 400
        assignments.append({
            "employee_id": employee["id"], "employee_name": employee["name"], "employee_no": employee["employee_no"],
            "shift_id": shift["shift_id"], "shift_name": shift["name"], "work_date": work_date,
        })
    assignments, error_message = validate_plan_assignments(plan, assignments)
    if error_message:
        return None, f"导入校验失败：{error_message}", 400
    # 校验函数会返回最小化的数据结构；预览需要补回用户可读的工号和名称。
    for item in assignments:
        employee = query_one("SELECT employee_no, name FROM employees WHERE id=%s", (item["employee_id"],))
        shift = query_one("SELECT name FROM shift_types WHERE shift_id=%s", (item["shift_id"],))
        item.update({"employee_no": employee["employee_no"], "employee_name": employee["name"], "shift_name": shift["name"]})
    return assignments, f"已识别 {len(assignments)} 条排班", 200


def import_schedule_file(plan_id, uploaded_file):
    """把导入文件转换为排班记录并执行现有的统一排班校验后写入。"""
    assignments, message, status_code = prepare_schedule_import(plan_id, uploaded_file)
    if status_code != 200:
        return None, message, status_code
    saved, save_error = append_schedule_assignments_atomic(plan_id, assignments)
    if not saved:
        return None, save_error, 409
    record_schedule_snapshot(plan_id, "导入排班")
    return get_plan_assignments(plan_id, serialize_dates=True), f"成功导入 {len(assignments)} 条排班", 201


def create_employee_task(employee, task_type, title, content):
    task_id, _ = execute("INSERT INTO employee_tasks (employee_id, task_type, title, content) VALUES (%s, %s, %s, %s)", (employee["id"], task_type, title, content))
    return task_id


def manual_task_targets(user, data):
    """Resolve and authorize manual-task recipients before any task is created."""
    target_type = str(data.get("target_type", "")).strip()
    if target_type == "employee":
        raw_ids = data.get("employee_ids")
        if not isinstance(raw_ids, list) or not raw_ids or len(raw_ids) > 2000:
            return None, "请选择 1 至 2000 名员工"
        try:
            employee_ids = sorted({int(item) for item in raw_ids})
        except (TypeError, ValueError):
            return None, "员工编号无效"
        placeholders = ",".join(["%s"] * len(employee_ids))
        employees = query_all(
            f"SELECT id,name,employee_no,department FROM employees WHERE id IN ({placeholders}) AND employment_status='在职' ORDER BY id",
            employee_ids,
        )
    elif target_type == "department":
        department = str(data.get("department", "")).strip()
        if not department:
            return None, "请选择目标部门"
        employees = query_all(
            "SELECT id,name,employee_no,department FROM employees WHERE department=%s AND employment_status='在职' AND role<>'super_admin' ORDER BY id",
            (department,),
        )
    else:
        return None, "接收范围必须是员工或部门"
    if not employees:
        return None, "没有符合条件的在职员工"
    if user["role"] == "department_admin" and any(item["department"] != user["department"] for item in employees):
        return None, "部门经理只能向本部门员工发布待办"
    return employees, None


def publish_manual_tasks(user, data):
    """发布人工待办的共用逻辑：网页会话与智能体入口共用，按调用方角色校验范围。"""
    if user["role"] not in ("department_admin", "hr_admin", "super_admin"): return jsonify({"message": "无权限发布人工待办"}), 403
    title = str(data.get("title", "")).strip()
    content = str(data.get("content", "")).strip()
    priority = str(data.get("priority", "普通")).strip()
    due_date = str(data.get("due_date", "")).strip() or None
    if not title or len(title) > 100: return jsonify({"message": "标题不能为空且不能超过 100 个字"}), 400
    if not content or len(content) > 1000: return jsonify({"message": "内容不能为空且不能超过 1000 个字"}), 400
    if priority not in ("普通", "重要", "紧急"): return jsonify({"message": "优先级无效"}), 400
    if due_date and not parse_iso_date(due_date): return jsonify({"message": "截止日期格式无效"}), 400
    if due_date and parse_iso_date(due_date) < date.today(): return jsonify({"message": "截止日期不能早于今天"}), 400
    employees, error_message = manual_task_targets(user, data)
    if error_message: return jsonify({"message": error_message}), 400
    client_request_id = str(data.get("idempotency_key") or request.headers.get("X-Request-ID") or g.request_id).strip()[:64]
    if not REQUEST_ID_PATTERN.fullmatch(client_request_id): return jsonify({"message": "请求编号无效，请刷新后重试"}), 400
    task_ids = []
    try:
        with connect_database() as connection:
            connection.begin()
            try:
                with connection.cursor() as cursor:
                    for employee in employees:
                        cursor.execute(
                            """INSERT INTO employee_tasks
                               (employee_id,task_type,title,content,status,priority,due_date,publisher_id,client_request_id,published_at)
                               VALUES (%s,'manual',%s,%s,'待处理',%s,%s,%s,%s,NOW())""",
                            (employee["id"], title, content, priority, due_date, user["id"], client_request_id),
                        )
                        task_ids.append(cursor.lastrowid)
                connection.commit()
            except Exception:
                connection.rollback()
                raise
    except pymysql.IntegrityError:
        existing = query_all("SELECT task_id FROM employee_tasks WHERE publisher_id=%s AND client_request_id=%s", (user["id"], client_request_id))
        if existing:
            return jsonify({"success": True, "message": "该发布请求已处理，未重复创建待办", "count": len(existing), "duplicate": True}), 200
        raise
    notify_users_safely(
        [employee["id"] for employee in employees], "task_published", "收到新的待办",
        f"{title} · {priority}{f' · 截止 {due_date}' if due_date else ''}",
        "/pages/my-todos.html", related_type="employee_task", related_id=task_ids[0],
        dedupe_key=f"manual-task:{task_ids[0]}",
    )
    return jsonify({"success": True, "message": f"已向 {len(task_ids)} 名员工发布待办", "count": len(task_ids)}), 201


@api_bp.route("/api/tasks/manual", methods=["POST"])
def create_manual_tasks():
    """Create one employee task per recipient; one notification points back to the task center."""
    user = current_user()
    if user is None: return jsonify({"message": "请先登录"}), 401
    return publish_manual_tasks(user, request.get_json(silent=True) or {})


@api_bp.route("/api/tasks/manual/targets", methods=["GET"])
def get_manual_task_targets():
    """Return only recipients visible to the current task publisher."""
    user = current_user()
    if user is None: return jsonify({"message": "请先登录"}), 401
    if user["role"] not in ("department_admin", "hr_admin", "super_admin"):
        return jsonify({"message": "无权限读取待办接收人"}), 403
    if user["role"] == "department_admin":
        departments = [{"name": user["department"]}] if user.get("department") else []
        employees = query_all(
            "SELECT id,name,employee_no,department FROM employees WHERE department=%s AND employment_status='在职' ORDER BY id",
            (user["department"],),
        )
    else:
        departments = query_all("SELECT name FROM departments ORDER BY department_id")
        employees = query_all(
            "SELECT id,name,employee_no,department FROM employees WHERE employment_status='在职' AND role<>'super_admin' ORDER BY department,id"
        )
    return jsonify({"success": True, "departments": departments, "employees": employees})


@api_bp.route("/api/tasks/manual/published", methods=["GET"])
def list_published_manual_tasks():
    user = current_user()
    if user is None: return jsonify({"message": "请先登录"}), 401
    if user["role"] not in ("department_admin", "hr_admin", "super_admin"): return jsonify({"message": "无权限查看已发布待办"}), 403
    conditions = ["t.task_type='manual'", "t.publisher_id=%s"]
    parameters = [user["id"]]
    if user["role"] == "department_admin":
        conditions.append("e.department=%s")
        parameters.append(user["department"])
    tasks = query_all(
        """SELECT t.task_id,t.title,t.content,t.status,
                  CASE WHEN t.status='待处理' AND t.due_date IS NOT NULL AND t.due_date<CURDATE() THEN '已逾期' ELSE t.status END AS display_status,
                  t.priority,t.due_date,t.created_at,t.submitted_at,t.completed_at,
                  e.name AS employee_name,e.employee_no,e.department
           FROM employee_tasks t JOIN employees e ON e.id=t.employee_id
           WHERE """ + " AND ".join(conditions) + " ORDER BY t.task_id DESC LIMIT 500",
        parameters,
    )
    return jsonify({"success": True, "tasks": tasks})


def notification_recipient_ids(*roles, department=None, exclude_id=None):
    """按岗位和部门解析消息接收人，自动排除已离职账户。"""
    conditions, parameters = ["employment_status<>'已离职'"], []
    if roles:
        conditions.append(f"role IN ({','.join(['%s'] * len(roles))})")
        parameters.extend(roles)
    if department is not None:
        conditions.append("department=%s")
        parameters.append(department)
    if exclude_id is not None:
        conditions.append("id<>%s")
        parameters.append(exclude_id)
    return [item["id"] for item in query_all(
        "SELECT id FROM employees WHERE " + " AND ".join(conditions), parameters,
    )]


def notify_users_safely(recipient_ids, notification_type, title, content, target_url, **metadata):
    """消息失败不回滚已成功的主营业务，但必须写入服务端错误日志。"""
    try:
        return create_notifications(
            recipient_ids, notification_type, title, content, target_url, **metadata,
        )
    except pymysql.MySQLError:
        app.logger.exception("站内消息写入失败 type=%s title=%s", notification_type, title)
        return 0


def current_user():
    """根据会话获取当前登录用户，返回时不包含密码。"""
    employee_id = session.get("employee_id")
    if not employee_id:
        return None
    employee = find_employee(employee_id)
    return public_employee(employee) if employee else None


def verify_password(password, stored_password):
    """校验 bcrypt / Werkzeug 哈希；历史明文登录成功后会立即升级。"""
    if stored_password.startswith(("$2a$", "$2b$", "$2y$")):
        return bcrypt.checkpw(password.encode("utf-8"), stored_password.encode("utf-8"))
    if stored_password.startswith(("pbkdf2:", "scrypt:")):
        return check_password_hash(stored_password, password)
    return hmac.compare_digest(password, stored_password)


def roles_required(*allowed_roles, message=""):
    """带参数的权限装饰器。

    ``allowed_roles`` 和 ``message`` 被内部包装函数闭包捕获，新增岗位权限时只需
    声明角色集合，不必复制整段会话、权限和响应逻辑。
    """
    def decorator(view):
        @wraps(view)
        def wrapped_view(*args, **kwargs):
            user = current_user()
            if user is None:
                return jsonify({"message": "请先登录"}), 401
            if allowed_roles and user["role"] not in allowed_roles:
                return jsonify({"message": message}), 403
            return view(*args, **kwargs)
        return wrapped_view
    return decorator


login_required = roles_required()


super_admin_required = roles_required("super_admin", message="无权限：仅超级管理员可以执行此操作")


department_admin_required = roles_required("department_admin", "super_admin", message="无权限：仅部门管理员可以执行此操作")


hr_required = roles_required("hr_admin", "super_admin", message="无权限：仅人事部门可以执行此操作")


schedule_owner_required = roles_required("department_admin", "hr_admin", "super_admin", message="无权限：仅部门经理、人事部门或老板可以管理班表")


def office_session_token():
    """从智能办公请求中取出短期身份凭证。"""
    office_token = request.headers.get("X-Office-Session", "").strip()
    authorization = request.headers.get("Authorization", "")
    if not office_token and authorization.lower().startswith("bearer "):
        office_token = authorization[7:].strip()
    return office_token


def office_session_user(office_token):
    """验证智能办公短期身份凭证并返回当前在职用户。"""
    if not office_token:
        return None
    serializer = URLSafeTimedSerializer(SECRET_KEY, salt="smart-office-session")
    try:
        payload = serializer.loads(office_token, max_age=OFFICE_TOKEN_MAX_AGE_SECONDS)
    except (BadSignature, SignatureExpired, TypeError):
        return None
    if not isinstance(payload, dict) or payload.get("aud") != "smart-office" or not payload.get("employee_id"):
        return None
    try:
        employee_id = int(payload["employee_id"])
    except (TypeError, ValueError, OverflowError):
        return None
    employee = query_one(
        "SELECT id FROM employees WHERE id=%s AND employment_status='在职'",
        (employee_id,),
    )
    return public_employee(find_employee(employee["id"])) if employee else None


def agent_current_user():
    """解析由登录会话签发的短期身份凭证，不信任聊天内容中的工号。"""
    supplied_key = request.headers.get("X-API-Key", "")
    office_token = office_session_token()
    if not AGENT_API_KEY or not supplied_key or not office_token:
        return None
    if not hmac.compare_digest(supplied_key, AGENT_API_KEY):
        return None
    return office_session_user(office_token)


def agent_auth_required(view):
    """保护外部智能体接口，并要求调用方提供登录会话签发的身份凭证。"""
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        user = agent_current_user()
        if user is None:
            return jsonify({"message": "API Key 无效、缺失或智能办公身份凭证无效或已过期"}), 401
        return view(user, *args, **kwargs)
    return wrapped_view


def validate_json_fields(*field_names):
    """校验 JSON 请求中的必填字段，避免每个路由重复编写相同判断。"""
    def decorator(view):
        @wraps(view)
        def wrapped_view(*args, **kwargs):
            if request.method == "OPTIONS":
                return view(*args, **kwargs)
            data = request.get_json(silent=True)
            if not isinstance(data, dict):
                return jsonify({"message": "请使用 JSON 格式提交数据"}), 400
            missing = [field for field in field_names if not str(data.get(field, "")).strip()]
            if missing:
                return jsonify({"message": f"缺少必填参数：{'、'.join(missing)}"}), 400
            return view(*args, **kwargs)
        return wrapped_view
    return decorator


def api_operation_log(operation):
    """记录接口操作摘要，不记录密码、原因等请求正文。"""
    def decorator(view):
        @wraps(view)
        def wrapped_view(*args, **kwargs):
            user = current_user()
            # 登录前会话可能属于上一次用户；登录日志统一视为匿名请求，避免归属误导。
            actor = "anonymous" if operation == "用户登录" else (user["employee_no"] if user else "anonymous")
            try:
                result = view(*args, **kwargs)
            except Exception:
                app.logger.exception("API操作异常 request_id=%s operation=%s actor=%s method=%s path=%s", g.request_id, operation, actor, request.method, request.path)
                raise
            status_code = result[1] if isinstance(result, tuple) and len(result) > 1 and isinstance(result[1], int) else getattr(result, "status_code", 200)
            app.logger.info("API操作 request_id=%s operation=%s actor=%s method=%s path=%s status=%s", g.request_id, operation, actor, request.method, request.path, status_code)
            return result
        return wrapped_view
    return decorator


def may_manage_department(user, department):
    return user["role"] in ("hr_admin", "super_admin") or (user["role"] == "department_admin" and user["department"] == department)


@app.before_request
def remember_audit_actor():
    """在注销前保留操作者身份，供响应阶段落库审计。"""
    supplied_request_id = request.headers.get("X-Request-ID", "").strip()
    g.request_id = supplied_request_id if REQUEST_ID_PATTERN.fullmatch(supplied_request_id) else secrets.token_hex(16)
    g.audit_actor = current_user()
    if app.config.get("TESTING") or request.method in {"GET", "HEAD", "OPTIONS"}:
        return None
    if request.path == "/api/auth/login" or not request.path.startswith("/api/"):
        return None
    if not session.get("employee_id"):
        return None
    token = request.headers.get("X-CSRF-Token", "")
    if not token or not hmac.compare_digest(token, session.get("csrf_token", "")):
        return jsonify({"message": "请求校验已失效，请刷新页面后重试"}), 403
    return None


@app.after_request
def secure_response(response):
    """同源部署的基础安全响应头与关键写操作审计。"""
    response.headers["X-Request-ID"] = getattr(g, "request_id", secrets.token_hex(16))
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "geolocation=(), microphone=(), camera=()"
    # The embedded office workspace is a separate local service on port 8080.
    # Keep the restrictive defaults while explicitly allowing only that iframe.
    response.headers["Content-Security-Policy"] = "default-src 'self'; frame-src 'self' http://127.0.0.1:8080 http://localhost:8080; base-uri 'self'; frame-ancestors 'none'; form-action 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; object-src 'none'"
    if IS_PRODUCTION:
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    if request.path.startswith("/api/") and request.method not in {"GET", "HEAD", "OPTIONS"}:
        actor = current_user() or getattr(g, "audit_actor", None)
        try:
            audit_details = getattr(g, "audit_details", {})
            if not isinstance(audit_details, dict):
                audit_details = {}
            execute(
                """INSERT INTO audit_logs (request_id, actor_id, actor_employee_no, action, http_method, request_path, status_code, client_ip, user_agent, details)
                   VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                (g.request_id, actor["id"] if actor else None, actor["employee_no"] if actor else "anonymous", request.endpoint or "unknown",
                 request.method, request.path[:255], response.status_code, request.remote_addr or "", request.user_agent.string[:255],
                 json.dumps(audit_details, ensure_ascii=False, default=str)[:4000]),
            )
        except pymysql.MySQLError:
            app.logger.exception("审计日志写入失败 request_id=%s path=%s", g.request_id, request.path)
    return response


@app.errorhandler(pymysql.MySQLError)
def handle_database_error(error):
    app.logger.error("数据库操作失败：%s", error)
    return jsonify({"message": "数据库操作失败，请检查 MySQL 服务和连接配置"}), 500


@app.errorhandler(413)
def handle_payload_too_large(_error):
    return jsonify({"message": f"请求体不能超过 {MAX_CONTENT_LENGTH // 1024 // 1024} MB"}), 413


@api_bp.route("/api/auth/login", methods=["POST", "OPTIONS"])
@api_operation_log("用户登录")
@validate_json_fields("employee_no", "password")
def login():
    if request.method == "OPTIONS": return "", 204
    # 登录失败也不能沿用上一个账号的会话，避免账户切换时出现权限残留。
    session.clear()
    employee, error_message, status_code = authenticate_employee(request.get_json(silent=True))
    if error_message: return jsonify({"message": error_message}), status_code
    session["employee_id"] = employee["id"]
    session["csrf_token"] = secrets.token_urlsafe(32)
    session.permanent = True
    return jsonify({"success": True, "message": "登录成功", "user": public_employee(employee), "csrf_token": session["csrf_token"]})


@api_bp.route("/api/health", methods=["GET"])
@api_bp.route("/api/health/ready", methods=["GET"])
def health_check():
    """就绪探针：应用与数据库都可用时才返回成功。"""
    query_one("SELECT 1 AS alive")
    return jsonify({"status": "ok"})


@api_bp.route("/api/health/live", methods=["GET"])
def liveness_check():
    """存活探针不访问外部依赖，避免数据库故障触发无效重启。"""
    return jsonify({"status": "ok"})


def _assistant_leave_query(user, scope, include_history=False):
    """按当前登录用户的角色构造请假查询范围，前端无权指定员工或部门。"""
    conditions, parameters = ["r.request_type='leave'"], []
    if scope == "mine":
        conditions.append("r.employee_id=%s")
        parameters.append(user["id"])
    elif scope == "department" and user["role"] == "department_admin":
        conditions.append("r.department=%s")
        parameters.append(user["department"])
    elif scope == "all" and user["role"] in ("hr_admin", "super_admin"):
        pass
    elif scope == "pending" and user["role"] == "department_admin":
        conditions.extend(["r.department=%s", "r.employee_id<>%s", "r.status='待部门审批'"])
        parameters.extend([user["department"], user["id"]])
    elif scope == "pending" and user["role"] in ("hr_admin", "super_admin"):
        conditions.append("r.status='待人事审批'")
    else:
        return None, None
    if not include_history:
        # 日常统计关注进行中或已批准的请假；历史驳回仅在用户明确要求时返回。
        conditions.append("r.status IN ('待部门审批','待人事审批','已批准')")
    where_sql = " AND ".join(conditions)
    summary = query_one(
        "SELECT COUNT(*) AS total,COUNT(DISTINCT r.employee_id) AS employee_count "
        "FROM schedule_requests r WHERE " + where_sql, parameters,
    )
    records = query_all(
        """SELECT r.request_id,e.name AS employee_name,r.department,r.work_date,r.reason,r.status
           FROM schedule_requests r JOIN employees e ON e.id=r.employee_id
           WHERE """ + where_sql + " ORDER BY r.request_id DESC LIMIT 20", parameters,
    )
    return summary, records


def _assistant_leave_answer(scope, summary, records):
    scope_names = {"mine": "我的", "department": "本部门", "all": "全员", "pending": "待审批"}
    lines = [f"{scope_names[scope]}请假申请共 {int(summary['total'])} 条，涉及 {int(summary['employee_count'])} 名员工。"]
    for record in records:
        work_date = record["work_date"].isoformat() if hasattr(record["work_date"], "isoformat") else str(record["work_date"])
        lines.append(f"#{record['request_id']} {record['employee_name']} · {work_date} · {record['status']} · {record['reason']}")
    if not records:
        lines.append("当前没有符合条件的请假记录。")
    return "\n".join(lines)


def _assistant_leave_choice_actions(records, decision):
    """为当前审批人可见的申请生成安全的快捷选项，浏览器只回传受控文案。"""
    verb = "批准" if decision == "同意" else "驳回"
    return [{
        "label": f"{verb} #{record['request_id']} · {record['employee_name']}",
        "message": f"{verb}请假申请 {record['request_id']}",
    } for record in records[:6]]


def _assistant_leave_days(value):
    """解析简短请假天数，避免把自然语言日期误当作天数。"""
    chinese_numbers = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
    return int(value) if str(value).isdigit() else chinese_numbers.get(str(value), 1)


def _assistant_leave_dates(user, start_date, days):
    """从已发布班次中选择后续可申请日期，不允许对无班次日期创建请假。"""
    rows = query_all(
        """SELECT a.work_date FROM schedule_assignments a
           JOIN schedule_plans p ON p.plan_id=a.plan_id
           WHERE a.employee_id=%s AND p.status='已发布' AND a.work_date>=%s
             AND NOT EXISTS (
                 SELECT 1 FROM schedule_requests r
                 WHERE r.employee_id=a.employee_id AND r.work_date=a.work_date
                   AND r.status IN ('待部门审批','待人事审批','已批准')
             )
           ORDER BY a.work_date LIMIT %s""",
        (user["id"], start_date, days),
    )
    return [row["work_date"].isoformat() if hasattr(row["work_date"], "isoformat") else str(row["work_date"]) for row in rows]


def _assistant_submit_leave_requests(user, work_dates, reason):
    """一次创建多个单日请假申请，复用页面申请的班次和状态约束。"""
    if user["role"] == "super_admin":
        return None, "超级管理员无需提交个人排班申请"
    initial_status = "待人事审批" if user["role"] == "department_admin" else "待部门审批"
    request_ids = []
    with connect_database() as connection:
        try:
            connection.begin()
            with connection.cursor() as cursor:
                for work_date in work_dates:
                    cursor.execute(
                        """SELECT a.assignment_id,a.shift_id,p.version FROM schedule_assignments a
                           JOIN schedule_plans p ON p.plan_id=a.plan_id
                           WHERE a.employee_id=%s AND a.work_date=%s AND p.status='已发布' FOR UPDATE""",
                        (user["id"], work_date),
                    )
                    assignment = cursor.fetchone()
                    if assignment is None:
                        raise ValueError(f"{work_date} 没有已发布班次，无法提交请假")
                    cursor.execute(
                        """SELECT request_id FROM schedule_requests WHERE employee_id=%s AND work_date=%s
                           AND status IN ('待部门审批','待人事审批','已批准') FOR UPDATE""",
                        (user["id"], work_date),
                    )
                    if cursor.fetchone():
                        raise ValueError(f"{work_date} 已有进行中或已批准的申请")
                    cursor.execute(
                        """INSERT INTO schedule_requests
                           (employee_id,department,request_type,work_date,reason,assignment_id,original_shift_id,assignment_version,status)
                           VALUES (%s,%s,'leave',%s,%s,%s,%s,%s,%s)""",
                        (user["id"], user["department"], work_date, reason, assignment["assignment_id"],
                         assignment["shift_id"], assignment["version"], initial_status),
                    )
                    request_ids.append(cursor.lastrowid)
            connection.commit()
        except Exception:
            connection.rollback()
            raise
    if user["role"] == "department_admin":
        recipients = notification_recipient_ids("hr_admin", "super_admin", exclude_id=user["id"])
        target_url = "/pages/admin-schedule-requests.html"
    else:
        recipients = notification_recipient_ids("department_admin", department=user["department"], exclude_id=user["id"])
        target_url = "/pages/department-requests.html"
    notify_users_safely(
        recipients, "schedule_request", "新的请假申请",
        f"{user['name']}提交了 {len(request_ids)} 天请假申请，请及时审批。",
        target_url, related_type="schedule_request", related_id=request_ids[0],
        dedupe_key=f"assistant-leave-submitted:{request_ids[0]}",
    )
    return request_ids, None


@api_bp.route("/api/assistant/operations", methods=["POST"])
@api_operation_log("AI 助手业务操作")
@login_required
def assistant_operations():
    """处理助手内可审计的员工开户、请假、待办查询与请假审批。"""
    data = request.get_json(silent=True) or {}
    question = str(data.get("question", "")).strip()[:1000]
    if not question:
        return jsonify({"message": "请输入你的问题"}), 400
    user = current_user()
    operation_conversation_id = str(data.get("conversation_id") or "").strip() or None

    @after_this_request
    def persist_operation(response):
        return _persist_assistant_operation_response(user, question, operation_conversation_id, response)

    pending = session.get("assistant_pending_leave_review")
    pending_submission = session.get("assistant_pending_leave_submission")
    pending_details = session.get("assistant_pending_leave_details")
    pending_choice = session.get("assistant_pending_leave_choice")
    pending_registration_nonce = session.get("assistant_pending_employee_registration")
    pending_profile_update = session.get("assistant_pending_profile_update")
    pending_hr_profile_update = session.get("assistant_pending_hr_profile_update")
    pending_offboarding = session.get("assistant_pending_offboarding")
    pending_hr_offboarding = session.get("assistant_pending_hr_offboarding")
    pending_handover = session.get("assistant_pending_handover")

    # 员工开户是高风险写操作：必须由人事/超级管理员发起，并经过明确确认。
    if pending_registration_nonce and re.fullmatch(r"(?:确认(?:注册|开户)?|是|好的|执行|同意)", question):
        if user["role"] not in ("hr_admin", "super_admin"):
            session.pop("assistant_pending_employee_registration", None)
            return jsonify({"handled": True, "answer": "当前登录账号无权注册员工，仅人事或超级管理员可以执行。"}), 403
        session.pop("assistant_pending_employee_registration", None)
        registration_data = _assistant_take_registration(user["id"], pending_registration_nonce)
        if registration_data is None:
            return jsonify({"handled": True, "answer": "员工注册确认已过期，请重新提供注册信息。"}), 409
        employee, error_message, status_code = create_employee(registration_data)
        if error_message:
            return jsonify({"handled": True, "answer": error_message}), status_code
        app.logger.info(
            "AI 助手员工开户 request_id=%s actor=%s employee=%s",
            g.request_id, user["employee_no"], registration_data["employee_no"],
        )
        g.audit_details = {
            "tool": "employee_create", "target_employee_id": employee["id"],
            "before": None,
            "after": {"employee_no": employee["employee_no"], "name": employee["name"], "phone": f"{registration_data['phone'][:3]}****{registration_data['phone'][-4:]}"},
        }
        return jsonify({
            "handled": True,
            "answer": f"员工 {employee['name']} 已注册成功，工号：{employee['employee_no']}。",
            "action": "employee_register",
        }), 201

    if pending_registration_nonce and re.search(r"取消|算了|不要", question):
        session.pop("assistant_pending_employee_registration", None)
        _assistant_take_registration(user["id"], pending_registration_nonce)
        return jsonify({"handled": True, "answer": "已取消员工注册，不会写入任何员工数据。"})

    if re.search(r"(?:修改|更新|更换|变更|设置).{0,12}(?:工号|姓名|部门|岗位|角色)|(?:工号|姓名|部门|岗位|角色).{0,8}(?:修改|更新|更换|变更)", question):
        return jsonify({"handled": True, "answer": "工号、姓名、部门、岗位和角色属于核心员工信息，不能由员工通过 AI 自助修改，请联系人力资源部门办理。"})

    if _assistant_registration_intent(question):
        if user["role"] not in ("hr_admin", "super_admin"):
            return jsonify({"handled": True, "answer": "当前登录账号无权注册员工，仅人事或超级管理员可以执行。"}), 403
        registration_data = _assistant_registration_data(question)
        auto_employee_no = not registration_data.get("employee_no")
        if auto_employee_no:
            registration_data["employee_no"] = next_employee_no()
        missing = [
            label for key, label in (
                ("name", "姓名"), ("password", "密码"), ("phone", "手机号"),
            ) if not registration_data.get(key)
        ]
        if missing:
            return jsonify({
                "handled": True,
                "answer": f"请补充员工注册信息：{'、'.join(missing)}。示例：注册员工 小廖，工号 EMP006，密码 Abcd@123，手机号 13800138000。",
                "requires_details": True,
            })
        registration_data["confirm_password"] = registration_data["confirm_password"] or registration_data["password"]
        _, validation_error = validate_registration_data(registration_data)
        if validation_error:
            return jsonify({"handled": True, "answer": validation_error}), 400
        if query_one("SELECT id FROM employees WHERE employee_no=%s", (registration_data["employee_no"],)):
            return jsonify({"handled": True, "answer": "该工号已注册，请更换工号。"}), 409
        nonce = _assistant_store_registration(user["id"], registration_data)
        session["assistant_pending_employee_registration"] = nonce
        masked_phone = f"{registration_data['phone'][:3]}****{registration_data['phone'][-4:]}"
        email_text = f"，邮箱：{registration_data['email']}" if registration_data.get("email") else ""
        auto_text = "，系统已自动生成工号" if auto_employee_no else ""
        return jsonify({
            "handled": True,
            "answer": f"准备注册员工：{registration_data['name']}，工号：{registration_data['employee_no']}，手机号：{masked_phone}{email_text}{auto_text}。确认后将创建员工账号并保存密码哈希。",
            "requires_confirmation": True,
            "confirm_text": "确认注册",
        })

    if pending_profile_update and re.fullmatch(r"(?:确认(?:修改|更新)?|是|好的|执行|同意)", question):
        session.pop("assistant_pending_profile_update", None)
        changed, error_message = _assistant_update_own_profile(user, pending_profile_update)
        if error_message:
            return jsonify({"handled": True, "answer": error_message}), 400
        if not changed:
            return jsonify({"handled": True, "answer": "个人资料没有变化，未执行修改。"})
        return jsonify({"handled": True, "answer": "你的个人资料已更新成功。", "action": "profile_update"})

    if pending_hr_profile_update and re.fullmatch(r"(?:确认(?:修改|更新)?|是|好的|执行|同意)", question):
        session.pop("assistant_pending_hr_profile_update", None)
        target = find_employee(pending_hr_profile_update["target_id"])
        if target is None:
            return jsonify({"handled": True, "answer": "目标员工不存在，未执行修改。"}), 404
        updated, error_message, status_code = _assistant_apply_hr_profile_update(user, target["id"], pending_hr_profile_update["changes"])
        if error_message:
            return jsonify({"handled": True, "answer": error_message}), status_code
        return jsonify({"handled": True, "answer": f"员工 {updated['name']} 的档案已更新成功。", "action": "hr_profile_update"})

    if pending_offboarding and re.fullmatch(r"(?:确认(?:离职|提交)?|是|好的|执行|同意)", question):
        session.pop("assistant_pending_offboarding", None)
        task_id, error_message, status_code = _assistant_create_offboarding_request(user, pending_offboarding["reason"])
        if error_message:
            return jsonify({"handled": True, "answer": error_message}), status_code
        return jsonify({"handled": True, "answer": f"离职申请已提交，待部门经理审批（申请任务编号：{task_id}）。", "action": "offboarding_request"}), status_code

    if pending_hr_offboarding and re.fullmatch(r"(?:确认(?:离职|提交)?|是|好的|执行|同意)", question):
        session.pop("assistant_pending_hr_offboarding", None)
        updated, error_message, status_code = _assistant_hr_start_offboarding(
            user, pending_hr_offboarding["target_id"], pending_hr_offboarding["reason"],
        )
        if error_message:
            return jsonify({"handled": True, "answer": error_message}), status_code
        return jsonify({"handled": True, "answer": f"已为员工 {updated['name']} 发起离职流程，等待超级管理员确认。", "action": "hr_offboarding_request"}), 201

    if pending_handover and re.fullmatch(r"(?:确认(?:提交|交接)?|是|好的|执行|同意)", question):
        session.pop("assistant_pending_handover", None)
        if user["role"] != "employee":
            return jsonify({"handled": True, "answer": "当前角色不能提交个人交接。"}), 403
        if pending_handover.get("task_id"):
            if not _assistant_handover_task(user):
                return jsonify({"handled": True, "answer": "交接任务已被处理，请刷新后查看最新状态。"}), 409
            changed = _assistant_submit_handover(user, pending_handover["task_id"], pending_handover["content"])
            if not changed:
                return jsonify({"handled": True, "answer": "交接任务已被处理，请刷新后查看最新状态。"}), 409
            return jsonify({"handled": True, "answer": "交接说明已提交人事审核。", "action": "handover_submit"})
        task_id, error_message, status_code = _assistant_create_handover_request(user, pending_handover["content"])
        if error_message:
            return jsonify({"handled": True, "answer": error_message}), status_code
        return jsonify({"handled": True, "answer": f"工作交接申请已提交人事审核（任务编号：{task_id}）。", "action": "handover_request"}), status_code

    if (pending_profile_update or pending_hr_profile_update or pending_offboarding or pending_hr_offboarding or pending_handover) and re.search(r"取消|算了|不要", question):
        session.pop("assistant_pending_profile_update", None)
        session.pop("assistant_pending_hr_profile_update", None)
        session.pop("assistant_pending_offboarding", None)
        session.pop("assistant_pending_hr_offboarding", None)
        session.pop("assistant_pending_handover", None)
        return jsonify({"handled": True, "answer": "已取消本次操作，不会修改任何数据。"})

    if _assistant_profile_update_intent(question) and _assistant_employee_target(question):
        target = _assistant_employee_target(question)
        if user["role"] not in ("hr_admin", "super_admin"):
            return jsonify({"handled": True, "answer": "你只能通过 AI 修改自己的手机号或邮箱，修改其他员工档案请联系人力资源部门。"})
        changes = _assistant_hr_profile_changes(question)
        if not changes:
            return jsonify({"handled": True, "answer": "请提供要修改的员工档案字段，例如“修改员工 EMP006 的手机号为 13800138000”。", "requires_details": True})
        rows = _assistant_find_employee_target(target)
        if not rows:
            return jsonify({"handled": True, "answer": "未找到目标员工，请核对姓名或工号。"}), 404
        if len(rows) > 1:
            return jsonify({"handled": True, "answer": "找到多名同名员工，请使用工号指定目标员工。"}), 409
        target_employee = rows[0]
        session["assistant_pending_hr_profile_update"] = {"target_id": target_employee["id"], "changes": changes}
        changed_text = "，".join(f"{field}改为 {value}" for field, value in changes.items() if field != "role")
        if "role" in changes:
            role_name = ROLE_NAMES.get(changes["role"], changes["role"])
            changed_text = f"{changed_text}，" if changed_text else ""
            changed_text += f"角色改为 {role_name}"
        return jsonify({"handled": True, "answer": f"准备修改员工 {target_employee['name']}（{target_employee['employee_no']}）的档案：{changed_text}。确认后立即保存。", "requires_confirmation": True, "confirm_text": "确认修改"})

    if _assistant_profile_update_intent(question):
        changes = _assistant_profile_update_data(question)
        if not changes:
            return jsonify({"handled": True, "answer": "请提供新的手机号或邮箱，例如“把我的手机号改为 13800138000”。", "requires_details": True})
        if "phone" in changes and not PHONE_PATTERN.fullmatch(changes["phone"]):
            return jsonify({"handled": True, "answer": "手机号格式无效，请输入 11 位手机号。"}), 400
        if "email" in changes and changes["email"] and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", changes["email"]):
            return jsonify({"handled": True, "answer": "邮箱格式无效，请重新提供。"}), 400
        session["assistant_pending_profile_update"] = changes
        summary = []
        if "phone" in changes:
            summary.append(f"手机号改为 {changes['phone'][:3]}****{changes['phone'][-4:]}")
        if "email" in changes:
            summary.append(f"邮箱改为 {changes['email']}")
        return jsonify({"handled": True, "answer": f"准备修改你的个人资料：{'，'.join(summary)}。确认后立即保存。", "requires_confirmation": True, "confirm_text": "确认修改"})

    if re.search(r"(?:发起|申请|提交|我要|帮我|办理).{0,8}离职|离职申请", question):
        reason_match = re.search(r"(?:原因|理由)\s*(?:是|为|[:：])?\s*(.+)$", question)
        reason = reason_match.group(1).strip(" ，,。；;") if reason_match else ""
        if not reason:
            return jsonify({"handled": True, "answer": "请提供离职原因，例如“发起离职，原因是个人职业规划”。", "requires_details": True})
        target = _assistant_employee_target(question)
        if target and user["role"] in ("hr_admin", "super_admin"):
            rows = _assistant_find_employee_target(target)
            if not rows:
                return jsonify({"handled": True, "answer": "未找到目标员工，请核对姓名或工号。"}), 404
            if len(rows) > 1:
                return jsonify({"handled": True, "answer": "找到多名同名员工，请使用工号指定目标员工。"}), 409
            session["assistant_pending_hr_offboarding"] = {"target_id": rows[0]["id"], "reason": reason[:500]}
            return jsonify({"handled": True, "answer": f"准备为员工 {rows[0]['name']}（{rows[0]['employee_no']}）发起离职流程，原因：{reason[:500]}。确认后提交超级管理员审批。", "requires_confirmation": True, "confirm_text": "确认离职"})
        if user["role"] != "employee":
            return jsonify({"handled": True, "answer": "当前角色不能发起个人离职申请。"}), 403
        session["assistant_pending_offboarding"] = {"reason": reason[:500]}
        return jsonify({"handled": True, "answer": f"准备提交个人离职申请，原因：{reason[:500]}。提交后将进入部门审批流程，确认后继续。", "requires_confirmation": True, "confirm_text": "确认离职"})

    if re.search(r"(?:发起|申请|提交|我要|帮我|办理).{0,8}(?:交接)|(?:提交|完成).{0,8}交接", question):
        if user["role"] != "employee":
            return jsonify({"handled": True, "answer": "当前角色不能提交个人交接。"}), 403
        task = _assistant_handover_task(user)
        content_match = re.search(r"(?:交接内容|交接说明|说明|内容)\s*(?:是|为|[:：])?\s*(.+)$", question)
        content = content_match.group(1).strip(" ，,。；;") if content_match else ""
        if len(content) < 2:
            return jsonify({"handled": True, "answer": "请补充交接说明，例如“交接内容：已整理客户资料并移交给李四”。", "requires_details": True})
        session["assistant_pending_handover"] = {"task_id": task["task_id"] if task else None, "content": content[:2000]}
        action_text = "提交现有离职交接待办" if task else "发起工作交接申请"
        return jsonify({"handled": True, "answer": f"准备{action_text}：{content[:2000]}。确认后将提交人事审核。", "requires_confirmation": True, "confirm_text": "确认提交"})

    if pending_submission and re.fullmatch(r"(?:确认(?:提交)?|是|好的|执行|同意)", question):
        session.pop("assistant_pending_leave_submission", None)
        try:
            request_ids, error_message = _assistant_submit_leave_requests(
                user, pending_submission["work_dates"], pending_submission["reason"],
            )
        except ValueError as error:
            return jsonify({"handled": True, "answer": str(error)}), 409
        return jsonify({"handled": True, "answer": f"已提交 {len(request_ids)} 条请假申请，编号：{'、'.join(map(str, request_ids))}。"})
    approval_confirmation = pending and (
        (pending["decision"] == "同意" and re.fullmatch(r"(?:确认(?:批准|同意)?|是|好的|执行|同意|批准)", question))
        or (pending["decision"] == "驳回" and re.fullmatch(r"(?:确认(?:驳回|拒绝)?|是|好的|执行|驳回|拒绝)", question))
    )
    if approval_confirmation:
        session.pop("assistant_pending_leave_review", None)
        response = _agent_review_leave_request(user, int(pending["request_id"]), pending["decision"], pending.get("note", ""))
        payload = response.get_json() or {}
        return jsonify({"handled": True, "answer": payload.get("message", "操作完成"), "action": "leave_review"}), response.status_code
    if (pending or pending_submission or pending_details or pending_choice) and re.search(r"取消|不批|算了", question):
        session.pop("assistant_pending_leave_review", None)
        session.pop("assistant_pending_leave_submission", None)
        session.pop("assistant_pending_leave_details", None)
        session.pop("assistant_pending_leave_choice", None)
        return jsonify({"handled": True, "answer": "已取消本次请假审批，不会修改任何记录。"})

    if pending_choice:
        if re.search(r"查询|查看|多少|记录|待办|任务", question):
            session.pop("assistant_pending_leave_choice", None)
        else:
            _, records = _assistant_leave_query(user, "pending")
            selection = re.search(r"(?:申请(?:编号)?|编号|#)\s*(\d+)\s*(?:号)?", question)
            if selection is None and re.fullmatch(r"\d+", question):
                selection = re.fullmatch(r"\d+", question)
            selected = None
            if selection:
                selected = next((item for item in records or [] if item["request_id"] == int(selection.group(1))), None)
            elif re.search(r"第一个|第一条", question) and records:
                selected = records[0]
            elif re.search(r"第二个|第二条", question) and records and len(records) > 1:
                selected = records[1]
            elif records:
                selected = next((item for item in records if item["employee_name"] in question), None)
            if selected is None:
                return jsonify({
                    "handled": True,
                    "answer": "请说明要处理哪一条，例如“第一个”“申请 6”或直接说员工姓名。",
                    "actions": _assistant_leave_choice_actions(records or [], pending_choice["decision"]),
                })
            session.pop("assistant_pending_leave_choice", None)
            session["assistant_pending_leave_review"] = {
                "request_id": selected["request_id"], "decision": pending_choice["decision"], "note": "智能助手对话审批",
            }
            verb = "批准" if pending_choice["decision"] == "同意" else "驳回"
            work_date = selected["work_date"].isoformat() if hasattr(selected["work_date"], "isoformat") else str(selected["work_date"])
            return jsonify({
                "handled": True,
                "answer": f"已选中 #{selected['request_id']}：{selected['employee_name']} · {work_date}。确认后立即{verb}。",
                "confirm_text": f"确认{verb}",
            })

    if pending_details:
        # 用户改问查询、待办或审批时，不能继续把新问题强行解释为请假日期和原因。
        if re.search(r"查询|查看|多少|记录|待办|任务|审批|批准|驳回", question):
            session.pop("assistant_pending_leave_details", None)
            pending_details = None
        else:
            start_match = re.search(r"\d{4}-\d{2}-\d{2}", question)
            if not start_match or not parse_iso_date(start_match.group(0)):
                return jsonify({"handled": True, "answer": "请提供有效的开始日期和原因，例如：2026-09-02，个人事务。"})
            reason = question.replace(start_match.group(0), "").strip("，,。；;：: ")
            if len(reason) < 2:
                return jsonify({"handled": True, "answer": "还需要请假原因，例如：2026-09-02，个人事务。"})
            days = int(pending_details["days"])
            work_dates = _assistant_leave_dates(user, start_match.group(0), days)
            session.pop("assistant_pending_leave_details", None)
            if len(work_dates) < days:
                return jsonify({"handled": True, "answer": f"从 {start_match.group(0)} 起只找到 {len(work_dates)} 天可申请的已发布班次，无法提交 {days} 天请假。"})
            session["assistant_pending_leave_submission"] = {"work_dates": work_dates, "reason": reason}
            return jsonify({
                "handled": True,
            "answer": f"准备提交 {len(work_dates)} 天请假：{'、'.join(work_dates)}。原因：{reason}。确认后将创建申请。",
            "requires_confirmation": True,
            "confirm_text": "确认提交",
            })

    if re.search(r"批准|同意|通过|驳回|拒绝", question):
        _, pending_records = _assistant_leave_query(user, "pending")
        if pending_records is None:
            return jsonify({"handled": True, "answer": "当前登录账号没有请假审批权限，只能查看自己的请假记录。"}), 403
        # 申请编号必须带明确前缀，避免“市场部员工2”中的 2 被误认为申请编号。
        request_match = re.search(r"(?:申请(?:编号)?|编号|#)\s*(\d+)\s*(?:号)?", question)
        decision = "驳回" if re.search(r"驳回|拒绝", question) else "同意"
        record = None
        if request_match:
            record = next((item for item in pending_records if item["request_id"] == int(request_match.group(1))), None)
        else:
            # 名称、工号都只在当前用户可审批的待办里匹配，避免用自然语言越权定位员工。
            named_records = [item for item in pending_records if item["employee_name"] in question]
            if len(named_records) == 1:
                record = named_records[0]
            elif len(named_records) > 1:
                return jsonify({"handled": True, "answer": "找到多条同名员工的待审批请假，请补充申请编号后再操作。"})
        if record is None:
            if request_match:
                return jsonify({"handled": True, "answer": "该申请不存在、已处理或不属于当前账号的审批范围。"})
            if pending_records:
                candidates = "\n".join(f"#{item['request_id']} {item['employee_name']} · {item['work_date']}" for item in pending_records)
                session["assistant_pending_leave_choice"] = {"decision": decision}
                verb = "批准" if decision == "同意" else "驳回"
                return jsonify({
                    "handled": True,
                    "answer": f"请告诉我需要{verb}哪一条请假申请：\n{candidates}",
                    "actions": _assistant_leave_choice_actions(pending_records, decision),
                })
            return jsonify({"handled": True, "answer": "当前没有待审批的请假申请。"})
        session["assistant_pending_leave_review"] = {
            "request_id": record["request_id"], "decision": decision, "note": "智能助手对话审批",
        }
        action_name = "批准" if decision == "同意" else "驳回"
        work_date = record["work_date"].isoformat() if hasattr(record["work_date"], "isoformat") else str(record["work_date"])
        return jsonify({
            "handled": True,
            "answer": f"将{action_name} #{record['request_id']}：{record['employee_name']} · {work_date} 的请假申请。确认后立即执行。",
            "requires_confirmation": True,
            "confirm_text": f"确认{action_name}",
        })

    if re.search(r"(?:帮我|我要|我想|申请|提交).{0,10}请假|请假\s*(?:\d+|[一二三四五六七八九十]+)\s*(?:天|日)", question):
        duration = re.search(r"(\d+|[一二三四五六七八九十]+)\s*(?:天|日)", question)
        days = _assistant_leave_days(duration.group(1)) if duration else 1
        session["assistant_pending_leave_details"] = {"days": days}
        return jsonify({
            "handled": True,
            "answer": f"可以为你提交 {days} 天请假申请。请继续提供开始日期和请假原因，例如：2026-09-02，个人事务。提交前我会展示具体日期并请你确认。",
            "requires_details": True,
        })

    if re.fullmatch(r"(?:请假|休假|我要请假|帮我请假)", question):
        return jsonify({
            "handled": True,
            "answer": "你是想查看请假记录，还是提交新的请假申请？可以直接说“我的请假记录”或“帮我请假三天”。",
            "actions": [
                {"label": "查看我的请假", "message": "我的请假记录"},
                {"label": "提交请假", "message": "帮我请假"},
            ],
        })

    if re.search(r"请假|休假", question):
        if user["role"] == "employee":
            scope = "mine"
        elif re.search(r"待审批|待处理", question):
            scope = "pending"
        elif re.search(r"全部|全员|多少员工|多少人|所有", question):
            scope = "department" if user["role"] == "department_admin" else "all"
        elif re.search(r"本部门|部门", question):
            scope = "department" if user["role"] == "department_admin" else "all"
        else:
            scope = "mine"
        summary, records = _assistant_leave_query(user, scope, include_history=bool(re.search(r"历史|已驳回|全部记录", question)))
        if summary is None:
            return jsonify({"handled": True, "answer": "当前账号无权查询该范围的请假数据。"}), 403
        return jsonify({"handled": True, "answer": _assistant_leave_answer(scope, summary, records), "scope": scope})

    if re.search(r"待办|任务|事项", question):
        include_history = bool(re.search(r"历史|全部|已完成", question))
        conditions = ["employee_id=%s"]
        if not include_history:
            conditions.append("status NOT IN ('已完成','已驳回','部门已驳回','老板已驳回')")
        tasks = query_all(
            """SELECT task_id,title,status,priority,due_date FROM employee_tasks WHERE """
            + " AND ".join(conditions) + " ORDER BY task_id DESC LIMIT 20", (user["id"],)
        )
        lines = [f"我的待办共 {len(tasks)} 条。"]
        for task in tasks:
            due_date = task["due_date"].isoformat() if hasattr(task.get("due_date"), "isoformat") else (task.get("due_date") or "无截止日期")
            lines.append(f"#{task['task_id']} {task['title']} · {task['status']} · {task['priority']} · 截止：{due_date}")
        if not tasks:
            lines.append("当前没有待办事项。")
        return jsonify({"handled": True, "answer": "\n".join(lines), "scope": "mine"})
    return jsonify({"handled": False})


@api_bp.route("/api/assistant/chat", methods=["POST"])
@api_operation_log("AI 助手咨询")
@login_required
def assistant_chat():
    """只读运营助手：服务层仅接收已授权的最小业务上下文。"""
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"message": "请使用 JSON 格式提交问题"}), 400
    page = str(data.get("page", "工作台")).strip()[:80] or "工作台"
    output_format = data.get("output_format", "normal")
    if data.get("stream") is True:
        iterator, error_message, status_code = stream_assistant(
            current_user(), data.get("question"), page, data.get("conversation_id"), output_format, data.get("memory_type"), data.get("memory_options"),
        )
        if error_message:
            return jsonify({"message": error_message}), status_code

        def sse_event(event, payload):
            return f"event: {event}\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"

        def generate_stream():
            yield sse_event("meta", {"output_format": str(output_format).lower()})
            try:
                for event_name, payload in iterator:
                    yield sse_event(event_name, payload)
                    if event_name == "error":
                        return
                yield sse_event("done", {"success": True, "output_format": str(output_format).lower()})
            except Exception:
                yield sse_event("error", {"message": "AI 助手暂时不可用，请稍后再试"})
        return Response(generate_stream(), mimetype="text/event-stream", headers={
            "Cache-Control": "no-cache, no-transform", "X-Accel-Buffering": "no",
        })
    result, error_message, status_code = ask_assistant(
        current_user(), data.get("question"), page, data.get("conversation_id"), output_format, data.get("memory_type"), data.get("memory_options"),
    )
    if error_message:
        return jsonify({"message": error_message}), status_code
    return jsonify({"success": True, **result})


@api_bp.route("/api/assistant/database-query", methods=["POST"])
@api_operation_log("AI 助手业务数据库查询")
@roles_required("department_admin", "hr_admin", "super_admin", message="无权限：仅部门管理员、人事和老板可以查询业务数据库")
@validate_json_fields("sql")
def assistant_database_query():
    """受控 SQL 查询入口：只读、角色隔离、结果限额且不记录 SQL 原文。"""
    user = current_user()
    try:
        result = run_database_query(user, request.get_json()["sql"])
    except DatabaseQueryError as error:
        return jsonify({"message": str(error)}), 400
    app.logger.info(
        "业务数据库查询 request_id=%s actor=%s scope=%s tables=%s rows=%s",
        g.request_id, user["employee_no"], result["scope"], ",".join(result["tables"]), result["row_count"],
    )
    return jsonify({"success": True, **result})


@api_bp.route("/api/assistant/data-query", methods=["POST"])
@api_operation_log("AI 助手自然语言数据查询")
@roles_required("department_admin", "hr_admin", "super_admin", message="无权限：仅部门管理员、人事和老板可以查询业务数据")
@validate_json_fields("question")
def assistant_data_query():
    """管理员直接用自然语言查询；服务端只执行固定的受控查询模板。"""
    user = current_user()
    try:
        result = run_database_intent(user, request.get_json()["question"])
    except DatabaseQueryError as error:
        return jsonify({"message": str(error)}), 400
    if result is None:
        return jsonify({"success": True, "handled": False})
    app.logger.info(
        "自然语言数据查询 request_id=%s actor=%s intent=%s scope=%s tables=%s rows=%s",
        g.request_id, user["employee_no"], result["intent"], result["scope"],
        ",".join(result["tables"]), result["row_count"],
    )
    return jsonify({"success": True, "handled": True, **result})


@api_bp.route("/api/assistant/conversations/<conversation_id>", methods=["GET"])
@api_operation_log("读取 AI 对话")
@login_required
def assistant_conversation(conversation_id):
    """只允许读取当前登录用户自己的近期对话消息。"""
    result, error_message, status_code = get_conversation_messages(current_user()["id"], conversation_id)
    if error_message:
        return jsonify({"message": error_message}), status_code
    return jsonify({"success": True, **result})


@api_bp.route("/api/assistant/conversations", methods=["GET"])
@api_operation_log("读取 AI 会话列表")
@login_required
def assistant_conversation_list():
    return jsonify({"success": True, "conversations": list_conversations(current_user()["id"])})


@api_bp.route("/api/assistant/conversations/<conversation_id>", methods=["PUT"])
@api_operation_log("修改 AI 会话")
@login_required
def assistant_conversation_update(conversation_id):
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"message": "请使用 JSON 格式提交修改内容"}), 400
    user_id = current_user()["id"]
    if "title" in data:
        result, error_message, status_code = update_conversation_title(user_id, conversation_id, data["title"])
    elif "memory_summary" in data:
        result, error_message, status_code = update_conversation_summary(user_id, conversation_id, data["memory_summary"])
    else:
        return jsonify({"message": "仅支持修改 title 或 memory_summary"}), 400
    if error_message:
        return jsonify({"message": error_message}), status_code
    return jsonify({"success": True, **result})


@api_bp.route("/api/assistant/conversations/<conversation_id>", methods=["DELETE"])
@api_operation_log("删除 AI 会话")
@login_required
def assistant_conversation_delete(conversation_id):
    error_message, status_code = delete_conversation(current_user()["id"], conversation_id)
    if error_message:
        return jsonify({"message": error_message}), status_code
    return jsonify({"success": True})


@api_bp.route("/api/auth/logout", methods=["POST"])
@login_required
def logout():
    session.clear()
    return jsonify({"success": True, "message": "已退出登录"})


@api_bp.route("/api/auth/me", methods=["GET"])
@api_operation_log("读取当前用户")
@login_required
def get_current_user():
    return jsonify({"success": True, "user": current_user(), "csrf_token": session.get("csrf_token")})


@api_bp.route("/api/agent/session-token", methods=["GET"])
@login_required
def issue_agent_session_token():
    """为当前登录用户签发仅供本机智能办公助手使用的短期凭证。"""
    user = current_user()
    serializer = URLSafeTimedSerializer(SECRET_KEY, salt="smart-office-session")
    token = serializer.dumps({"aud": "smart-office", "employee_id": user["id"]})
    return jsonify({
        "success": True,
        "token": token,
        "expires_in": OFFICE_TOKEN_MAX_AGE_SECONDS,
        "user": user,
    })


@api_bp.route("/api/auth/csrf", methods=["GET"])
@login_required
def get_csrf_token():
    token = session.get("csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["csrf_token"] = token
    return jsonify({"success": True, "csrf_token": token})


@api_bp.route("/api/notifications", methods=["GET"])
@api_operation_log("读取站内消息")
@login_required
def get_notifications():
    user = current_user()
    try:
        limit = min(max(int(request.args.get("limit", "50")), 1), 100)
    except (TypeError, ValueError):
        return jsonify({"message": "limit 必须是 1 至 100 的整数"}), 400
    unread_only = request.args.get("unread_only", "").lower() in ("1", "true", "yes")
    conditions, parameters = ["recipient_id=%s"], [user["id"]]
    if unread_only:
        conditions.append("is_read=0")
    where_sql = " WHERE " + " AND ".join(conditions)
    records = query_all(
        "SELECT * FROM notifications" + where_sql + " ORDER BY created_at DESC,notification_id DESC LIMIT %s",
        [*parameters, limit],
    )
    unread_count = query_one(
        "SELECT COUNT(*) AS total FROM notifications WHERE recipient_id=%s AND is_read=0",
        (user["id"],),
    )["total"]
    return jsonify({
        "success": True,
        "unread_count": unread_count,
        "notifications": [serialize_notification(record) for record in records],
    })


@api_bp.route("/api/notifications/unread-count", methods=["GET"])
@login_required
def get_notification_unread_count():
    total = query_one(
        "SELECT COUNT(*) AS total FROM notifications WHERE recipient_id=%s AND is_read=0",
        (current_user()["id"],),
    )["total"]
    return jsonify({"success": True, "unread_count": total})


@api_bp.route("/api/notifications/<int:notification_id>/read", methods=["POST"])
@login_required
def read_notification(notification_id):
    user = current_user()
    record = query_one(
        "SELECT * FROM notifications WHERE notification_id=%s AND recipient_id=%s",
        (notification_id, user["id"]),
    )
    if record is None:
        return jsonify({"message": "消息不存在或不属于当前用户"}), 404
    if not record["is_read"]:
        execute(
            "UPDATE notifications SET is_read=1,read_at=NOW() WHERE notification_id=%s AND recipient_id=%s",
            (notification_id, user["id"]),
        )
    record = query_one(
        "SELECT * FROM notifications WHERE notification_id=%s AND recipient_id=%s",
        (notification_id, user["id"]),
    )
    return jsonify({"success": True, "notification": serialize_notification(record)})


@api_bp.route("/api/notifications/read-all", methods=["POST"])
@login_required
def read_all_notifications():
    _, updated = execute(
        "UPDATE notifications SET is_read=1,read_at=NOW() WHERE recipient_id=%s AND is_read=0",
        (current_user()["id"],),
    )
    return jsonify({"success": True, "message": "全部消息已标为已读", "updated": updated})


@api_bp.route("/api/audit-logs", methods=["GET"])
@super_admin_required
def get_audit_logs():
    """超级管理员按时间查看关键操作，审计记录只允许追加，不提供修改或删除接口。"""
    try:
        page = max(1, int(request.args.get("page", "1")))
        page_size = min(100, max(1, int(request.args.get("page_size", "30"))))
    except ValueError:
        return jsonify({"message": "page 和 page_size 必须是整数"}), 400
    actor_employee_no = request.args.get("actor_employee_no", "").strip()
    action = request.args.get("action", "").strip()
    conditions, parameters = [], []
    if actor_employee_no:
        conditions.append("actor_employee_no LIKE %s")
        parameters.append(f"%{actor_employee_no}%")
    if action:
        conditions.append("action LIKE %s")
        parameters.append(f"%{action}%")
    where_sql = " WHERE " + " AND ".join(conditions) if conditions else ""
    total = query_one("SELECT COUNT(*) AS total FROM audit_logs" + where_sql, parameters)["total"]
    records = query_all(
        "SELECT * FROM audit_logs" + where_sql + " ORDER BY audit_id DESC LIMIT %s OFFSET %s",
        [*parameters, page_size, (page - 1) * page_size],
    )
    return jsonify({"success": True, "total": total, "page": page, "page_size": page_size, "logs": records})


@api_bp.route("/api/employees/register", methods=["POST", "OPTIONS"])
@api_operation_log("员工开户注册")
@hr_required
@validate_json_fields("name", "password", "confirm_password", "phone")
def register_employee():
    if request.method == "OPTIONS": return "", 204
    data = request.get_json(silent=True) or {}
    if not str(data.get("employee_no", "")).strip():
        data["employee_no"] = next_employee_no()
    employee, error_message, status_code = create_employee(data)
    if error_message: return jsonify({"message": error_message}), status_code
    return jsonify({"message": "员工注册成功，工号已由系统自动生成", "employee": employee}), status_code


@api_bp.route("/api/employees", methods=["GET"])
@hr_required
def get_employees():
    conditions, parameters = [], []
    for field in ("department", "employee_no", "name", "employment_status"):
        value = request.args.get(field, "").strip()
        if value:
            conditions.append(f"{field} LIKE %s"); parameters.append(f"%{value}%")
    sql = "SELECT * FROM employees" + (" WHERE " + " AND ".join(conditions) if conditions else "") + " ORDER BY id"
    employees = query_all(sql, parameters)
    return jsonify({"total": len(employees), "employees": [public_employee(employee) for employee in employees]})


@api_bp.route("/api/employees/<int:employee_id>", methods=["GET"])
@hr_required
def get_employee(employee_id):
    employee = find_employee(employee_id)
    if employee is None: return jsonify({"message": "员工不存在"}), 404
    return jsonify({"success": True, "employee": public_employee(employee)})


@api_bp.route("/api/employees/<int:employee_id>", methods=["PUT"])
@hr_required
def update_employee(employee_id):
    data = request.get_json(silent=True)
    if isinstance(data, dict) and data.get("role") in {"super_admin", "hr_admin", "department_admin"} and current_user()["role"] != "super_admin":
        return jsonify({"message": "无权限：仅超级管理员可以授予管理角色"}), 403
    employee, error_message, status_code = update_employee_data(employee_id, data)
    if error_message: return jsonify({"message": error_message}), status_code
    return jsonify({"success": True, "employee": employee}), status_code


# ==================== 排班基础信息：部门、班次、日历和规则 ====================

@api_bp.route("/api/departments", methods=["GET", "POST"])
@login_required
def departments():
    if request.method == "GET":
        return jsonify({"success": True, "departments": query_all("SELECT * FROM departments ORDER BY department_id")})
    if current_user()["role"] not in ("hr_admin", "super_admin"):
        return jsonify({"message": "无权限：仅人事部门可以维护部门"}), 403
    data = request.get_json(silent=True)
    if not isinstance(data, dict): return jsonify({"message": "请使用 JSON 格式提交部门信息"}), 400
    name, description = str(data.get("name", "")).strip(), str(data.get("description", "")).strip()
    if not name: return jsonify({"message": "部门名称不能为空"}), 400
    parent_id = data.get("parent_id")
    if parent_id not in (None, ""):
        try: parent_id = int(parent_id)
        except (TypeError, ValueError): return jsonify({"message": "上级部门编号无效"}), 400
        if not query_one("SELECT department_id FROM departments WHERE department_id=%s", (parent_id,)):
            return jsonify({"message": "上级部门不存在"}), 404
    else:
        parent_id = None
    if query_one("SELECT department_id FROM departments WHERE name=%s", (name,)):
        return jsonify({"message": "该部门已存在"}), 409
    department_id, _ = execute("INSERT INTO departments (name, parent_id, description) VALUES (%s, %s, %s)", (name, parent_id, description))
    return jsonify({"success": True, "department": query_one("SELECT * FROM departments WHERE department_id=%s", (department_id,))}), 201


@api_bp.route("/api/departments/<int:department_id>", methods=["PUT", "DELETE"])
@hr_required
def department_detail(department_id):
    department = query_one("SELECT * FROM departments WHERE department_id=%s", (department_id,))
    if department is None: return jsonify({"message": "部门不存在"}), 404
    if request.method == "DELETE":
        if query_one("SELECT id FROM employees WHERE department=%s LIMIT 1", (department["name"],)) or query_one("SELECT department_id FROM departments WHERE parent_id=%s LIMIT 1", (department_id,)):
            return jsonify({"message": "部门仍有关联员工或下级部门，无法删除"}), 409
        execute("DELETE FROM departments WHERE department_id=%s", (department_id,))
        return jsonify({"success": True, "message": "部门已删除"})
    data = request.get_json(silent=True)
    if not isinstance(data, dict): return jsonify({"message": "请使用 JSON 格式提交部门信息"}), 400
    name, description = str(data.get("name", department["name"])).strip(), str(data.get("description", department["description"])).strip()
    if not name: return jsonify({"message": "部门名称不能为空"}), 400
    if name != department["name"] and query_one("SELECT department_id FROM departments WHERE name=%s", (name,)):
        return jsonify({"message": "该部门名称已存在"}), 409
    execute("UPDATE departments SET name=%s, description=%s WHERE department_id=%s", (name, description, department_id))
    if name != department["name"]: execute("UPDATE employees SET department=%s WHERE department=%s", (name, department["name"]))
    return jsonify({"success": True, "department": query_one("SELECT * FROM departments WHERE department_id=%s", (department_id,))})


@api_bp.route("/api/shifts", methods=["GET", "POST"])
@login_required
def shifts():
    if request.method == "GET":
        return jsonify({"success": True, "shifts": [serialize_shift(shift) for shift in query_all("SELECT * FROM shift_types WHERE is_active=1 ORDER BY shift_id")]})
    if current_user()["role"] not in ("hr_admin", "super_admin"):
        return jsonify({"message": "无权限：仅人事部门可以维护班次"}), 403
    data = request.get_json(silent=True)
    if not isinstance(data, dict): return jsonify({"message": "请使用 JSON 格式提交班次信息"}), 400
    shift, error_message = validate_shift_data(data)
    if error_message: return jsonify({"message": error_message}), 400
    if query_one("SELECT shift_id FROM shift_types WHERE name=%s", (shift["name"],)):
        return jsonify({"message": "该班次名称已存在"}), 409
    shift_id, _ = execute("INSERT INTO shift_types (name,start_time,end_time,duration_hours,color) VALUES (%s,%s,%s,%s,%s)", tuple(shift.values()))
    return jsonify({"success": True, "shift": serialize_shift(query_one("SELECT * FROM shift_types WHERE shift_id=%s", (shift_id,)))}), 201


@api_bp.route("/api/shifts/<int:shift_id>", methods=["PUT", "DELETE"])
@hr_required
def shift_detail(shift_id):
    existing = query_one("SELECT * FROM shift_types WHERE shift_id=%s", (shift_id,))
    if existing is None: return jsonify({"message": "班次不存在"}), 404
    if request.method == "DELETE":
        execute("UPDATE shift_types SET is_active=0 WHERE shift_id=%s", (shift_id,))
        return jsonify({"success": True, "message": "班次已停用"})
    data = request.get_json(silent=True)
    if not isinstance(data, dict): return jsonify({"message": "请使用 JSON 格式提交班次信息"}), 400
    shift, error_message = validate_shift_data({**existing, **data})
    if error_message: return jsonify({"message": error_message}), 400
    if shift["name"] != existing["name"] and query_one("SELECT shift_id FROM shift_types WHERE name=%s", (shift["name"],)):
        return jsonify({"message": "该班次名称已存在"}), 409
    execute("UPDATE shift_types SET name=%s,start_time=%s,end_time=%s,duration_hours=%s,color=%s WHERE shift_id=%s", (*shift.values(), shift_id))
    return jsonify({"success": True, "shift": serialize_shift(query_one("SELECT * FROM shift_types WHERE shift_id=%s", (shift_id,)))})


@api_bp.route("/api/work-calendar", methods=["GET", "PUT"])
@login_required
def work_calendar():
    if request.method == "GET":
        start, end = request.args.get("start", ""), request.args.get("end", "")
        if (start and not parse_iso_date(start)) or (end and not parse_iso_date(end)): return jsonify({"message": "日期必须为 YYYY-MM-DD"}), 400
        sql, parameters = "SELECT * FROM work_calendar", []
        if start and end: sql += " WHERE work_date BETWEEN %s AND %s"; parameters = [start, end]
        calendar_records = query_all(sql + " ORDER BY work_date", parameters)
        for record in calendar_records:
            if hasattr(record.get("work_date"), "isoformat"):
                record["work_date"] = record["work_date"].isoformat()
        return jsonify({"success": True, "calendar": calendar_records})
    if current_user()["role"] not in ("hr_admin", "super_admin"):
        return jsonify({"message": "无权限：仅人事部门可以维护工作日历"}), 403
    data = request.get_json(silent=True)
    if not isinstance(data, dict) or not parse_iso_date(data.get("work_date")): return jsonify({"message": "work_date 必须为 YYYY-MM-DD"}), 400
    if not isinstance(data.get("is_workday"), bool): return jsonify({"message": "is_workday 必须为布尔值"}), 400
    execute("""INSERT INTO work_calendar (work_date,is_workday,label) VALUES (%s,%s,%s)
               ON DUPLICATE KEY UPDATE is_workday=VALUES(is_workday),label=VALUES(label)""", (data["work_date"], data["is_workday"], str(data.get("label", "")).strip()))
    return jsonify({"success": True, "calendar": query_one("SELECT * FROM work_calendar WHERE work_date=%s", (data["work_date"],))})


@api_bp.route("/api/scheduling-rules", methods=["GET", "PUT"])
@login_required
def scheduling_rules():
    rule = query_one("SELECT * FROM scheduling_rules WHERE rule_name='默认排班规则'")
    if request.method == "GET": return jsonify({"success": True, "rule": rule})
    if current_user()["role"] not in ("hr_admin", "super_admin"):
        return jsonify({"message": "无权限：仅人事部门可以维护排班规则"}), 403
    data = request.get_json(silent=True)
    if not isinstance(data, dict): return jsonify({"message": "请使用 JSON 格式提交排班规则"}), 400
    try:
        max_days = int(data.get("max_consecutive_days", rule["max_consecutive_days"]))
        daily_hours = float(data.get("max_daily_overtime_hours", rule["max_daily_overtime_hours"]))
        monthly_hours = float(data.get("max_monthly_overtime_hours", rule["max_monthly_overtime_hours"]))
        min_staff = int(data.get("min_staff_per_position", rule["min_staff_per_position"]))
        min_rest_hours = float(data.get("min_rest_hours", rule.get("min_rest_hours", 10)))
    except (TypeError, ValueError): return jsonify({"message": "排班规则数值格式无效"}), 400
    if not 1 <= max_days <= 31 or not 0 <= min_rest_hours <= 24 or daily_hours < 0 or monthly_hours < 0 or min_staff < 1:
        return jsonify({"message": "排班规则超出允许范围"}), 400
    execute("""UPDATE scheduling_rules SET max_consecutive_days=%s,min_rest_hours=%s,
               max_daily_overtime_hours=%s,max_monthly_overtime_hours=%s,min_staff_per_position=%s
               WHERE rule_id=%s""", (max_days, min_rest_hours, daily_hours, monthly_hours, min_staff, rule["rule_id"]))
    return jsonify({"success": True, "rule": query_one("SELECT * FROM scheduling_rules WHERE rule_id=%s", (rule["rule_id"],))})


# ==================== 排班计划：草稿、校验、发布 ====================

@api_bp.route("/api/schedule-cycles", methods=["GET"])
@schedule_owner_required
def schedule_cycles():
    anchor = parse_iso_date(request.args.get("anchor", date.today().isoformat()))
    if not anchor:
        return jsonify({"message": "anchor 必须为 YYYY-MM-DD"}), 400
    cycle = get_schedule_cycle(anchor)
    return jsonify({"success": True, "cycle": serialize_cycle(cycle)})


@api_bp.route("/api/scheduling-cycle-policy", methods=["GET"])
@schedule_owner_required
def scheduling_cycle_policy():
    current = get_schedule_cycle(date.today())
    pending = query_one("""SELECT r.*, e.name AS requested_by_name FROM scheduling_cycle_change_requests r
                           JOIN employees e ON e.id=r.requested_by WHERE r.status='待老板审批'
                           ORDER BY r.request_id DESC LIMIT 1""")
    payload = {"success": True, "default_start_day": int(current["start_day"]), "current_cycle": serialize_cycle(current)}
    if pending and current_user()["role"] in ("hr_admin", "super_admin"):
        payload["pending"] = serialize_cycle_request(pending)
    return jsonify(payload)


@api_bp.route("/api/scheduling-cycle-policy/requests", methods=["POST"])
@hr_required
def request_scheduling_cycle_change():
    data = request.get_json(silent=True) or {}
    try:
        start_day = int(data.get("start_day"))
    except (TypeError, ValueError):
        return jsonify({"message": "周期开始日必须是 1 至 28 的整数"}), 400
    if not 1 <= start_day <= 28:
        return jsonify({"message": "周期开始日仅支持 1 至 28 日，避免短月日期缺失"}), 400
    current, transition_start, effective_start = cycle_change_preview(start_day)
    if start_day == int(current["start_day"]):
        return jsonify({"message": "新规则与当前周期规则相同，无需提交"}), 400
    if query_one("SELECT request_id FROM scheduling_cycle_change_requests WHERE status='待老板审批'"):
        return jsonify({"message": "已有待老板审批的周期变更，请先完成该审批"}), 409
    request_id, _ = execute("""INSERT INTO scheduling_cycle_change_requests
        (proposed_start_day,requested_by,transition_start,effective_start) VALUES (%s,%s,%s,%s)""",
        (start_day, current_user()["id"], transition_start, effective_start))
    notify_users_safely(
        notification_recipient_ids("super_admin", exclude_id=current_user()["id"]),
        "cycle_policy_review", "排班周期规则等待审批",
        f"人事申请将排班周期起始日调整为每月 {start_day} 日，请审核切换影响。",
        "/pages/schedule.html", related_type="cycle_change_request", related_id=request_id,
        dedupe_key=f"cycle-policy-submitted:{request_id}",
    )
    record = query_one("SELECT * FROM scheduling_cycle_change_requests WHERE request_id=%s", (request_id,))
    return jsonify({"success": True, "message": "周期规则变更已提交老板审批；批准前不会影响任何部门班表。", "request": serialize_cycle_request(record)}), 201


@api_bp.route("/api/scheduling-cycle-policy/requests/<int:request_id>/review", methods=["POST"])
@roles_required("super_admin", message="无权限：仅老板可以审批周期规则变更")
def review_scheduling_cycle_change(request_id):
    record = query_one("SELECT * FROM scheduling_cycle_change_requests WHERE request_id=%s", (request_id,))
    if not record or record["status"] != "待老板审批":
        return jsonify({"message": "该周期变更不存在或已处理"}), 404
    data = request.get_json(silent=True) or {}
    decision, note = str(data.get("decision", "")).strip(), str(data.get("note", "")).strip()[:500]
    if decision not in ("approve", "reject"):
        return jsonify({"message": "decision 必须为 approve 或 reject"}), 400
    if decision == "reject":
        execute("UPDATE scheduling_cycle_change_requests SET status='已驳回',reviewed_by=%s,reviewed_at=NOW(),review_note=%s WHERE request_id=%s", (current_user()["id"], note, request_id))
        notify_users_safely(
            [record["requested_by"]], "cycle_policy_result", "排班周期规则变更已驳回",
            f"调整为每月 {record['proposed_start_day']} 日起算的申请未获批准"
            f"{f'。审批意见：{note}' if note else '。'}",
            "/pages/schedule.html", related_type="cycle_change_request", related_id=request_id,
            dedupe_key=f"cycle-policy-result:{request_id}:rejected",
        )
        return jsonify({"success": True, "message": "已驳回周期规则变更"})
    # 已发布班表属于员工已可见的正式承诺，绝不自动改变；未发布版本可安全归档并按新周期重排。
    published = query_one("""SELECT COUNT(*) AS total FROM schedule_plans
                             WHERE end_date >= %s AND status='已发布'""", (record["transition_start"],))
    if published["total"]:
        return jsonify({"message": f"切换日之后仍有 {published['total']} 份已发布班表。请先由人事撤回这些正式班表，再批准周期变更。"}), 409
    _, archived_total = execute("""UPDATE schedule_plans
        SET status='已归档', plan_key=NULL, hr_note='排班周期规则已变更，请按新周期重新排班', version=version+1
        WHERE end_date >= %s AND status IN ('草稿','已打回','待人事审核')""", (record["transition_start"],))
    execute("DELETE FROM scheduling_cycles WHERE start_date >= %s", (record["transition_start"],))
    if record["transition_start"] < record["effective_start"]:
        execute("""INSERT INTO scheduling_cycles (start_date,end_date,start_day,cycle_kind)
                   VALUES (%s,%s,%s,'transition')""", (record["transition_start"], record["effective_start"] - timedelta(days=1), record["proposed_start_day"]))
    start = record["effective_start"]
    target = date.today() + timedelta(days=1095)
    while start <= target:
        end = cycle_end_for_start(start, int(record["proposed_start_day"]))
        execute("INSERT INTO scheduling_cycles (start_date,end_date,start_day,cycle_kind) VALUES (%s,%s,%s,'standard')", (start, end, record["proposed_start_day"]))
        start = end + timedelta(days=1)
    execute("""UPDATE scheduling_cycle_change_requests SET status='已批准',reviewed_by=%s,reviewed_at=NOW(),review_note=%s
               WHERE request_id=%s""", (current_user()["id"], note, request_id))
    notify_users_safely(
        [record["requested_by"]], "cycle_policy_result", "排班周期规则变更已批准",
        f"新的排班周期规则已批准，将从 {record['effective_start']} 起按每月 {record['proposed_start_day']} 日起算。",
        "/pages/schedule.html", related_type="cycle_change_request", related_id=request_id,
        dedupe_key=f"cycle-policy-result:{request_id}:approved",
    )
    return jsonify({"success": True, "message": f"周期规则已批准并生效。已归档 {archived_total} 份未发布旧班表，系统已保留切换期；之后请按新周期重新排班。"})


@api_bp.route("/api/schedule-cycles/<int:cycle_id>/summary", methods=["GET"])
@schedule_owner_required
def schedule_cycle_summary(cycle_id):
    cycle = query_one("SELECT * FROM scheduling_cycles WHERE cycle_id=%s", (cycle_id,))
    if not cycle:
        return jsonify({"message": "排班周期不存在"}), 404
    user, department = current_user(), request.args.get("department", "").strip()
    if user["role"] == "department_admin": department = user["department"]
    if department and not may_manage_department(user, department):
        return jsonify({"message": "无权查看该部门周期统计"}), 403
    params, dept_filter = [cycle["start_date"], cycle["end_date"]], ""
    if department:
        dept_filter = " AND e.department=%s"; params.append(department)
    rows = query_all(f"""SELECT e.department, COUNT(DISTINCT e.id) AS employee_count,
                 COUNT(CASE WHEN lr.employee_id IS NULL THEN a.assignment_id END) AS assignment_count,
                 COUNT(DISTINCT CASE WHEN lr.employee_id IS NULL THEN CONCAT(a.employee_id,'-',a.work_date) END) AS scheduled_days,
                 COALESCE(SUM(CASE WHEN lr.employee_id IS NULL THEN s.duration_hours ELSE 0 END),0) AS scheduled_hours,
                 COUNT(DISTINCT CASE WHEN lr.employee_id IS NOT NULL THEN CONCAT(a.employee_id,'-',a.work_date) END) AS approved_leave_days
              FROM employees e LEFT JOIN schedule_plans p ON p.department=e.department AND p.start_date=%s AND p.end_date=%s
                   AND p.plan_key IS NOT NULL AND p.status='已发布'
              LEFT JOIN schedule_assignments a ON a.plan_id=p.plan_id AND a.employee_id=e.id AND a.work_date BETWEEN %s AND %s
              LEFT JOIN shift_types s ON s.shift_id=a.shift_id
              LEFT JOIN (SELECT employee_id,work_date FROM schedule_requests
                         WHERE request_type='leave' AND status='已批准' GROUP BY employee_id,work_date) lr
                ON lr.employee_id=a.employee_id AND lr.work_date=a.work_date
              WHERE e.role<>'super_admin' AND e.employment_status<>'已离职'{dept_filter}
              GROUP BY e.department ORDER BY e.department""", [cycle["start_date"], cycle["end_date"], *params])
    rest_days = query_one("SELECT COUNT(*) AS total FROM work_calendar WHERE work_date BETWEEN %s AND %s AND is_workday=0", (cycle["start_date"], cycle["end_date"]))["total"]
    cycle_days = (cycle["end_date"] - cycle["start_date"]).days + 1
    work_days = max(cycle_days - rest_days, 0)
    statuses = query_all("""SELECT department,status,COUNT(*) AS total FROM schedule_plans
                            WHERE start_date=%s AND end_date=%s GROUP BY department,status""", (cycle["start_date"], cycle["end_date"]))
    status_map = {}
    for item in statuses: status_map.setdefault(item["department"], {})[item["status"]] = item["total"]
    draft_parameters, draft_filter = [cycle["start_date"], cycle["end_date"]], ""
    if department:
        draft_filter = " AND p.department=%s"
        draft_parameters.append(department)
    draft_rows = query_all(f"""SELECT p.department,COUNT(a.assignment_id) AS assignment_count,
                                      COALESCE(SUM(s.duration_hours),0) AS scheduled_hours
                               FROM schedule_plans p
                               LEFT JOIN schedule_assignments a ON a.plan_id=p.plan_id
                               LEFT JOIN shift_types s ON s.shift_id=a.shift_id
                               WHERE p.start_date=%s AND p.end_date=%s AND p.plan_key IS NOT NULL
                                 AND p.status IN ('草稿','已打回'){draft_filter}
                               GROUP BY p.department""", draft_parameters)
    draft_map = {item["department"]: item for item in draft_rows}
    for row in rows:
        row["scheduled_hours"] = float(row["scheduled_hours"] or 0)
        draft = draft_map.get(row["department"], {})
        row["draft_assignment_count"] = int(draft.get("assignment_count") or 0)
        row["draft_scheduled_hours"] = float(draft.get("scheduled_hours") or 0)
        expected = row["employee_count"] * work_days
        row["work_days"] = work_days
        row["coverage_rate"] = round(row["scheduled_days"] / expected * 100, 1) if expected else 0
        row["unassigned_employee_days"] = max(expected - row["scheduled_days"], 0)
        row["avg_shifts_per_employee"] = round(row["assignment_count"] / row["employee_count"], 1) if row["employee_count"] else 0
        row["avg_hours_per_employee"] = round(row["scheduled_hours"] / row["employee_count"], 1) if row["employee_count"] else 0
        row["plan_statuses"] = status_map.get(row["department"], {})
    return jsonify({"success": True, "cycle": serialize_cycle(cycle), "departments": rows,
                    "totals": {"employees": sum(r["employee_count"] for r in rows),
                               "assignments": sum(r["assignment_count"] for r in rows),
                               "hours": sum(r["scheduled_hours"] for r in rows),
                               "draft_assignments": sum(r["draft_assignment_count"] for r in rows),
                               "draft_hours": sum(r["draft_scheduled_hours"] for r in rows)}})

@api_bp.route("/api/schedule-plans", methods=["GET", "POST"])
@schedule_owner_required
def schedule_plans():
    user = current_user()
    if request.method == "GET":
        department = request.args.get("department", "").strip()
        status = request.args.get("status", "").strip()
        try: limit = min(max(int(request.args.get("limit", "100")), 1), 100)
        except (TypeError, ValueError): limit = 100
        if user["role"] == "department_admin": department = user["department"]
        sql, parameters = "SELECT * FROM schedule_plans", []
        conditions = ["plan_key IS NOT NULL"]
        if department: conditions.append("department=%s"); parameters.append(department)
        if status: conditions.append("status=%s"); parameters.append(status)
        if conditions: sql += " WHERE " + " AND ".join(conditions)
        return jsonify({"success": True, "plans": [serialize_schedule_plan(plan) for plan in query_all(sql + " ORDER BY updated_at DESC, plan_id DESC LIMIT %s", parameters + [limit])]})
    data = request.get_json(silent=True)
    if not isinstance(data, dict): return jsonify({"message": "请使用 JSON 格式提交排班计划"}), 400
    name, department = str(data.get("name", "")).strip(), str(data.get("department", "")).strip()
    if not may_manage_department(user, department): return jsonify({"message": "部门经理只能创建本部门班表"}), 403
    start_date, end_date = parse_iso_date(data.get("start_date")), parse_iso_date(data.get("end_date"))
    if not name or not start_date or not end_date: return jsonify({"message": "计划名称、开始日期和结束日期不能为空，日期格式为 YYYY-MM-DD"}), 400
    if start_date > end_date: return jsonify({"message": "开始日期不能晚于结束日期"}), 400
    cycle, cycle_error = validate_cycle_range(start_date, end_date)
    if cycle_error: return jsonify({"message": cycle_error}), 400
    if department and not query_one("SELECT department_id FROM departments WHERE name=%s", (department,)):
        return jsonify({"message": "计划部门不存在，请先创建部门"}), 404
    plan_key = schedule_plan_key(department, start_date, end_date)
    if query_one("SELECT plan_id FROM schedule_plans WHERE plan_key=%s", (plan_key,)):
        return jsonify({"message": "该部门在此排班周期已有班表，请打开现有草稿继续编辑"}), 409
    plan_id, _ = execute("INSERT INTO schedule_plans (name,department,start_date,end_date,cycle_id,plan_key,created_by) VALUES (%s,%s,%s,%s,%s,%s,%s)", (name, department, start_date, end_date, cycle["cycle_id"], plan_key, current_user()["id"]))
    return jsonify({"success": True, "plan": query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))}), 201


@api_bp.route("/api/schedule-plans/<int:plan_id>", methods=["GET", "DELETE"])
@schedule_owner_required
def schedule_plan_detail(plan_id):
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None: return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]): return jsonify({"message": "无权访问该部门班表"}), 403
    if request.method == "DELETE":
        if plan["status"] != "草稿": return jsonify({"message": "只能删除草稿；已提交、已打回或已发布班表请按对应流程处理"}), 409
        execute("DELETE FROM schedule_assignments WHERE plan_id=%s", (plan_id,))
        execute("DELETE FROM schedule_plans WHERE plan_id=%s", (plan_id,))
        return jsonify({"success": True, "message": "排班计划已删除"})
    return jsonify({"success": True, "plan": plan, "assignments": get_plan_assignments(plan_id, serialize_dates=True)})


@api_bp.route("/api/schedule-plans/<int:plan_id>/versions", methods=["GET"])
@schedule_owner_required
def schedule_plan_versions(plan_id):
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None:
        return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]):
        return jsonify({"message": "无权查看该部门班表版本"}), 403
    versions = query_all(
        """SELECT v.snapshot_id,v.plan_version,v.action,v.status,v.note,v.created_at,
                  e.name AS actor_name,v.snapshot_json
           FROM schedule_plan_versions v LEFT JOIN employees e ON e.id=v.actor_id
           WHERE v.plan_id=%s ORDER BY v.snapshot_id DESC LIMIT 30""",
        (plan_id,),
    )
    for version in versions:
        if hasattr(version.get("created_at"), "isoformat"):
            version["created_at"] = version["created_at"].isoformat(timespec="minutes")
        if isinstance(version.get("snapshot_json"), str):
            version["snapshot"] = json.loads(version.pop("snapshot_json"))
        else:
            version["snapshot"] = version.pop("snapshot_json")
    return jsonify({"success": True, "versions": versions})


@api_bp.route("/api/schedule-plans/<int:plan_id>/assignments", methods=["POST", "GET"])
@schedule_owner_required
def plan_assignments(plan_id):
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None: return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]): return jsonify({"message": "无权访问该部门班表"}), 403
    if request.method == "GET": return jsonify({"success": True, "assignments": get_plan_assignments(plan_id, serialize_dates=True)})
    if plan["status"] not in ("草稿", "已打回"): return jsonify({"message": "仅草稿或已打回的计划可以修改排班"}), 409
    data = request.get_json(silent=True)
    assignments, error_message = validate_plan_assignments(plan, data.get("assignments") if isinstance(data, dict) else None)
    if error_message: return jsonify({"message": error_message}), 400
    for assignment in assignments:
        execute("INSERT INTO schedule_assignments (plan_id,employee_id,work_date,shift_id) VALUES (%s,%s,%s,%s)", (plan_id, assignment["employee_id"], assignment["work_date"], assignment["shift_id"]))
    return jsonify({"success": True, "assignments": get_plan_assignments(plan_id, serialize_dates=True)}), 201


@api_bp.route("/api/schedule-view", methods=["GET"])
@schedule_owner_required
def get_department_schedule_view():
    """按部门与日期范围返回员工班次，用于部门班表展示。"""
    department = request.args.get("department", "").strip()
    start_date, end_date = request.args.get("start", ""), request.args.get("end", "")
    if (start_date and not parse_iso_date(start_date)) or (end_date and not parse_iso_date(end_date)):
        return jsonify({"message": "日期必须为 YYYY-MM-DD"}), 400
    if start_date and end_date and start_date > end_date:
        return jsonify({"message": "开始日期不能晚于结束日期"}), 400
    conditions, parameters = [], []
    if department:
        conditions.append("e.department=%s")
        parameters.append(department)
    if start_date:
        conditions.append("a.work_date >= %s")
        parameters.append(start_date)
    if end_date:
        conditions.append("a.work_date <= %s")
        parameters.append(end_date)
    sql = """SELECT a.assignment_id, a.work_date, e.id AS employee_id, e.name AS employee_name,
                    e.department, e.position, s.name AS shift_name, s.start_time, s.end_time,
                    p.plan_id, p.name AS plan_name, p.status AS plan_status
             FROM schedule_assignments a
             JOIN employees e ON e.id=a.employee_id
             JOIN shift_types s ON s.shift_id=a.shift_id
             JOIN schedule_plans p ON p.plan_id=a.plan_id"""
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    rows = query_all(sql + " ORDER BY e.department, e.id, a.work_date", parameters)
    for row in rows:
        for field in ("start_time", "end_time"):
            value = row[field]
            if hasattr(value, "total_seconds"):
                seconds = int(value.total_seconds())
                row[field] = f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}"
    return jsonify({"success": True, "schedules": rows})


def schedule_plan_key(department, start_date, end_date):
    """一个部门在一个完整周期只能有一份活动班表；历史重复数据保留为无 key 的只读记录。"""
    return f"{department}|{start_date}|{end_date}"


def replace_schedule_assignments(plan_id, expected_version, assignments):
    """用一次事务替换班次，避免删旧班次后写入失败造成半份草稿。"""
    with connect_database() as connection:
        try:
            connection.begin()
            with connection.cursor() as cursor:
                cursor.execute("""UPDATE schedule_plans SET version=version+1,updated_at=NOW(),
                                  force_submit_reason='',force_submitted_by=NULL,force_submitted_at=NULL
                                  WHERE plan_id=%s AND version=%s""", (plan_id, expected_version))
                if cursor.rowcount != 1:
                    connection.rollback()
                    return False
                cursor.execute("DELETE FROM schedule_assignments WHERE plan_id=%s", (plan_id,))
                if assignments:
                    cursor.executemany(
                        "INSERT INTO schedule_assignments (plan_id,employee_id,work_date,shift_id) VALUES (%s,%s,%s,%s)",
                        [(plan_id, item["employee_id"], item["work_date"], item["shift_id"]) for item in assignments],
                    )
            connection.commit()
            return True
        except Exception:
            connection.rollback()
            raise


def append_schedule_assignments_atomic(plan_id, assignments):
    """原子追加一批班次并推进计划版本，任意一条失败都会整体回滚。"""
    with connect_database() as connection:
        try:
            connection.begin()
            with connection.cursor() as cursor:
                cursor.execute("SELECT status FROM schedule_plans WHERE plan_id=%s FOR UPDATE", (plan_id,))
                plan = cursor.fetchone()
                if plan is None:
                    connection.rollback()
                    return False, "排班计划不存在"
                if plan["status"] not in ("草稿", "已打回"):
                    connection.rollback()
                    return False, "班表状态已变化，请刷新后重试"
                cursor.executemany(
                    "INSERT INTO schedule_assignments (plan_id,employee_id,work_date,shift_id) VALUES (%s,%s,%s,%s)",
                    [(plan_id, item["employee_id"], item["work_date"], item["shift_id"]) for item in assignments],
                )
                cursor.execute("""UPDATE schedule_plans SET version=version+1,updated_at=NOW(),
                                  force_submit_reason='',force_submitted_by=NULL,force_submitted_at=NULL
                                  WHERE plan_id=%s""", (plan_id,))
            connection.commit()
            return True, None
        except pymysql.IntegrityError:
            connection.rollback()
            return False, "批量班次与现有班表重复，请刷新后重试"
        except Exception:
            connection.rollback()
            raise


def find_schedule_board_plan(department, start_date, end_date):
    return query_one(
        """SELECT * FROM schedule_plans WHERE department=%s AND start_date=%s AND end_date=%s
           AND plan_key IS NOT NULL ORDER BY plan_id DESC LIMIT 1""",
        (department, start_date, end_date),
    )


@api_bp.route("/api/schedule-board", methods=["GET", "PUT"])
@schedule_owner_required
def schedule_board():
    """部门班表：前端只面向部门和日期；排班周期在后台自动维护。"""
    user = current_user()
    if request.method == "GET":
        department, start_date, end_date = request.args.get("department", "").strip(), request.args.get("start", ""), request.args.get("end", "")
        if not department or not parse_iso_date(start_date) or not parse_iso_date(end_date) or start_date > end_date:
            return jsonify({"message": "请提供有效的部门、开始日期和结束日期"}), 400
        _, cycle_error = validate_cycle_range(parse_iso_date(start_date), parse_iso_date(end_date))
        if cycle_error: return jsonify({"message": cycle_error}), 400
        if not may_manage_department(user, department):
            return jsonify({"message": "部门经理只能查看本部门班表"}), 403
        employees = query_all(
            """SELECT id, name, employee_no, department, position FROM employees
               WHERE department=%s AND role<>'super_admin' AND employment_status='在职' ORDER BY id""",
            (department,),
        )
        plan = find_schedule_board_plan(department, start_date, end_date)
        assignments = get_plan_assignments(plan["plan_id"], serialize_dates=True) if plan else []
        return jsonify({"success": True, "plan": plan, "employees": employees, "assignments": assignments})

    data = request.get_json(silent=True)
    if not isinstance(data, dict): return jsonify({"message": "请使用 JSON 格式保存班表"}), 400
    department, start_date, end_date = str(data.get("department", "")).strip(), str(data.get("start_date", "")), str(data.get("end_date", ""))
    if not department or not parse_iso_date(start_date) or not parse_iso_date(end_date) or start_date > end_date:
        return jsonify({"message": "请提供有效的部门、开始日期和结束日期"}), 400
    cycle, cycle_error = validate_cycle_range(parse_iso_date(start_date), parse_iso_date(end_date))
    if cycle_error: return jsonify({"message": cycle_error}), 400
    if not may_manage_department(user, department):
        return jsonify({"message": "部门经理只能维护本部门班表"}), 403
    if not query_one("SELECT department_id FROM departments WHERE name=%s", (department,)):
        return jsonify({"message": "部门不存在"}), 404
    plan = find_schedule_board_plan(department, start_date, end_date)
    if plan and plan["status"] not in ("草稿", "已打回"): return jsonify({"message": "班表正在审核或已发布，不能继续修改"}), 409
    if plan is None:
        plan_key = schedule_plan_key(department, start_date, end_date)
        try:
            plan_id, _ = execute("INSERT INTO schedule_plans (name,department,start_date,end_date,cycle_id,plan_key,created_by) VALUES (%s,%s,%s,%s,%s,%s,%s)", (f"{department} {start_date} 至 {end_date} 排班", department, start_date, end_date, cycle["cycle_id"], plan_key, current_user()["id"]))
            plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
        except pymysql.IntegrityError:
            return jsonify({"message": "该班表刚被其他人创建，请重新加载后继续编辑", "conflict": True}), 409
    entries = data.get("assignments", [])
    if not isinstance(entries, list): return jsonify({"message": "assignments 必须是数组"}), 400
    assignments, error_message = validate_plan_assignments(plan, entries, allow_existing=True, allow_consecutive_warning=True) if entries else ([], None)
    if error_message: return jsonify({"message": error_message}), 400
    try: expected_version = int(data.get("version", plan["version"]))
    except (TypeError, ValueError): return jsonify({"message": "班表版本无效，请刷新后重试"}), 400
    if not replace_schedule_assignments(plan["plan_id"], expected_version, assignments):
        return jsonify({"message": "班表已被其他人更新，请重新加载后再保存", "conflict": True}), 409
    record_schedule_snapshot(plan["plan_id"], "保存草稿")
    return jsonify({"success": True, "message": "部门班表草稿已保存", "plan": query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan["plan_id"],)), "assignments": get_plan_assignments(plan["plan_id"], serialize_dates=True)})


@api_bp.route("/api/schedule-plans/<int:plan_id>/assignments/batch", methods=["POST"])
@schedule_owner_required
def create_batch_assignments(plan_id):
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None: return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]): return jsonify({"message": "无权维护该部门班表"}), 403
    if plan["status"] not in ("草稿", "已打回"): return jsonify({"message": "仅草稿或已打回的计划可以修改排班"}), 409
    data = request.get_json(silent=True)
    if not isinstance(data, dict): return jsonify({"message": "请使用 JSON 格式提交批量排班信息"}), 400
    assignments, error_message = build_batch_assignments(plan, data.get("employee_ids"), data.get("shift_id"), data.get("start_date"), data.get("end_date"), data.get("skip_rest_days", True), data.get("skip_weekdays"))
    if error_message: return jsonify({"message": error_message}), 400
    saved, save_error = append_schedule_assignments_atomic(plan_id, assignments)
    if not saved:
        return jsonify({"message": save_error, "conflict": True}), 409
    record_schedule_snapshot(plan_id, "批量排班")
    return jsonify({"success": True, "message": f"已批量添加 {len(assignments)} 条排班", "assignments": get_plan_assignments(plan_id, serialize_dates=True)}), 201


@api_bp.route("/api/schedule-plans/<int:plan_id>/assignments/auto-preview", methods=["POST"])
@schedule_owner_required
def preview_auto_schedule(plan_id):
    """生成自动排班沙盘结果；用户确认并保存前不改变数据库。"""
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None: return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]): return jsonify({"message": "无权维护该部门班表"}), 403
    if plan["status"] not in ("草稿", "已打回"): return jsonify({"message": "仅草稿或已打回计划可以生成自动排班"}), 409
    data = request.get_json(silent=True) or {}
    assignments, error_message = build_auto_schedule_preview(plan, data.get("shift_ids"), data.get("skip_rest_days", True))
    if error_message: return jsonify({"message": error_message}), 400
    return jsonify({"success": True, "count": len(assignments), "assignments": [{**item, "work_date": item["work_date"].isoformat()} for item in assignments]})


@api_bp.route("/api/schedule-plans/<int:plan_id>/assignments/import-template", methods=["GET"])
@schedule_owner_required
def download_schedule_import_template(plan_id):
    plan = query_one("SELECT plan_id,department FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if not plan: return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]): return jsonify({"message": "无权下载该部门模板"}), 403
    # 模板使用真正的 XLSX，避免用户编辑 CSV 时遇到编码和日期格式问题。
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "排班数据"
    sheet.append(["员工工号", "排班日期", "班次名称"])
    sheet.append(["EMP10001", date.today().isoformat(), "早班"])
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        from copy import copy
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    guide = workbook.create_sheet("填写说明")
    guide.append(["项目", "要求"])
    guide.append(["员工工号", "必须填写当前部门已有员工的工号"])
    guide.append(["排班日期", "必须填写当前排班周期内的日期，格式建议为 YYYY-MM-DD"])
    guide.append(["班次名称", "必须填写班次库中已启用的班次名称"])
    guide.append(["注意", "请勿修改第一行表头；空行会自动忽略；上传后请在预览弹窗确认"])
    guide.freeze_panes = "A2"
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 72
    for cell in guide[1]:
        from copy import copy
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    output = BytesIO()
    workbook.save(output)
    return Response(output.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=schedule-import-template.xlsx"})


@api_bp.route("/api/schedule-plans/<int:plan_id>/assignments/export", methods=["GET"])
@schedule_owner_required
def export_schedule_assignments(plan_id):
    """导出当前已保存班表，避免把浏览器中的未保存修改误当成正式数据。"""
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None:
        return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]):
        return jsonify({"message": "无权导出该部门班表"}), 403
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "排班数据"
    sheet.append(["员工工号", "员工姓名", "排班日期", "班次名称", "开始时间", "结束时间"])
    for item in get_plan_assignments(plan_id, serialize_dates=True):
        sheet.append([item["employee_no"], item["employee_name"], item["work_date"], item["shift_name"], item["start_time"], item["end_time"]])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in (("A", 16), ("B", 16), ("C", 16), ("D", 16), ("E", 14), ("F", 14)):
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    filename = f"{plan['department']}-{plan['start_date']}-{plan['end_date']}-排班.xlsx"
    return Response(output.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename, safe='')}"})


@api_bp.route("/api/schedule-templates", methods=["GET", "POST"])
@schedule_owner_required
def schedule_templates():
    """管理可跨设备使用的排班模板；模板按创建人和部门隔离。"""
    user = current_user()
    if request.method == "GET":
        department = request.args.get("department", "").strip()
        if not department or not may_manage_department(user, department):
            return jsonify({"message": "无权查看该部门模板"}), 403
        rows = query_all("""SELECT template_id,name,department,template_json,created_at
                            FROM schedule_templates WHERE created_by=%s AND department=%s
                            ORDER BY updated_at DESC, template_id DESC LIMIT 20""", (user["id"], department))
        templates = []
        for row in rows:
            payload = row["template_json"]
            if isinstance(payload, str): payload = json.loads(payload)
            templates.append({"template_id": row["template_id"], "name": row["name"], "department": row["department"], "assignments": payload, "created_at": row["created_at"].isoformat()[:10] if hasattr(row["created_at"], "isoformat") else str(row["created_at"])[:10]})
        return jsonify({"success": True, "templates": templates})
    data = request.get_json(silent=True) or {}
    department, name, assignments = str(data.get("department", "")).strip(), str(data.get("name", "")).strip()[:100], data.get("assignments")
    if not department or not may_manage_department(user, department): return jsonify({"message": "无权保存该部门模板"}), 403
    if not name or not isinstance(assignments, list) or not assignments or len(assignments) > 2000:
        return jsonify({"message": "模板名称和排班内容不能为空，最多保存 2000 条记录"}), 400
    if any(not isinstance(item, dict) or not item.get("employee_no") or not item.get("shift_name") or not isinstance(item.get("day_index"), int) or item["day_index"] < 0 or item["day_index"] > 366 for item in assignments):
        return jsonify({"message": "模板记录格式无效"}), 400
    template_id, _ = execute("INSERT INTO schedule_templates (department,name,template_json,created_by) VALUES (%s,%s,%s,%s)", (department, name, json.dumps(assignments, ensure_ascii=False), user["id"]))
    return jsonify({"success": True, "template_id": template_id, "message": "排班模板已保存到云端"}), 201


@api_bp.route("/api/schedule-templates/<int:template_id>", methods=["DELETE"])
@schedule_owner_required
def delete_schedule_template(template_id):
    user = current_user()
    template = query_one("SELECT template_id,department FROM schedule_templates WHERE template_id=%s AND created_by=%s", (template_id, user["id"]))
    if template is None: return jsonify({"message": "模板不存在或无权删除"}), 404
    execute("DELETE FROM schedule_templates WHERE template_id=%s AND created_by=%s", (template_id, user["id"]))
    return jsonify({"success": True, "message": "排班模板已删除"})


@api_bp.route("/api/schedule-plans/<int:plan_id>/assignments/import", methods=["POST"])
@schedule_owner_required
def import_schedule_assignments(plan_id):
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None: return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]): return jsonify({"message": "无权维护该部门班表"}), 403
    assignments, message, status_code = import_schedule_file(plan_id, request.files.get("file"))
    if status_code != 201: return jsonify({"message": message}), status_code
    return jsonify({"success": True, "message": message, "assignments": assignments}), status_code


@api_bp.route("/api/schedule-plans/<int:plan_id>/assignments/import-preview", methods=["POST"])
@schedule_owner_required
def preview_schedule_assignments(plan_id):
    """上传文件后只返回识别结果，用户确认前不改变排班数据。"""
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None:
        return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]):
        return jsonify({"message": "无权维护该部门班表"}), 403
    assignments, message, status_code = prepare_schedule_import(plan_id, request.files.get("file"))
    if status_code != 200:
        return jsonify({"message": message}), status_code
    preview_assignments = [{**item, "work_date": item["work_date"].isoformat() if hasattr(item["work_date"], "isoformat") else item["work_date"]} for item in assignments]
    return jsonify({
        "success": True,
        "message": message,
        "count": len(assignments),
        "assignments": preview_assignments,
        "preview_rows": [{key: item[key] for key in ("employee_no", "employee_name", "work_date", "shift_name")} for item in preview_assignments],
    })


@api_bp.route("/api/schedule-plans/<int:plan_id>/assignments/import-confirm", methods=["POST"])
@schedule_owner_required
def confirm_schedule_assignments(plan_id):
    """再次校验预览数据并原子写入，防止预览期间数据变化。"""
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None:
        return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]):
        return jsonify({"message": "无权维护该部门班表"}), 403
    data = request.get_json(silent=True) or {}
    raw_assignments = data.get("assignments")
    assignments, error_message = validate_plan_assignments(plan, raw_assignments)
    if error_message:
        return jsonify({"message": f"确认导入校验失败：{error_message}"}), 400
    saved, save_error = append_schedule_assignments_atomic(plan_id, assignments)
    if not saved:
        return jsonify({"message": save_error, "conflict": True}), 409
    record_schedule_snapshot(plan_id, "导入排班")
    return jsonify({"success": True, "message": f"成功导入 {len(assignments)} 条排班", "plan": query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,)), "assignments": get_plan_assignments(plan_id, serialize_dates=True)}), 201


@api_bp.route("/api/schedule-plans/<int:plan_id>/publish", methods=["POST"])
@hr_required
def publish_schedule_plan(plan_id):
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None: return jsonify({"message": "排班计划不存在"}), 404
    if plan["status"] not in ("草稿", "待人事审核"): return jsonify({"message": "仅草稿或待人事审核的班表可以发布"}), 409
    assignments = get_plan_assignments(plan_id)
    if not assignments: return jsonify({"message": "计划没有排班记录，不能发布"}), 400
    data = {"assignments": [{
        "employee_id": item["employee_id"], "shift_id": item["shift_id"],
        "work_date": item["work_date"].isoformat(), "is_on_leave": item.get("is_on_leave", False),
    } for item in assignments]}
    _, error_message = validate_plan_assignments(plan, data["assignments"], allow_existing=True, allow_consecutive_warning=True)
    if error_message: return jsonify({"message": f"发布校验失败：{error_message}"}), 400
    policy_issues = collect_schedule_policy_issues(plan, data["assignments"])
    if policy_issues and not plan.get("force_submit_reason"):
        return jsonify({"message": f"发布校验失败：{policy_issues[0]}；请由部门经理说明原因后强制提交人事复核"}), 400
    _, changed = execute("""UPDATE schedule_plans SET status='已发布',published_at=NOW(),hr_id=%s,version=version+1
                            WHERE plan_id=%s AND status IN ('草稿','待人事审核')""", (current_user()["id"], plan_id))
    if changed != 1:
        return jsonify({"message": "班表状态已被其他审核人修改，请刷新后重试", "conflict": True}), 409
    record_schedule_snapshot(plan_id, "审核发布", plan.get("force_submit_reason") or "")
    event_version = int(plan.get("version") or 1) + 1
    assigned_employee_ids = [
        item["employee_id"] for item in query_all(
            """SELECT DISTINCT a.employee_id FROM schedule_assignments a
               JOIN employees e ON e.id=a.employee_id
               WHERE a.plan_id=%s AND e.employment_status<>'已离职' AND e.role<>'department_admin'""",
            (plan_id,),
        )
    ]
    notify_users_safely(
        assigned_employee_ids, "schedule_published", "新的班表已发布",
        f"{plan['department']} {plan['start_date']} 至 {plan['end_date']} 的班表已发布，请查看自己的班次。",
        "/pages/my-schedule.html", related_type="schedule_plan", related_id=plan_id,
        dedupe_key=f"schedule-published-personal:{plan_id}:{event_version}",
    )
    notify_users_safely(
        notification_recipient_ids("department_admin", department=plan["department"]),
        "schedule_published", "本部门班表已发布",
        f"{plan['department']} {plan['start_date']} 至 {plan['end_date']} 的班表已由人事审核发布。",
        "/pages/schedule.html", related_type="schedule_plan", related_id=plan_id,
        dedupe_key=f"schedule-published-management:{plan_id}:{event_version}",
    )
    return jsonify({"success": True, "message": "人事已审核并发布班表，员工日历已同步", "plan": query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))})


@api_bp.route("/api/schedule-plans/<int:plan_id>/submit", methods=["POST"])
@department_admin_required
def submit_schedule_plan(plan_id):
    user = current_user()
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None: return jsonify({"message": "排班计划不存在"}), 404
    if user["role"] != "department_admin" or plan["department"] != user["department"]:
        return jsonify({"message": "部门经理只能提交本部门班表"}), 403
    if plan["status"] not in ("草稿", "已打回"): return jsonify({"message": "仅草稿或已打回班表可以提交人事审核"}), 409
    data = request.get_json(silent=True) or {}
    force_submit, force_reason = data.get("force") is True, str(data.get("reason", "")).strip()[:500]
    raw_assignments = get_plan_assignments(plan_id)
    assignments, error_message = validate_plan_assignments(plan, raw_assignments, allow_existing=True, allow_consecutive_warning=True)
    if error_message:
        return jsonify({"message": f"提交校验失败：{error_message}", "force_available": False}), 400
    policy_issues = collect_schedule_policy_issues(plan, assignments)
    forceable = bool(policy_issues)
    if policy_issues and not (force_submit and len(force_reason) >= 10):
        return jsonify({"message": f"提交校验失败：{policy_issues[0]}", "issues": policy_issues[:10], "force_available": True}), 400
    if force_submit and not forceable:
        return jsonify({"message": "当前班表没有需要例外复核的软规则问题，无需强制提交"}), 400
    if force_submit and len(force_reason) < 10:
        return jsonify({"message": "强制提交请填写至少 10 个字的业务原因，供人事复核"}), 400
    _, changed = execute("""UPDATE schedule_plans SET status='待人事审核',submitted_at=NOW(),hr_note='',
               force_submit_reason=%s,force_submitted_by=%s,
               force_submitted_at=IF(%s, NOW(), NULL),version=version+1
               WHERE plan_id=%s AND status IN ('草稿','已打回')""",
            (force_reason if force_submit else "", user["id"] if force_submit else None, force_submit, plan_id))
    if changed != 1:
        return jsonify({"message": "班表状态已被其他人修改，请刷新后重试", "conflict": True}), 409
    record_schedule_snapshot(plan_id, "强制提交" if force_submit else "提交审核", force_reason)
    event_version = int(plan.get("version") or 1) + 1
    notify_users_safely(
        notification_recipient_ids("hr_admin", "super_admin", exclude_id=user["id"]),
        "schedule_review", "有新的班表等待审核",
        f"{plan['department']}提交了 {plan['start_date']} 至 {plan['end_date']} 的班表"
        f"{'，并申请例外复核' if force_submit else ''}。",
        "/pages/schedule.html", related_type="schedule_plan", related_id=plan_id,
        dedupe_key=f"schedule-submitted:{plan_id}:{event_version}",
    )
    return jsonify({"success": True, "message": "班表已强制提交人事复核" if force_submit else "班表已提交人事审核", "plan": query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))})


@api_bp.route("/api/schedule-plans/<int:plan_id>/revoke-submission", methods=["POST"])
@department_admin_required
def revoke_schedule_submission(plan_id):
    user = current_user()
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None: return jsonify({"message": "排班计划不存在"}), 404
    if plan["department"] != user["department"]: return jsonify({"message": "部门经理只能撤销本部门班表"}), 403
    if plan["status"] != "待人事审核": return jsonify({"message": "仅待人事审核的班表可以撤销提交"}), 409
    _, changed = execute("UPDATE schedule_plans SET status='草稿',submitted_at=NULL,version=version+1 WHERE plan_id=%s AND status='待人事审核'", (plan_id,))
    if changed != 1: return jsonify({"message": "班表状态已变化，请刷新后重试", "conflict": True}), 409
    record_schedule_snapshot(plan_id, "撤销提交")
    return jsonify({"success": True, "message": "已撤销提交，班表已恢复为草稿"})


@api_bp.route("/api/schedule-plans/<int:plan_id>/return", methods=["POST"])
@hr_required
def return_schedule_plan(plan_id):
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None: return jsonify({"message": "排班计划不存在"}), 404
    if plan["status"] != "待人事审核": return jsonify({"message": "仅待人事审核的班表可以打回"}), 409
    data = request.get_json(silent=True) or {}
    note = str(data.get("note", "")).strip()
    if not note: return jsonify({"message": "请填写打回原因，方便部门经理修改"}), 400
    _, changed = execute("UPDATE schedule_plans SET status='已打回',hr_id=%s,hr_note=%s,version=version+1 WHERE plan_id=%s AND status='待人事审核'", (current_user()["id"], note, plan_id))
    if changed != 1: return jsonify({"message": "班表已被其他审核人处理，请刷新后重试", "conflict": True}), 409
    record_schedule_snapshot(plan_id, "打回修改", note)
    notify_users_safely(
        notification_recipient_ids("department_admin", department=plan["department"]),
        "schedule_returned", "班表已被打回修改",
        f"{plan['department']} {plan['start_date']} 至 {plan['end_date']} 的班表已被打回。原因：{note}",
        "/pages/schedule.html", related_type="schedule_plan", related_id=plan_id,
        dedupe_key=f"schedule-returned:{plan_id}:{int(plan.get('version') or 1) + 1}",
    )
    return jsonify({"success": True, "message": "班表已打回部门经理修改"})


@api_bp.route("/api/schedule-plans/<int:plan_id>/withdraw", methods=["POST"])
@schedule_owner_required
def withdraw_schedule_plan(plan_id):
    plan = query_one("SELECT * FROM schedule_plans WHERE plan_id=%s", (plan_id,))
    if plan is None: return jsonify({"message": "排班计划不存在"}), 404
    if not may_manage_department(current_user(), plan["department"]): return jsonify({"message": "无权撤回该部门班表"}), 403
    if plan["status"] != "已发布": return jsonify({"message": "仅已发布计划可以撤回"}), 409
    _, changed = execute("UPDATE schedule_plans SET status='草稿',published_at=NULL,version=version+1 WHERE plan_id=%s AND status='已发布'", (plan_id,))
    if changed != 1: return jsonify({"message": "班表状态已变化，请刷新后重试", "conflict": True}), 409
    record_schedule_snapshot(plan_id, "撤回发布")
    affected_ids = [
        item["employee_id"] for item in query_all(
            "SELECT DISTINCT employee_id FROM schedule_assignments WHERE plan_id=%s", (plan_id,),
        )
    ]
    notify_users_safely(
        affected_ids, "schedule_withdrawn", "已发布班表被撤回",
        f"{plan['department']} {plan['start_date']} 至 {plan['end_date']} 的班表已撤回，原班次暂不再生效。",
        "/pages/my-schedule.html", related_type="schedule_plan", related_id=plan_id,
        dedupe_key=f"schedule-withdrawn:{plan_id}:{int(plan.get('version') or 1) + 1}",
    )
    return jsonify({"success": True, "message": "排班计划已撤回"})


@api_bp.route("/api/onboarding", methods=["POST", "OPTIONS"])
@api_operation_log("发起入职申请")
@hr_required
@validate_json_fields("employee_id", "department", "position")
def create_onboarding_record():
    if request.method == "OPTIONS": return "", 204
    data = request.get_json(silent=True)
    if isinstance(data, dict) and data.get("role") in {"super_admin", "hr_admin", "department_admin"} and current_user()["role"] != "super_admin":
        return jsonify({"message": "无权限：仅超级管理员可以授予管理角色"}), 403
    record, error_message, status_code = create_onboarding(data)
    if error_message: return jsonify({"message": error_message}), status_code
    return jsonify({"success": True, "record": record}), status_code


@api_bp.route("/api/onboarding/queue", methods=["GET"])
@hr_required
def get_onboarding_queue():
    onboarding_type = request.args.get("type", "").strip()
    sql, parameters = "SELECT * FROM onboarding_records WHERE status = '待办理'", []
    if onboarding_type:
        sql += " AND onboarding_type = %s"; parameters.append(onboarding_type)
    return jsonify({"success": True, "queue": query_all(sql + " ORDER BY FIELD(priority, '紧急', '普通'), queue_number", parameters)})


@api_bp.route("/api/onboarding/handle", methods=["POST", "OPTIONS"])
@hr_required
def handle_next_onboarding():
    if request.method == "OPTIONS": return "", 204
    data = request.get_json(silent=True)
    onboarding_type = str(data.get("type", "")).strip() if isinstance(data, dict) else ""
    if not onboarding_type: return jsonify({"message": "请提交办理类别 type"}), 400
    record = query_one("""SELECT * FROM onboarding_records WHERE status = '待办理' AND onboarding_type = %s ORDER BY FIELD(priority, '紧急', '普通'), queue_number LIMIT 1""", (onboarding_type,))
    if record is None: return jsonify({"success": False, "message": "暂无待办理员工"})
    return finish_onboarding_record(record)


def finish_onboarding_record(record):
    """由人事办结一条已获批准的入职记录。"""
    execute("UPDATE onboarding_records SET status = '已办理' WHERE record_id = %s", (record["record_id"],))
    execute("UPDATE employees SET employment_status='在职' WHERE id=%s", (record["employee_id"],))
    record["status"] = "已办理"
    return jsonify({"success": True, "message": f"已为 {record['employee_name']} 办理{record['onboarding_type']}", "record": record})


@api_bp.route("/api/onboarding/queue/<int:record_id>/handle", methods=["POST"])
@hr_required
def handle_onboarding_record(record_id):
    record = query_one("SELECT * FROM onboarding_records WHERE record_id=%s AND status='待办理'", (record_id,))
    if record is None: return jsonify({"message": "该入职事项不存在、尚未获批或已办理"}), 404
    return finish_onboarding_record(record)


@api_bp.route("/api/employee-records", methods=["POST", "OPTIONS"])
@hr_required
def create_employee_record():
    if request.method == "OPTIONS": return "", 204
    record, error_message, status_code = create_employee_archive(request.get_json(silent=True))
    if error_message: return jsonify({"message": error_message}), status_code
    return jsonify({"success": True, "record": record}), status_code


@api_bp.route("/api/employee-records", methods=["GET"])
@hr_required
def get_employee_records():
    employee_id = request.args.get("employee_id", "").strip()
    try: records = query_all("SELECT * FROM employee_records WHERE employee_id = %s ORDER BY record_id", (int(employee_id),)) if employee_id else query_all("SELECT * FROM employee_records ORDER BY record_id")
    except ValueError: return jsonify({"message": "employee_id 必须是有效编号"}), 400
    return jsonify({"success": True, "total": len(records), "records": [serialize_record(record) for record in records]})


@api_bp.route("/api/employee-records/<int:record_id>", methods=["GET"])
@hr_required
def get_employee_record(record_id):
    record = query_one("SELECT * FROM employee_records WHERE record_id = %s", (record_id,))
    if record is None: return jsonify({"message": "档案不存在"}), 404
    return jsonify({"success": True, "record": serialize_record(record)})


@api_bp.route("/api/employee-records/<int:record_id>", methods=["PUT", "OPTIONS"])
@hr_required
def update_employee_record(record_id):
    if request.method == "OPTIONS": return "", 204
    record, error_message, status_code = update_employee_archive(record_id, request.get_json(silent=True))
    if error_message: return jsonify({"message": error_message}), status_code
    return jsonify({"success": True, "record": record}), status_code


@api_bp.route("/api/employee-records/<int:record_id>", methods=["DELETE", "OPTIONS"])
@hr_required
def delete_employee_record(record_id):
    if request.method == "OPTIONS": return "", 204
    _, affected_rows = execute("DELETE FROM employee_records WHERE record_id = %s", (record_id,))
    if not affected_rows: return jsonify({"message": "档案不存在"}), 404
    return jsonify({"success": True, "message": "档案已删除"})


@api_bp.route("/api/my/tasks", methods=["GET"])
@api_operation_log("读取个人待办")
@login_required
def get_my_tasks():
    user = current_user()
    tasks = query_all("SELECT *, CASE WHEN status='待处理' AND due_date IS NOT NULL AND due_date<CURDATE() THEN '已逾期' ELSE status END AS display_status FROM employee_tasks WHERE employee_id=%s ORDER BY task_id DESC", (user["id"],))
    return jsonify({"success": True, "tasks": tasks})


@api_bp.route("/api/my/tasks/<int:task_id>/submit", methods=["POST"])
@login_required
def submit_my_task(task_id):
    user = current_user(); data = request.get_json(silent=True) or {}
    task = query_one("SELECT * FROM employee_tasks WHERE task_id=%s AND employee_id=%s", (task_id, user["id"]))
    if task is None: return jsonify({"message": "待办不存在"}), 404
    if task["status"] != "待处理": return jsonify({"message": "该待办已提交"}), 400
    content = str(data.get("content", "员工已提交处理结果")).strip()
    if task["task_type"] in ("offboarding_handover", "handover_request"):
        _, changed = execute("UPDATE employee_tasks SET status='待人事审核', content=CONCAT(content, '\n员工交接说明：', %s), submitted_at=NOW() WHERE task_id=%s AND employee_id=%s AND status='待处理'", (content, task_id, user["id"]))
        if changed != 1: return jsonify({"message": "该待办刚刚已被提交，请刷新页面"}), 409
        return jsonify({"success": True, "message": "交接已提交人事审核"})
    if task["task_type"] == "manual":
        _, changed = execute("UPDATE employee_tasks SET status='已完成', content=CONCAT(content, '\n员工处理结果：', %s), submitted_at=NOW(), completed_at=NOW() WHERE task_id=%s AND employee_id=%s AND status='待处理'", (content, task_id, user["id"]))
        if changed != 1: return jsonify({"message": "该待办刚刚已被提交，请刷新页面"}), 409
        if task.get("publisher_id") and task["publisher_id"] != user["id"]:
            notify_users_safely([task["publisher_id"]], "task_completed", "员工已完成待办", f"{user['name']} 已提交：{task['title']}", "/pages/task-center.html", related_type="employee_task", related_id=task_id, dedupe_key=f"manual-task-completed:{task_id}")
        return jsonify({"success": True, "message": "待办已完成并提交结果"})
    _, changed = execute("UPDATE employee_tasks SET status='已提交', content=CONCAT(content, '\n员工反馈：', %s), submitted_at=NOW() WHERE task_id=%s AND employee_id=%s AND status='待处理'", (content, task_id, user["id"]))
    if changed != 1: return jsonify({"message": "该待办刚刚已被提交，请刷新页面"}), 409
    employee = find_employee(user["id"])
    execute("INSERT INTO admin_todos (task_id, employee_id, employee_name, title, content) VALUES (%s,%s,%s,%s,%s)", (task_id, user["id"], employee["name"], task["title"], content))
    return jsonify({"success": True, "message": "已提交给超级管理员确认"})


@api_bp.route("/api/offboarding-requests", methods=["POST"])
@login_required
def create_offboarding_request():
    user, data = current_user(), request.get_json(silent=True) or {}
    reason = str(data.get("reason", "")).strip()
    if not reason: return jsonify({"message": "请填写离职原因"}), 400
    task_id, error_message, status_code = _assistant_create_offboarding_request(user, reason)
    if error_message: return jsonify({"message": error_message}), status_code
    return jsonify({"success": True, "message": "离职申请已提交给部门经理审批"}), 201


@api_bp.route("/api/department/offboarding-requests", methods=["GET"])
@department_admin_required
def department_offboarding_requests():
    user = current_user()
    if user["role"] != "department_admin": return jsonify({"message": "仅部门经理可以审批员工离职"}), 403
    tasks = query_all("""SELECT t.*, e.name AS employee_name, e.employee_no FROM employee_tasks t
        JOIN employees e ON e.id=t.employee_id WHERE e.department=%s AND e.id<>%s
        AND t.task_type='offboarding_request' AND t.status='待部门审批' ORDER BY t.created_at""", (user["department"], user["id"]))
    return jsonify({"success": True, "requests": tasks})


@api_bp.route("/api/department/offboarding-requests/<int:task_id>/review", methods=["POST"])
@department_admin_required
def review_department_offboarding_request(task_id):
    user, data = current_user(), request.get_json(silent=True) or {}
    if user["role"] != "department_admin": return jsonify({"message": "仅部门经理可以审批员工离职"}), 403
    decision, note = data.get("decision"), str(data.get("note", "")).strip()
    if decision not in ("同意", "驳回"): return jsonify({"message": "decision 必须为同意或驳回"}), 400
    task = query_one("""SELECT t.*, e.name AS employee_name FROM employee_tasks t JOIN employees e ON e.id=t.employee_id
        WHERE t.task_id=%s AND e.department=%s AND e.id<>%s AND t.task_type='offboarding_request' AND t.status='待部门审批'""", (task_id, user["department"], user["id"]))
    if task is None: return jsonify({"message": "离职申请不存在或已处理"}), 404
    if decision == "驳回":
        execute("UPDATE employee_tasks SET status='部门已驳回',content=CONCAT(content,'\n经理意见：',%s) WHERE task_id=%s", (note, task_id))
        execute("UPDATE employees SET employment_status='在职' WHERE id=%s", (task["employee_id"],))
        return jsonify({"success": True, "message": "已驳回离职申请"})
    execute("UPDATE employee_tasks SET status='待老板批准',content=CONCAT(content,'\n经理意见：',%s) WHERE task_id=%s", (note, task_id))
    execute("INSERT INTO admin_todos (task_id,employee_id,employee_name,title,content) VALUES (%s,%s,%s,%s,%s)", (task_id, task["employee_id"], task["employee_name"], "离职申请审批", task["content"]))
    return jsonify({"success": True, "message": "已提交超级管理员最终批准"})


@api_bp.route("/api/hr/offboarding-handovers", methods=["GET"])
@hr_required
def hr_offboarding_handovers():
    tasks = query_all("""SELECT t.*, e.name AS employee_name, e.employee_no, e.department FROM employee_tasks t
        JOIN employees e ON e.id=t.employee_id WHERE t.task_type IN ('offboarding_handover','handover_request') AND t.status='待人事审核' ORDER BY t.submitted_at""")
    return jsonify({"success": True, "handovers": tasks})


@api_bp.route("/api/hr/offboarding-handovers/<int:task_id>/review", methods=["POST"])
@hr_required
def review_hr_offboarding_handover(task_id):
    data = request.get_json(silent=True) or {}; decision, note = data.get("decision"), str(data.get("note", "")).strip()
    if decision not in ("同意", "驳回"): return jsonify({"message": "decision 必须为同意或驳回"}), 400
    task = query_one("SELECT * FROM employee_tasks WHERE task_id=%s AND task_type IN ('offboarding_handover','handover_request') AND status='待人事审核'", (task_id,))
    if task is None: return jsonify({"message": "交接不存在或已处理"}), 404
    if decision == "同意":
        execute("UPDATE employee_tasks SET status='已完成',completed_at=NOW(),content=CONCAT(content,'\n人事意见：',%s) WHERE task_id=%s", (note, task_id))
        if task["task_type"] == "offboarding_handover":
            execute("UPDATE employees SET employment_status='已离职' WHERE id=%s", (task["employee_id"],))
            return jsonify({"success": True, "message": "人事已审核通过，员工已正式离职"})
        return jsonify({"success": True, "message": "工作交接申请已审核通过"})
    execute("UPDATE employee_tasks SET status='待处理',content=CONCAT(content,'\n人事退回：',%s) WHERE task_id=%s", (note, task_id))
    return jsonify({"success": True, "message": "已退回员工补充交接"})


@api_bp.route("/api/admin/todos", methods=["GET"])
@super_admin_required
def get_admin_todos():
    return jsonify({"success": True, "todos": query_all("SELECT * FROM admin_todos WHERE status='待确认' ORDER BY todo_id DESC")})


@api_bp.route("/api/admin/todos/<int:todo_id>/review", methods=["POST"])
@super_admin_required
def review_admin_todo(todo_id):
    data = request.get_json(silent=True) or {}; decision = data.get("decision")
    if decision not in ("确认", "驳回"): return jsonify({"message": "decision 必须为确认或驳回"}), 400
    todo = query_one("SELECT * FROM admin_todos WHERE todo_id=%s AND status='待确认'", (todo_id,))
    if todo is None: return jsonify({"message": "待办不存在或已处理"}), 404
    task = query_one("SELECT * FROM employee_tasks WHERE task_id=%s", (todo["task_id"],))
    execute("UPDATE admin_todos SET status=%s, reviewed_at=NOW() WHERE todo_id=%s", (decision, todo_id))
    execute("UPDATE employee_tasks SET status=%s, completed_at=NOW() WHERE task_id=%s", ("已完成" if decision == "确认" else "待处理", task["task_id"]))
    if decision == "确认" and task["task_type"] == "onboarding":
        execute("UPDATE employees SET employment_status='入职办理中' WHERE id=%s", (todo["employee_id"],))
        execute("UPDATE onboarding_records SET status='待办理' WHERE employee_id=%s AND status='待超管批准'", (todo["employee_id"],))
    if decision == "驳回" and task["task_type"] == "onboarding":
        execute("UPDATE employees SET employment_status='待分配' WHERE id=%s", (todo["employee_id"],))
        execute("UPDATE onboarding_records SET status='已驳回' WHERE employee_id=%s AND status='待超管批准'", (todo["employee_id"],))
    if decision == "确认" and task["task_type"] == "offboarding_request":
        employee = find_employee(todo["employee_id"])
        create_employee_task(employee, "offboarding_handover", "请完成离职交接", "老板已批准离职申请。请完成工作、资产和资料交接后提交人事审核。")
        execute("UPDATE employees SET employment_status='离职交接中' WHERE id=%s", (todo["employee_id"],))
    if decision == "驳回" and task["task_type"] == "offboarding_request":
        execute("UPDATE employee_tasks SET status='老板已驳回' WHERE task_id=%s", (task["task_id"],))
        execute("UPDATE employees SET employment_status='在职' WHERE id=%s", (todo["employee_id"],))
    if decision == "确认" and task["task_type"] == "offboarding": execute("UPDATE employees SET employment_status='已离职' WHERE id=%s", (todo["employee_id"],))
    if decision == "驳回" and task["task_type"] == "offboarding": execute("UPDATE employees SET employment_status='在职' WHERE id=%s", (todo["employee_id"],))
    if decision == "确认" and task["task_type"] == "handover": execute("UPDATE employees SET employment_status='在职' WHERE id=%s", (todo["employee_id"],))
    return jsonify({"success": True, "message": f"已{decision}员工事项"})


# ==================== 员工班次日历与两级审批 ====================

def serialize_schedule_request(record):
    """为申请记录补充展示所需的班次与审批人信息。"""
    if record is None:
        return None
    result = record.copy()
    if hasattr(result.get("work_date"), "isoformat"):
        result["work_date"] = result["work_date"].isoformat()
    result["request_type_name"] = {"shift_change": "换班", "leave": "请假"}.get(result["request_type"], result["request_type"])
    return result


def serialize_schedule_plan(plan):
    """统一输出 ISO 日期，供日期输入框和审核队列直接使用。"""
    if plan is None:
        return None
    result = plan.copy()
    for field in ("start_date", "end_date"):
        if hasattr(result.get(field), "isoformat"):
            result[field] = result[field].isoformat()
    return result


def finalize_schedule_request(request_id, expected_status, decision, reviewer_id, reviewer_kind, note=""):
    """在一个事务中锁定申请和原班次；换班终审时重新计算关键工时约束。"""
    status = "已批准" if decision == "同意" else ("部门已驳回" if reviewer_kind == "manager" else "超管已驳回")
    with connect_database() as connection:
        try:
            connection.begin()
            with connection.cursor() as cursor:
                cursor.execute(
                    "SELECT * FROM schedule_requests WHERE request_id=%s AND status=%s FOR UPDATE",
                    (request_id, expected_status),
                )
                record = cursor.fetchone()
                if record is None:
                    connection.rollback()
                    return None, None, "申请已被其他审批人处理，请刷新后重试"
                plan_id = None
                if decision == "同意":
                    cursor.execute(
                        """SELECT a.assignment_id,a.plan_id,a.shift_id,p.version,p.status
                           FROM schedule_assignments a JOIN schedule_plans p ON p.plan_id=a.plan_id
                           WHERE a.employee_id=%s AND a.work_date=%s AND p.status='已发布' FOR UPDATE""",
                        (record["employee_id"], record["work_date"]),
                    )
                    assignment = cursor.fetchone()
                    if assignment is None:
                        connection.rollback()
                        return None, None, "原班表已撤回或班次已不存在，申请已失效"
                    if record.get("assignment_id") and assignment["assignment_id"] != record["assignment_id"]:
                        connection.rollback()
                        return None, None, "原班次已被替换，请员工基于最新班表重新申请"
                    if record.get("original_shift_id") and assignment["shift_id"] != record["original_shift_id"]:
                        connection.rollback()
                        return None, None, "原班次已经变化，请员工基于最新班表重新申请"
                    plan_id = assignment["plan_id"]
                    if record["request_type"] == "shift_change":
                        if assignment["shift_id"] == record["desired_shift_id"]:
                            connection.rollback()
                            return None, None, "目标班次与当前班次相同，无需调整"
                        cursor.execute("SELECT * FROM shift_types WHERE shift_id=%s AND is_active=1", (record["desired_shift_id"],))
                        desired_shift = cursor.fetchone()
                        if desired_shift is None:
                            connection.rollback()
                            return None, None, "目标班次已停用，请重新选择"
                        context_start = record["work_date"] - timedelta(days=40)
                        context_end = record["work_date"] + timedelta(days=40)
                        cursor.execute(
                            """SELECT a.assignment_id,a.work_date,s.name AS shift_name,s.start_time,s.end_time,s.duration_hours,
                                      CASE WHEN a.assignment_id=%s THEN 1 ELSE 0 END AS _candidate
                               FROM schedule_assignments a JOIN schedule_plans p ON p.plan_id=a.plan_id
                               JOIN shift_types s ON s.shift_id=a.shift_id
                               LEFT JOIN (SELECT employee_id,work_date FROM schedule_requests
                                          WHERE request_type='leave' AND status='已批准' GROUP BY employee_id,work_date) lr
                                 ON lr.employee_id=a.employee_id AND lr.work_date=a.work_date
                               WHERE a.employee_id=%s AND p.status='已发布' AND lr.employee_id IS NULL
                                 AND a.work_date BETWEEN %s AND %s FOR UPDATE""",
                            (assignment["assignment_id"], record["employee_id"], context_start, context_end),
                        )
                        work_rows = cursor.fetchall()
                        for row in work_rows:
                            row["employee_id"] = record["employee_id"]
                            if row["assignment_id"] == assignment["assignment_id"]:
                                row.update(
                                    shift_name=desired_shift["name"], start_time=desired_shift["start_time"],
                                    end_time=desired_shift["end_time"], duration_hours=desired_shift["duration_hours"],
                                )
                        cursor.execute("SELECT * FROM scheduling_rules WHERE rule_name='默认排班规则'")
                        rule = cursor.fetchone()
                        violation = find_minimum_rest_violation(work_rows, float(rule.get("min_rest_hours") or 0), require_candidate=True)
                        if violation:
                            connection.rollback()
                            return None, None, "换班后与相邻班次重叠或班间休息不足"
                        work_dates = {row["work_date"] for row in work_rows}
                        consecutive = 0
                        for work_date in sorted(work_dates):
                            consecutive = consecutive + 1 if work_date - timedelta(days=1) in work_dates else 1
                            if consecutive > int(rule["max_consecutive_days"]):
                                connection.rollback()
                                return None, None, f"换班后将连续工作超过 {rule['max_consecutive_days']} 天"
                        target_overtime = max(float(desired_shift["duration_hours"] or 0) - 8, 0)
                        if target_overtime > float(rule["max_daily_overtime_hours"]):
                            connection.rollback()
                            return None, None, "目标班次会造成单日加班超限"
                        month = record["work_date"].replace(day=1)
                        next_month = add_months(month, 1)
                        monthly_overtime = sum(
                            max(float(row["duration_hours"] or 0) - 8, 0)
                            for row in work_rows if month <= row["work_date"] < next_month
                        )
                        if monthly_overtime > float(rule["max_monthly_overtime_hours"]):
                            connection.rollback()
                            return None, None, "换班后会造成当月累计加班超限"
                        cursor.execute(
                            "UPDATE schedule_assignments SET shift_id=%s WHERE assignment_id=%s",
                            (record["desired_shift_id"], assignment["assignment_id"]),
                        )
                    cursor.execute("UPDATE schedule_plans SET version=version+1,updated_at=NOW() WHERE plan_id=%s", (plan_id,))
                review_columns = (
                    "manager_id=%s,manager_note=%s,manager_reviewed_at=NOW()"
                    if reviewer_kind == "manager"
                    else "admin_id=%s,admin_note=%s,admin_reviewed_at=NOW()"
                )
                cursor.execute(
                    f"UPDATE schedule_requests SET status=%s,{review_columns} WHERE request_id=%s AND status=%s",
                    (status, reviewer_id, note, request_id, expected_status),
                )
                if cursor.rowcount != 1:
                    connection.rollback()
                    return None, None, "申请状态已变化，请刷新后重试"
            connection.commit()
            record["status"] = status
            return record, plan_id, None
        except Exception:
            connection.rollback()
            raise


@api_bp.route("/api/my-schedule", methods=["GET"])
@api_operation_log("读取个人班次")
@login_required
def get_my_schedule():
    user = current_user()
    start, end = request.args.get("start", ""), request.args.get("end", "")
    if not parse_iso_date(start) or not parse_iso_date(end) or start > end:
        return jsonify({"message": "请提供有效的开始和结束日期"}), 400
    schedules = query_all(
        """SELECT a.assignment_id, a.work_date, s.shift_id, s.name AS shift_name, s.start_time, s.end_time, s.color
           FROM schedule_assignments a JOIN schedule_plans p ON p.plan_id=a.plan_id
           JOIN shift_types s ON s.shift_id=a.shift_id
           WHERE a.employee_id=%s AND p.status='已发布' AND a.work_date BETWEEN %s AND %s ORDER BY a.work_date""",
        (user["id"], start, end),
    )
    leave_days = query_all(
        """SELECT work_date, request_id FROM schedule_requests
           WHERE employee_id=%s AND request_type='leave' AND status='已批准' AND work_date BETWEEN %s AND %s""",
        (user["id"], start, end),
    )
    for schedule in schedules:
        if hasattr(schedule.get("work_date"), "isoformat"):
            schedule["work_date"] = schedule["work_date"].isoformat()
        for field in ("start_time", "end_time"):
            value = schedule[field]
            if hasattr(value, "total_seconds"):
                seconds = int(value.total_seconds())
                schedule[field] = f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}"
    for leave_day in leave_days:
        if hasattr(leave_day.get("work_date"), "isoformat"):
            leave_day["work_date"] = leave_day["work_date"].isoformat()
    return jsonify({"success": True, "schedules": schedules, "leave_days": leave_days})


@api_bp.route("/api/schedule-requests", methods=["POST"])
@api_operation_log("提交班次申请")
@login_required
@validate_json_fields("request_type", "work_date", "reason")
def create_schedule_request():
    user, data = current_user(), request.get_json(silent=True)
    if user["role"] == "super_admin":
        return jsonify({"message": "超级管理员无需提交个人排班申请"}), 403
    if not isinstance(data, dict):
        return jsonify({"message": "请使用 JSON 格式提交申请"}), 400
    request_type, work_date = str(data.get("request_type", "")).strip(), str(data.get("work_date", "")).strip()
    reason = str(data.get("reason", "")).strip()
    if request_type not in ("shift_change", "leave") or not parse_iso_date(work_date) or not reason:
        return jsonify({"message": "请填写申请类型、日期和原因"}), 400
    assignment = query_one(
        """SELECT a.assignment_id,a.shift_id,p.version AS assignment_version FROM schedule_assignments a JOIN schedule_plans p ON p.plan_id=a.plan_id
           WHERE a.employee_id=%s AND a.work_date=%s AND p.status='已发布'""", (user["id"], work_date)
    )
    if assignment is None:
        return jsonify({"message": "只能对已发布的个人班次提交申请"}), 400
    desired_shift_id = data.get("desired_shift_id")
    if request_type == "shift_change":
        try: desired_shift_id = int(desired_shift_id)
        except (TypeError, ValueError): return jsonify({"message": "换班申请请选择目标班次"}), 400
        if not query_one("SELECT shift_id FROM shift_types WHERE shift_id=%s AND is_active=1", (desired_shift_id,)):
            return jsonify({"message": "目标班次不存在或已停用"}), 400
        if desired_shift_id == assignment["shift_id"]:
            return jsonify({"message": "目标班次与当前班次相同，无需提交调整"}), 400
        if query_one("SELECT request_id FROM schedule_requests WHERE employee_id=%s AND work_date=%s AND request_type='leave' AND status='已批准'", (user["id"], work_date)):
            return jsonify({"message": "该日期已经批准请假，不能再申请调整班次"}), 409
    else:
        desired_shift_id = None
    if query_one("SELECT request_id FROM schedule_requests WHERE employee_id=%s AND work_date=%s AND status IN ('待部门审批','待人事审批')", (user["id"], work_date)):
        return jsonify({"message": "该日期已有待审批申请"}), 409
    initial_status = "待人事审批" if user["role"] == "department_admin" else "待部门审批"
    request_id, _ = execute(
        """INSERT INTO schedule_requests
           (employee_id,department,request_type,work_date,desired_shift_id,reason,assignment_id,original_shift_id,assignment_version,status)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
        (user["id"], user["department"], request_type, work_date, desired_shift_id, reason,
         assignment["assignment_id"], assignment["shift_id"], assignment["assignment_version"], initial_status)
    )
    request_name = "换班" if request_type == "shift_change" else "请假"
    if user["role"] == "department_admin":
        recipients = notification_recipient_ids("hr_admin", "super_admin", exclude_id=user["id"])
        target_url = "/pages/admin-schedule-requests.html"
    else:
        recipients = notification_recipient_ids("department_admin", department=user["department"], exclude_id=user["id"])
        target_url = "/pages/department-requests.html"
    notify_users_safely(
        recipients, "schedule_request", f"新的{request_name}申请",
        f"{user['name']}提交了 {work_date} 的{request_name}申请，请及时审批。",
        target_url, related_type="schedule_request", related_id=request_id,
        dedupe_key=f"schedule-request-submitted:{request_id}",
    )
    message = "经理申请已提交给人事部门审批" if user["role"] == "department_admin" else "申请已提交给部门经理审批"
    return jsonify({"success": True, "message": message, "request": serialize_schedule_request(query_one("SELECT * FROM schedule_requests WHERE request_id=%s", (request_id,)))}), 201


@api_bp.route("/api/schedule-requests/mine", methods=["GET"])
@api_operation_log("读取个人班次申请")
@login_required
def get_my_schedule_requests():
    user = current_user()
    records = query_all(
        """SELECT r.*, s.name AS desired_shift_name, m.name AS manager_name, a.name AS admin_name
           FROM schedule_requests r LEFT JOIN shift_types s ON s.shift_id=r.desired_shift_id
           LEFT JOIN employees m ON m.id=r.manager_id LEFT JOIN employees a ON a.id=r.admin_id
           WHERE r.employee_id=%s ORDER BY r.request_id DESC""", (user["id"],)
    )
    return jsonify({"success": True, "requests": [serialize_schedule_request(record) for record in records]})


@api_bp.route("/api/department/schedule-requests", methods=["GET"])
@department_admin_required
def get_department_schedule_requests():
    """返回经理可审批的申请，并提供同部门其他进行中请假的只读概览。"""
    user = current_user()
    if user["role"] == "super_admin":
        return jsonify({"message": "请使用超级管理员审批页面"}), 403
    records = query_all(
        """SELECT r.*, e.name AS employee_name, e.employee_no, s.name AS desired_shift_name
           FROM schedule_requests r JOIN employees e ON e.id=r.employee_id
           LEFT JOIN shift_types s ON s.shift_id=r.desired_shift_id
        WHERE r.department=%s AND r.employee_id<>%s AND r.status='待部门审批' ORDER BY r.created_at""", (user["department"], user["id"])
    )
    # 经理本人申请会直接进入人事审批；它应被看见，但绝不能出现在可由本人审批的列表中。
    other_active_records = query_all(
        """SELECT r.*, e.name AS employee_name, e.employee_no, s.name AS desired_shift_name
           FROM schedule_requests r JOIN employees e ON e.id=r.employee_id
           LEFT JOIN shift_types s ON s.shift_id=r.desired_shift_id
           WHERE r.department=%s AND r.request_type='leave'
             AND r.status IN ('待人事审批','已批准')
           ORDER BY r.work_date, r.request_id DESC""",
        (user["department"],),
    )
    return jsonify({
        "success": True,
        "requests": [serialize_schedule_request(record) for record in records],
        "leave_overview": {
            "awaiting_department_count": sum(1 for record in records if record["request_type"] == "leave"),
            "other_active_count": len(other_active_records),
            "other_active_requests": [serialize_schedule_request(record) for record in other_active_records],
        },
    })


@api_bp.route("/api/department/schedule-requests/<int:request_id>/review", methods=["POST"])
@department_admin_required
def review_department_schedule_request(request_id):
    user, data = current_user(), request.get_json(silent=True) or {}
    if user["role"] == "super_admin":
        return jsonify({"message": "请使用超级管理员审批页面"}), 403
    decision, note = data.get("decision"), str(data.get("note", "")).strip()
    if decision not in ("同意", "驳回"):
        return jsonify({"message": "decision 必须为同意或驳回"}), 400
    record = query_one("SELECT * FROM schedule_requests WHERE request_id=%s AND department=%s AND employee_id<>%s AND status='待部门审批'", (request_id, user["department"], user["id"]))
    if record is None:
        return jsonify({"message": "申请不存在、已处理或不属于本部门"}), 404
    record, plan_id, finalize_error = finalize_schedule_request(
        request_id, "待部门审批", decision, user["id"], "manager", note,
    )
    if finalize_error:
        return jsonify({"message": finalize_error, "conflict": True}), 409
    status = record["status"]
    if plan_id:
        record_schedule_snapshot(plan_id, "员工请假生效" if record["request_type"] == "leave" else "员工班次调整", note)
    request_name = "换班" if record["request_type"] == "shift_change" else "请假"
    notify_users_safely(
        [record["employee_id"]], "schedule_request_result",
        f"{request_name}申请{'已批准' if decision == '同意' else '已驳回'}",
        f"你在 {record['work_date']} 的{request_name}申请{'已批准并同步到日历' if decision == '同意' else '未通过审批'}"
        f"{f'。审批意见：{note}' if note else '。'}",
        "/pages/my-schedule.html", related_type="schedule_request", related_id=request_id,
        dedupe_key=f"schedule-request-result:{request_id}:{status}",
    )
    return jsonify({"success": True, "message": "申请已批准并同步到员工日历" if decision == "同意" else "已驳回该申请"})


@api_bp.route("/api/admin/schedule-requests", methods=["GET"])
@hr_required
def get_admin_schedule_requests():
    records = query_all(
        """SELECT r.*, e.name AS employee_name, e.employee_no, s.name AS desired_shift_name, m.name AS manager_name
           FROM schedule_requests r JOIN employees e ON e.id=r.employee_id
           LEFT JOIN shift_types s ON s.shift_id=r.desired_shift_id LEFT JOIN employees m ON m.id=r.manager_id
           WHERE r.status='待人事审批' ORDER BY r.created_at"""
    )
    return jsonify({"success": True, "requests": [serialize_schedule_request(record) for record in records]})


@api_bp.route("/api/admin/schedule-requests/<int:request_id>/review", methods=["POST"])
@hr_required
def review_admin_schedule_request(request_id):
    user, data = current_user(), request.get_json(silent=True) or {}
    decision, note = data.get("decision"), str(data.get("note", "")).strip()
    if decision not in ("同意", "驳回"):
        return jsonify({"message": "decision 必须为同意或驳回"}), 400
    record = query_one("SELECT * FROM schedule_requests WHERE request_id=%s AND status='待人事审批'", (request_id,))
    if record is None:
        return jsonify({"message": "申请不存在或尚未提交人事审批"}), 404
    record, plan_id, finalize_error = finalize_schedule_request(
        request_id, "待人事审批", decision, user["id"], "admin", note,
    )
    if finalize_error:
        return jsonify({"message": finalize_error, "conflict": True}), 409
    status = record["status"]
    if plan_id:
        record_schedule_snapshot(plan_id, "员工请假生效" if record["request_type"] == "leave" else "员工班次调整", note)
    request_name = "换班" if record["request_type"] == "shift_change" else "请假"
    notify_users_safely(
        [record["employee_id"]], "schedule_request_result",
        f"{request_name}申请{'已批准' if decision == '同意' else '已驳回'}",
        f"你在 {record['work_date']} 的{request_name}申请{'已批准并同步到日历' if decision == '同意' else '未通过审批'}"
        f"{f'。审批意见：{note}' if note else '。'}",
        "/pages/my-schedule.html", related_type="schedule_request", related_id=request_id,
        dedupe_key=f"schedule-request-result:{request_id}:{status}",
    )
    return jsonify({"success": True, "message": "申请已批准并同步到员工日历" if decision == "同意" else "已驳回该申请"})


# ==================== 智能体 OpenAPI 接口 ====================

@api_bp.route("/api/agent/me", methods=["GET"])
@agent_auth_required
def agent_current_user_profile(user):
    """返回智能体请求头指定的当前员工身份和岗位。"""
    return jsonify({"success": True, "user": user})


@api_bp.route("/api/agent/permission-role", methods=["GET"])
def agent_permission_role():
    """为 Harness 校验签名身份令牌的真实角色，不返回业务数据。"""
    user = office_session_user(office_session_token())
    if user is None:
        return jsonify({"message": "智能办公身份凭证无效或已过期"}), 401
    # 技术部需要能处理项目级技术问题，因此与超级管理员共享 Harness
    # 权限切换能力。授权依据来自签名令牌对应的数据库用户，绝不依赖 URL。
    can_switch = user["role"] == "super_admin" or user.get("department") == "技术部"
    return jsonify({
        "success": True,
        "role": user["role"],
        "department": user.get("department", ""),
        "can_switch": can_switch,
    })

@api_bp.route("/api/agent/leave-requests", methods=["POST"])
@agent_auth_required
def agent_submit_leave_request(user):
    """让外部智能体以当前员工身份提交请假申请。"""
    data = request.get_json(silent=True) or {}
    if data.get("confirmed") is not True:
        return jsonify({"message": "请先向用户确认请假日期和原因，再将 confirmed 设置为 true"}), 400
    raw_dates = data.get("work_dates")
    reason = str(data.get("reason", "")).strip()[:500]
    if not isinstance(raw_dates, list) or not raw_dates:
        return jsonify({"message": "work_dates 必须是至少包含一个日期的数组"}), 400
    work_dates = []
    for value in raw_dates[:31]:
        date_value = str(value).strip()
        if not parse_iso_date(date_value):
            return jsonify({"message": f"日期格式无效：{date_value}，应为 YYYY-MM-DD"}), 400
        if date_value not in work_dates:
            work_dates.append(date_value)
    if not reason:
        return jsonify({"message": "请假原因不能为空"}), 400
    try:
        request_ids, error_message = _assistant_submit_leave_requests(user, work_dates, reason)
    except ValueError as error:
        return jsonify({"message": str(error)}), 409
    if error_message:
        return jsonify({"message": error_message}), 403
    return jsonify({"success": True, "request_ids": request_ids, "work_dates": work_dates}), 201

@api_bp.route("/api/agent/leave-requests", methods=["GET"])
@agent_auth_required
def agent_list_leave_requests(user):
    """查询智能体当前操作者可见的本人、本部门或全量请假申请。"""
    scope = request.args.get("scope", "mine").strip().lower()
    status = request.args.get("status", "").strip()
    department = request.args.get("department", "").strip()
    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()
    conditions, parameters = ["r.request_type='leave'"], []
    if scope == "mine":
        conditions.append("r.employee_id=%s")
        parameters.append(user["id"])
    elif scope == "pending":
        if user["role"] == "department_admin":
            conditions.extend(["r.department=%s", "r.employee_id<>%s", "r.status='待部门审批'"])
            parameters.extend([user["department"], user["id"]])
        elif user["role"] in ("hr_admin", "super_admin"):
            conditions.append("r.status='待人事审批'")
        else:
            return jsonify({"message": "当前用户无权查询待审批请假"}), 403
    elif scope == "department":
        if user["role"] == "department_admin":
            department = user["department"]
        elif user["role"] not in ("hr_admin", "super_admin"):
            return jsonify({"message": "当前用户无权查询部门请假"}), 403
        if not department:
            return jsonify({"message": "department 不能为空"}), 400
        conditions.append("r.department=%s")
        parameters.append(department)
    elif scope == "all":
        if user["role"] not in ("hr_admin", "super_admin"):
            return jsonify({"message": "仅人事和超级管理员可以查询全部请假"}), 403
    else:
        return jsonify({"message": "scope 只能是 mine、department、all 或 pending"}), 400
    if status:
        if scope == "pending" and status not in ("待部门审批", "待人事审批"):
            return jsonify({"message": "pending 查询的 status 必须是待部门审批或待人事审批"}), 400
        conditions.append("r.status=%s")
        parameters.append(status)
    if start_date:
        if not parse_iso_date(start_date): return jsonify({"message": "start_date 格式无效"}), 400
        conditions.append("r.work_date>=%s")
        parameters.append(start_date)
    if end_date:
        if not parse_iso_date(end_date): return jsonify({"message": "end_date 格式无效"}), 400
        conditions.append("r.work_date<=%s")
        parameters.append(end_date)
    if start_date and end_date and start_date > end_date:
        return jsonify({"message": "start_date 不能晚于 end_date"}), 400
    where_sql = " AND ".join(conditions)
    rows = query_all(
        """SELECT r.request_id,r.employee_id,e.name AS employee_name,e.employee_no,
                  r.department,r.request_type,r.work_date,r.reason,r.status,r.created_at,
                  r.manager_note,r.admin_note
           FROM schedule_requests r JOIN employees e ON e.id=r.employee_id
           WHERE """ + where_sql + " ORDER BY r.request_id DESC LIMIT 200", parameters,
    )
    summary = query_one(
        "SELECT COUNT(*) AS total, COUNT(DISTINCT r.employee_id) AS employee_count "
        "FROM schedule_requests r WHERE " + where_sql, parameters,
    )
    return jsonify({
        "success": True, "scope": scope, "department": department or None,
        "total": int(summary["total"]), "employee_count": int(summary["employee_count"]),
        "requests": [serialize_schedule_request(row) for row in rows],
    })


@api_bp.route("/api/agent/tasks", methods=["GET"])
@agent_auth_required
def agent_list_tasks(user):
    """查询智能体当前操作者自己的待办事项。"""
    status = request.args.get("status", "").strip()
    conditions, parameters = ["employee_id=%s"], [user["id"]]
    if status:
        conditions.append("status=%s")
        parameters.append(status)
    tasks = query_all(
        """SELECT task_id,task_type,title,content,status,priority,due_date,
                  created_at,submitted_at,completed_at
           FROM employee_tasks WHERE """ + " AND ".join(conditions) + " ORDER BY task_id DESC LIMIT 200", parameters,
    )
    return jsonify({"success": True, "tasks": tasks})


@api_bp.route("/api/agent/tasks/manual", methods=["POST"])
@agent_auth_required
def agent_publish_manual_tasks(user):
    """让外部智能体以当前操作者身份发布人工待办，范围与角色权限一致。"""
    data = request.get_json(silent=True) or {}
    confirmation_error = _agent_confirmation(data, "发布人工待办")
    if confirmation_error:
        return confirmation_error
    return publish_manual_tasks(user, data)


def _agent_review_leave_request(user, request_id, decision, note):
    """统一执行智能体请假审批，确保和页面审批使用同一事务状态机。"""
    if user["role"] == "department_admin":
        expected_status, reviewer_kind = "待部门审批", "manager"
        record = query_one(
            "SELECT * FROM schedule_requests WHERE request_id=%s AND department=%s AND employee_id<>%s AND status=%s",
            (request_id, user["department"], user["id"], expected_status),
        )
    elif user["role"] in ("hr_admin", "super_admin"):
        expected_status, reviewer_kind = "待人事审批", "admin"
        record = query_one(
            "SELECT * FROM schedule_requests WHERE request_id=%s AND status=%s",
            (request_id, expected_status),
        )
    else:
        return jsonify({"message": "当前用户无权审批请假"}), 403
    if record is None or record["request_type"] != "leave":
        return jsonify({"message": "请假申请不存在、已处理或不属于当前审批范围"}), 404
    record, plan_id, error_message = finalize_schedule_request(
        request_id, expected_status, decision, user["id"], reviewer_kind, note,
    )
    if error_message:
        return jsonify({"message": error_message, "conflict": True}), 409
    if plan_id:
        record_schedule_snapshot(plan_id, "员工请假生效", note)
    notify_users_safely(
        [record["employee_id"]], "schedule_request_result",
        f"请假申请{'已批准' if decision == '同意' else '已驳回'}",
        f"你在 {record['work_date']} 的请假申请{'已批准并同步到日历' if decision == '同意' else '未通过审批'}"
        f"{f'。审批意见：{note}' if note else '。'}",
        "/pages/my-schedule.html", related_type="schedule_request", related_id=request_id,
        dedupe_key=f"schedule-request-result:{request_id}:{record['status']}",
    )
    app.logger.info("智能体审批请假 request_id=%s actor=%s decision=%s", request_id, user["employee_no"], decision)
    return jsonify({"success": True, "message": "请假已批准并同步到员工日历" if decision == "同意" else "请假已驳回", "request": serialize_schedule_request(record)})


@api_bp.route("/api/agent/leave-requests/<int:request_id>/approve", methods=["POST"])
@agent_auth_required
def agent_approve_leave_request(user, request_id):
    data = request.get_json(silent=True) or {}
    if data.get("confirmed") is not True:
        return jsonify({"message": "请在用户明确确认后，将 confirmed 设置为 true"}), 400
    return _agent_review_leave_request(user, request_id, "同意", str(data.get("note", "")).strip()[:500])


@api_bp.route("/api/agent/leave-requests/<int:request_id>/reject", methods=["POST"])
@agent_auth_required
def agent_reject_leave_request(user, request_id):
    data = request.get_json(silent=True) or {}
    if data.get("confirmed") is not True:
        return jsonify({"message": "请在用户明确确认后，将 confirmed 设置为 true"}), 400
    return _agent_review_leave_request(user, request_id, "驳回", str(data.get("note", "")).strip()[:500])


def _agent_confirmation(data, action):
    """拒绝未经用户确认的智能办公写入操作。"""
    if not isinstance(data, dict) or data.get("confirmed") is not True:
        return jsonify({"message": f"{action} 前必须先获得用户明确确认，并将 confirmed 设置为 true"}), 400
    return None


def _agent_schedule_rows(user, scope):
    """按智能办公身份返回可见的班次申请，查询范围不由模型决定。"""
    conditions, parameters = [], []
    if scope == "mine":
        conditions.append("r.employee_id=%s")
        parameters.append(user["id"])
    elif scope == "pending":
        if user["role"] == "department_admin":
            conditions.extend(["r.department=%s", "r.employee_id<>%s", "r.status='待部门审批'"])
            parameters.extend([user["department"], user["id"]])
        elif user["role"] in ("hr_admin", "super_admin"):
            conditions.append("r.status='待人事审批'")
        else:
            return None, "当前用户无权查看待审批班次申请"
    elif scope == "department":
        if user["role"] == "department_admin":
            conditions.append("r.department=%s")
            parameters.append(user["department"])
        elif user["role"] in ("hr_admin", "super_admin"):
            department = request.args.get("department", "").strip()
            if not department:
                return None, "人事或超级管理员查询部门申请时必须提供 department"
            conditions.append("r.department=%s")
            parameters.append(department)
        else:
            return None, "当前用户无权查看部门班次申请"
    else:
        return None, "scope 只能是 mine、pending 或 department"
    rows = query_all(
        """SELECT r.*, e.name AS employee_name, e.employee_no, s.name AS desired_shift_name,
                  m.name AS manager_name, a.name AS admin_name
           FROM schedule_requests r JOIN employees e ON e.id=r.employee_id
           LEFT JOIN shift_types s ON s.shift_id=r.desired_shift_id
           LEFT JOIN employees m ON m.id=r.manager_id LEFT JOIN employees a ON a.id=r.admin_id
           WHERE """ + " AND ".join(conditions) + " ORDER BY r.request_id DESC LIMIT 200",
        parameters,
    )
    return rows, None


@api_bp.route("/api/agent/overview", methods=["GET"])
@agent_auth_required
def agent_overview(user):
    """返回当前员工和其岗位真正可见的业务摘要，供每个智能办公会话初始化。"""
    summary = {
        "unread_notifications": int(query_one(
            "SELECT COUNT(*) AS total FROM notifications WHERE recipient_id=%s AND is_read=0", (user["id"],)
        )["total"]),
        "open_tasks": int(query_one(
            "SELECT COUNT(*) AS total FROM employee_tasks WHERE employee_id=%s AND status='待处理'", (user["id"],)
        )["total"]),
        "pending_leave_requests": int(query_one(
            "SELECT COUNT(*) AS total FROM schedule_requests WHERE employee_id=%s AND request_type='leave' "
            "AND status IN ('待部门审批','待人事审批')", (user["id"],)
        )["total"]),
        "pending_schedule_requests": int(query_one(
            "SELECT COUNT(*) AS total FROM schedule_requests WHERE employee_id=%s "
            "AND status IN ('待部门审批','待人事审批')", (user["id"],)
        )["total"]),
    }
    approvals = {}
    if user["role"] == "department_admin":
        approvals["department_schedule_requests"] = int(query_one(
            "SELECT COUNT(*) AS total FROM schedule_requests WHERE department=%s AND employee_id<>%s "
            "AND status='待部门审批'", (user["department"], user["id"])
        )["total"])
        approvals["department_offboarding_requests"] = int(query_one(
            """SELECT COUNT(*) AS total FROM employee_tasks t JOIN employees e ON e.id=t.employee_id
               WHERE e.department=%s AND e.id<>%s AND t.task_type='offboarding_request' AND t.status='待部门审批'""",
            (user["department"], user["id"]),
        )["total"])
    elif user["role"] == "hr_admin":
        approvals["schedule_requests"] = int(query_one(
            "SELECT COUNT(*) AS total FROM schedule_requests WHERE status='待人事审批'"
        )["total"])
        approvals["offboarding_handovers"] = int(query_one(
            "SELECT COUNT(*) AS total FROM employee_tasks WHERE task_type='offboarding_handover' AND status='待人事审核'"
        )["total"])
    elif user["role"] == "super_admin":
        approvals["schedule_requests"] = int(query_one(
            "SELECT COUNT(*) AS total FROM schedule_requests WHERE status='待人事审批'"
        )["total"])
        approvals["admin_todos"] = int(query_one(
            "SELECT COUNT(*) AS total FROM admin_todos WHERE status='待确认'"
        )["total"])
    return jsonify({"success": True, "user": user, "summary": summary, "approvals": approvals})


@api_bp.route("/api/agent/notifications", methods=["GET"])
@agent_auth_required
def agent_list_notifications(user):
    """读取当前智能办公身份收到的站内消息。"""
    unread_only = request.args.get("unread_only", "").strip().lower() in ("1", "true", "yes")
    try:
        limit = min(max(int(request.args.get("limit", "50")), 1), 100)
    except (TypeError, ValueError):
        return jsonify({"message": "limit 必须是 1 到 100 的整数"}), 400
    conditions, parameters = ["recipient_id=%s"], [user["id"]]
    if unread_only:
        conditions.append("is_read=0")
    rows = query_all(
        "SELECT * FROM notifications WHERE " + " AND ".join(conditions) + " ORDER BY created_at DESC,notification_id DESC LIMIT %s",
        [*parameters, limit],
    )
    unread_count = query_one(
        "SELECT COUNT(*) AS total FROM notifications WHERE recipient_id=%s AND is_read=0", (user["id"],)
    )["total"]
    return jsonify({"success": True, "unread_count": int(unread_count), "notifications": [serialize_notification(row) for row in rows]})


@api_bp.route("/api/agent/notifications/<int:notification_id>/read", methods=["POST"])
@agent_auth_required
def agent_read_notification(user, notification_id):
    """将当前员工自己的站内消息标记为已读。"""
    confirmation_error = _agent_confirmation(request.get_json(silent=True), "标记消息已读")
    if confirmation_error:
        return confirmation_error
    record = query_one("SELECT * FROM notifications WHERE notification_id=%s AND recipient_id=%s", (notification_id, user["id"]))
    if record is None:
        return jsonify({"message": "消息不存在或不属于当前用户"}), 404
    if not record["is_read"]:
        execute("UPDATE notifications SET is_read=1,read_at=NOW() WHERE notification_id=%s AND recipient_id=%s", (notification_id, user["id"]))
    record = query_one("SELECT * FROM notifications WHERE notification_id=%s AND recipient_id=%s", (notification_id, user["id"]))
    return jsonify({"success": True, "notification": serialize_notification(record)})


@api_bp.route("/api/agent/tasks/<int:task_id>/submit", methods=["POST"])
@agent_auth_required
def agent_submit_task(user, task_id):
    """以当前员工身份提交自己的待办，复用网页端的状态流转。"""
    data = request.get_json(silent=True) or {}
    confirmation_error = _agent_confirmation(data, "提交待办")
    if confirmation_error:
        return confirmation_error
    task = query_one("SELECT * FROM employee_tasks WHERE task_id=%s AND employee_id=%s", (task_id, user["id"]))
    if task is None:
        return jsonify({"message": "待办不存在或不属于当前用户"}), 404
    if task["status"] != "待处理":
        return jsonify({"message": "该待办已提交或已处理"}), 409
    content = str(data.get("content", "员工已提交处理结果")).strip()[:2000]
    if task["task_type"] in ("offboarding_handover", "handover_request"):
        _, changed = execute(
            "UPDATE employee_tasks SET status='待人事审核', content=CONCAT(content, '\n员工交接说明：', %s), submitted_at=NOW() "
            "WHERE task_id=%s AND employee_id=%s AND status='待处理'", (content, task_id, user["id"])
        )
        message = "交接已提交人事审核"
    elif task["task_type"] == "manual":
        _, changed = execute(
            "UPDATE employee_tasks SET status='已完成', content=CONCAT(content, '\n员工处理结果：', %s), submitted_at=NOW(), completed_at=NOW() "
            "WHERE task_id=%s AND employee_id=%s AND status='待处理'", (content, task_id, user["id"])
        )
        if changed == 1 and task.get("publisher_id") and task["publisher_id"] != user["id"]:
            notify_users_safely([task["publisher_id"]], "task_completed", "员工已完成待办", f"{user['name']} 已提交：{task['title']}", "/pages/task-center.html", related_type="employee_task", related_id=task_id, dedupe_key=f"manual-task-completed:{task_id}")
        message = "待办已完成并提交结果"
    else:
        _, changed = execute(
            "UPDATE employee_tasks SET status='已提交', content=CONCAT(content, '\n员工反馈：', %s), submitted_at=NOW() "
            "WHERE task_id=%s AND employee_id=%s AND status='待处理'", (content, task_id, user["id"])
        )
        if changed == 1:
            employee = find_employee(user["id"])
            execute("INSERT INTO admin_todos (task_id, employee_id, employee_name, title, content) VALUES (%s,%s,%s,%s,%s)", (task_id, user["id"], employee["name"], task["title"], content))
        message = "已提交给超级管理员确认"
    if changed != 1:
        return jsonify({"message": "该待办刚刚已被处理，请刷新后重试"}), 409
    return jsonify({"success": True, "message": message})


@api_bp.route("/api/agent/my-schedule", methods=["GET"])
@agent_auth_required
def agent_my_schedule(user):
    """读取当前员工自己的已发布班次和已批准请假。"""
    start, end = request.args.get("start", "").strip(), request.args.get("end", "").strip()
    if not parse_iso_date(start) or not parse_iso_date(end) or start > end:
        return jsonify({"message": "请提供有效的开始和结束日期"}), 400
    schedules = query_all(
        """SELECT a.assignment_id,a.work_date,s.shift_id,s.name AS shift_name,s.start_time,s.end_time,s.color
           FROM schedule_assignments a JOIN schedule_plans p ON p.plan_id=a.plan_id JOIN shift_types s ON s.shift_id=a.shift_id
           WHERE a.employee_id=%s AND p.status='已发布' AND a.work_date BETWEEN %s AND %s ORDER BY a.work_date""",
        (user["id"], start, end),
    )
    leave_days = query_all(
        "SELECT work_date,request_id FROM schedule_requests WHERE employee_id=%s AND request_type='leave' AND status='已批准' AND work_date BETWEEN %s AND %s",
        (user["id"], start, end),
    )
    for schedule in schedules:
        if hasattr(schedule.get("work_date"), "isoformat"):
            schedule["work_date"] = schedule["work_date"].isoformat()
        for field in ("start_time", "end_time"):
            value = schedule[field]
            if hasattr(value, "total_seconds"):
                seconds = int(value.total_seconds())
                schedule[field] = f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}"
    for leave_day in leave_days:
        if hasattr(leave_day.get("work_date"), "isoformat"):
            leave_day["work_date"] = leave_day["work_date"].isoformat()
    return jsonify({"success": True, "schedules": schedules, "leave_days": leave_days})


@api_bp.route("/api/agent/schedule-requests", methods=["GET"])
@agent_auth_required
def agent_list_schedule_requests(user):
    """读取当前岗位被允许查看的请假和换班申请。"""
    scope = request.args.get("scope", "mine").strip().lower()
    rows, error_message = _agent_schedule_rows(user, scope)
    if error_message:
        return jsonify({"message": error_message}), 403 if "无权" in error_message else 400
    return jsonify({"success": True, "scope": scope, "requests": [serialize_schedule_request(row) for row in rows]})


def _agent_review_schedule_request(user, request_id, decision, note):
    """按当前审批角色处理其范围内的请假或换班申请。"""
    if user["role"] == "department_admin":
        expected_status, reviewer_kind = "待部门审批", "manager"
        record = query_one("SELECT * FROM schedule_requests WHERE request_id=%s AND department=%s AND employee_id<>%s AND status=%s", (request_id, user["department"], user["id"], expected_status))
    elif user["role"] in ("hr_admin", "super_admin"):
        expected_status, reviewer_kind = "待人事审批", "admin"
        record = query_one("SELECT * FROM schedule_requests WHERE request_id=%s AND status=%s", (request_id, expected_status))
    else:
        return jsonify({"message": "当前用户无权审批班次申请"}), 403
    if record is None:
        return jsonify({"message": "申请不存在、已处理或不属于当前审批范围"}), 404
    record, plan_id, error_message = finalize_schedule_request(request_id, expected_status, decision, user["id"], reviewer_kind, note)
    if error_message:
        return jsonify({"message": error_message, "conflict": True}), 409
    if plan_id:
        record_schedule_snapshot(plan_id, "员工请假生效" if record["request_type"] == "leave" else "员工班次调整", note)
    request_name = "换班" if record["request_type"] == "shift_change" else "请假"
    notify_users_safely([record["employee_id"]], "schedule_request_result", f"{request_name}申请{'已批准' if decision == '同意' else '已驳回'}", f"你在 {record['work_date']} 的{request_name}申请{'已批准并同步到日历' if decision == '同意' else '未通过审批'}{f'。审批意见：{note}' if note else '。'}", "/pages/my-schedule.html", related_type="schedule_request", related_id=request_id, dedupe_key=f"schedule-request-result:{request_id}:{record['status']}")
    app.logger.info("智能体审批班次申请 request_id=%s actor=%s decision=%s", request_id, user["employee_no"], decision)
    return jsonify({"success": True, "message": "申请已批准并同步到员工日历" if decision == "同意" else "申请已驳回", "request": serialize_schedule_request(record)})


@api_bp.route("/api/agent/schedule-requests/<int:request_id>/approve", methods=["POST"])
@agent_auth_required
def agent_approve_schedule_request(user, request_id):
    data = request.get_json(silent=True) or {}
    confirmation_error = _agent_confirmation(data, "审批班次申请")
    if confirmation_error:
        return confirmation_error
    return _agent_review_schedule_request(user, request_id, "同意", str(data.get("note", "")).strip()[:500])


@api_bp.route("/api/agent/schedule-requests/<int:request_id>/reject", methods=["POST"])
@agent_auth_required
def agent_reject_schedule_request(user, request_id):
    data = request.get_json(silent=True) or {}
    confirmation_error = _agent_confirmation(data, "驳回班次申请")
    if confirmation_error:
        return confirmation_error
    return _agent_review_schedule_request(user, request_id, "驳回", str(data.get("note", "")).strip()[:500])


app.register_blueprint(api_bp)
if AUTO_INITIALIZE_DATABASE:
    initialize_database()

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=APP_PORT, debug=APP_DEBUG)
